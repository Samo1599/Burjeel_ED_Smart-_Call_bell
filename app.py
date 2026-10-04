import os, json, hashlib, secrets, base64, io, re, traceback, asyncio
from datetime import datetime, timezone, timedelta
from typing import Optional
from fastapi import FastAPI, Request, Depends, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse, Response
from starlette.middleware.sessions import SessionMiddleware
from itsdangerous import URLSafeTimedSerializer, BadSignature, SignatureExpired
from jinja2 import Environment, DictLoader, select_autoescape
import qrcode
import qrcode.image.svg
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from sqlalchemy import create_engine, Column, Integer, String, Boolean, DateTime, ForeignKey, Text, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import declarative_base, relationship, sessionmaker, Session

try:
    from pywebpush import webpush, WebPushException
except Exception:
    webpush = None
    WebPushException = Exception

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_URL = os.getenv('DATABASE_URL', f"sqlite:///{os.path.join(BASE_DIR,'callbell.db')}")
if DB_URL.startswith('postgres://'): DB_URL = DB_URL.replace('postgres://','postgresql+psycopg://',1)
if DB_URL.startswith('postgresql://'): DB_URL = DB_URL.replace('postgresql://','postgresql+psycopg://',1)
if DB_URL.startswith('postgresql+psycopg://'):
    engine = create_engine(
        DB_URL,
        connect_args={
            "options":"-csearch_path=burjeel_ed_call",
            "connect_timeout":10,
            "keepalives":1,
            "keepalives_idle":30,
            "keepalives_interval":10,
            "keepalives_count":3
        },
        pool_pre_ping=True,
        pool_recycle=120,
        pool_use_lifo=True,
        pool_size=5,
        max_overflow=5,
        pool_timeout=20,
        pool_reset_on_return='rollback',
        future=True
    )
    with engine.connect().execution_options(isolation_level="AUTOCOMMIT") as conn:
        conn.execute(text("CREATE SCHEMA IF NOT EXISTS burjeel_ed_call"))
else:
    engine = create_engine(DB_URL, connect_args={'check_same_thread':False} if DB_URL.startswith('sqlite') else {}, future=True)
SessionLocal = sessionmaker(bind=engine, expire_on_commit=False, future=True)
Base = declarative_base()
SLA_SECONDS = int(os.getenv('SLA_SECONDS','120'))
VAPID_PUBLIC_KEY = os.getenv('VAPID_PUBLIC_KEY','')
VAPID_PRIVATE_KEY = os.getenv('VAPID_PRIVATE_KEY','')
VAPID_SUBJECT = os.getenv('VAPID_SUBJECT','mailto:it@example.com')

app = FastAPI(title='Burjeel ED Smart Call')
SESSION_SECRET=os.getenv('SECRET_KEY','dev-secret-change-me')
app.add_middleware(SessionMiddleware, secret_key=SESSION_SECRET)
push_receipt_signer=URLSafeTimedSerializer(SESSION_SECRET,salt='push-display-receipt-v1')

TEMPLATES = {
    "base.html": "<!doctype html><html lang=\"en\"><head><meta charset=\"utf-8\"><meta name=\"viewport\" content=\"width=device-width,initial-scale=1,viewport-fit=cover\"><title>{% block title %}Burjeel ED Call{% endblock %}</title><link rel=\"manifest\" href=\"/manifest.webmanifest\"><meta name=\"theme-color\" content=\"#004f9e\"><link rel=\"apple-touch-icon\" href=\"/static/icons/icon-192.png\"><link rel=\"stylesheet\" href=\"/static/app.css\"></head><body class=\"{% block body_class %}{% endblock %}\"><header class=\"topbar\"><div class=\"brand\"><img src=\"/static/icons/icon-192.png\" alt=\"logo\"><div><b>Burjeel ED Call</b><span>Smart Call Bell</span></div></div>{% if current_user %}<div class=\"userbox\"><span>{{ current_user.name }}</span><small>{{ role_labels.get(current_user.role,current_user.role.replace('_',' ')|title) }}</small><a href=\"/logout\">Logout</a></div>{% endif %}</header>{% if current_user %}<nav class=\"nav\">{% if 'my_rooms' in current_permissions %}<a href=\"/nurse\">My Rooms</a>{% endif %}{% if 'live_board' in current_permissions %}<a href=\"/charge\">Live Board</a>{% endif %}{% if 'wallboard' in current_permissions %}<a href=\"/wallboard\">📺 Call Bell Screen</a>{% endif %}{% if 'management' in current_permissions %}<a href=\"/manager\">Management</a>{% endif %}{% if 'admin_control' in current_permissions %}<a href=\"/admin\">Admin</a>{% endif %}\n<div class=\"nav-alert-wrap\">\n  <button id=\"alertsNavButton\" class=\"nav-alert-btn\" type=\"button\" onclick=\"toggleAlertsDropdown(event)\">🔔 <span>Alerts</span> <b id=\"alertsBadge\" hidden>0</b></button>\n  <div id=\"alertsDropdown\" class=\"alerts-dropdown\" hidden>\n    <div class=\"alerts-drop-head\"><div><b>Alerts</b><small id=\"alertsDropStatus\">Recent notifications</small></div><button type=\"button\" onclick=\"markAllAlertsRead()\">Mark all read</button></div>\n    <div id=\"alertsDropList\" class=\"alerts-drop-list\"><div class=\"alerts-empty\">Loading…</div></div>\n    <a class=\"alerts-view-all\" href=\"/alerts\">View all alerts</a>\n  </div>\n</div></nav>{% endif %}<main class=\"container\">{% block content %}{% endblock %}</main><script>window.VAPID_PUBLIC_KEY='{{ vapid_public_key|default('') }}';</script><script src=\"/static/app.js\"></script>{% block scripts %}{% endblock %}</body></html>",
    "login.html": "{% extends 'base.html' %}{% block title %}Sign in - Burjeel ED Call{% endblock %}\n{% block content %}\n<section class=\"login-shell\"><div class=\"login-card\">\n<img class=\"hero-logo\" src=\"/static/icons/icon-512.png\"><h1>ED Smart Call Bell</h1><p>Secure staff access</p>\n{% if error %}<div class=\"alert danger\">{{ error }}</div>{% endif %}\n<form method=\"post\" action=\"/login\">\n<label>Email<input name=\"email\" type=\"email\" required placeholder=\"sara@demo.local\"></label>\n<label>Password<div class=\"password-wrap\"><input id=\"loginPassword\" name=\"password\" type=\"password\" required placeholder=\"••••••••\"><button type=\"button\" class=\"password-eye\" aria-label=\"Show password\" onclick=\"togglePassword('loginPassword',this)\">👁</button></div></label>\n<button class=\"btn primary wide\">Sign in</button>\n</form>\n<div class=\"login-security-note\">🔔 Nurse and Charge Nurse accounts verify notifications before entering the live workspace.</div>\n<div class=\"demo\"><b>Demo</b><span>sara@demo.local / Demo123!</span><span>charge@demo.local / Demo123!</span><span>manager@demo.local / Demo123!</span></div>\n</div></section>\n{% endblock %}",
    "notification_setup.html": "{% extends 'base.html' %}{% block title %}Notification Verification - Burjeel ED Call{% endblock %}\n{% block content %}\n<section class=\"notify-mini-shell\">\n  <div class=\"notify-mini-card\">\n    <div class=\"notify-mini-title\">🔔 <b>Notifications Required</b></div>\n    <div class=\"notify-mini-sub\">Notifications are required. Checking browser permission...</div>\n    <div id=\"iosInstallHint\" class=\"notify-platform-hint\" style=\"display:none\">📱 On iPhone/iPad, add Burjeel ED Call to the Home Screen and open it from the app icon before enabling notifications.</div>\n    <div class=\"notify-mini-steps\">\n      <div class=\"notify-mini-chip\" id=\"stepPermission\"><span class=\"chip-dot\"></span><b>Permission</b><small class=\"step-state\">Waiting</small></div>\n      <div class=\"notify-mini-chip\" id=\"stepSubscription\"><span class=\"chip-dot\"></span><b>Push Subscription</b><small class=\"step-state\">Waiting</small></div>\n      <div class=\"notify-mini-chip\" id=\"stepTest\"><span class=\"chip-dot\"></span><b>Test Alert</b><small class=\"step-state\">Waiting</small></div>\n      <div class=\"notify-mini-chip\" id=\"stepConfirmed\"><span class=\"chip-dot\"></span><b>Device Confirmed</b><small class=\"step-state\">Waiting</small></div>\n    </div>\n  </div>\n  <div id=\"notifyStatus\" class=\"notify-mini-status\">Preparing notification check...</div>\n  <button id=\"verifyBtn\" class=\"notify-verify-btn\" onclick=\"verifyNotifications()\">Verifying Notifications...</button>\n  <button id=\"continueBtn\" class=\"notify-verify-btn ready\" style=\"display:none\" onclick=\"location.href='{{next_url}}'\">Notifications Ready — Continue</button>\n</section>\n{% endblock %}\n{% block scripts %}<script>\nfunction markStep(id,state,text){const el=document.getElementById(id);el.classList.remove('ok','bad','working');el.classList.add(state);el.querySelector('.step-state').textContent=text}\nfunction isIOS(){return /iPad|iPhone|iPod/.test(navigator.userAgent)||(navigator.platform==='MacIntel'&&navigator.maxTouchPoints>1)}\nfunction isStandalone(){return window.matchMedia('(display-mode: standalone)').matches||window.navigator.standalone===true}\nasync function verifyNotifications(){\n const btn=document.getElementById('verifyBtn'),status=document.getElementById('notifyStatus'),iosHint=document.getElementById('iosInstallHint');\n btn.disabled=true;btn.textContent='Verifying Notifications...';\n try{\n  if(isIOS()&&!isStandalone()){iosHint.style.display='block';throw new Error('iPhone/iPad Web Push requires the Home Screen web app. Add this site to Home Screen, open it there, then retry.')}\n  if(!('serviceWorker' in navigator)||!('PushManager' in window)||!('Notification' in window))throw new Error('Web Push is not supported in this browser/device mode.');\n  markStep('stepPermission','working','Checking');\n  let p=Notification.permission;\n  if(p!=='granted')p=await Notification.requestPermission();\n  if(p!=='granted'){markStep('stepPermission','bad','Blocked');throw new Error('Notification permission is blocked. Enable notifications for this site, then retry.')}\n  markStep('stepPermission','ok','Allowed');\n  markStep('stepSubscription','working','Registering');\n  await navigator.serviceWorker.register('/sw.js',{updateViaCache:'none'});const reg=await navigator.serviceWorker.ready;\n  const key=''+(window.VAPID_PUBLIC_KEY||'');if(!key)throw new Error('Push keys are not configured on the server.');\n  const expectedKey=urlBase64ToUint8Array(key);\n  let sub=await reg.pushManager.getSubscription();\n  if(sub){\n    try{\n      const dsr=await fetch('/api/push/device-status',{method:'POST',cache:'no-store',headers:{'Content-Type':'application/json'},body:JSON.stringify({endpoint:sub.endpoint,device:pushDeviceMeta()})});\n      if(dsr.ok){\n        const dsd=await dsr.json();\n        if(dsd.verified&&sub.options&&sub.options.applicationServerKey&&new Uint8Array(sub.options.applicationServerKey).length===expectedKey.length&&new Uint8Array(sub.options.applicationServerKey).every((v,i)=>v===expectedKey[i])){\n          markStep('stepPermission','ok','Allowed');\n          markStep('stepSubscription','ok','Registered');\n          markStep('stepTest','ok','Previously verified');\n          markStep('stepConfirmed','ok','Confirmed');\n          status.textContent='✓ Notifications are already active on this device.';\n          btn.style.display='none';\n          document.getElementById('continueBtn').style.display='block';\n          return;\n        }\n      }\n    }catch(_){}\n  }\n  let replaceSub=false;\n  if(sub&&sub.options&&sub.options.applicationServerKey){\n    const current=new Uint8Array(sub.options.applicationServerKey);\n    replaceSub=current.length!==expectedKey.length||current.some((v,i)=>v!==expectedKey[i]);\n  }\n  if(sub&&replaceSub){try{await sub.unsubscribe()}catch(_){};sub=null}\n  if(!sub){sub=await reg.pushManager.subscribe({userVisibleOnly:true,applicationServerKey:expectedKey});await new Promise(r=>setTimeout(r,1200));}\n  const sr=await fetch('/api/push/subscribe',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({subscription:sub.toJSON?sub.toJSON():sub,device:pushDeviceMeta()})});\n  if(!sr.ok)throw new Error('Could not register this browser for push.');\n  markStep('stepSubscription','ok','Registered');\n  markStep('stepTest','working','Sending');markStep('stepConfirmed','working','Checking');\n  const tr=await fetch('/api/push/test',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({endpoint:sub.endpoint,device:pushDeviceMeta()})});\n  let td={};const raw=await tr.text();try{td=raw?JSON.parse(raw):{}}catch(_){td={ok:false,error:'Server returned an invalid response during notification test.'}}\n  if(!tr.ok||!td.ok){markStep('stepTest','bad','Failed');const msg=(td.error||'Test notification failed.')+(td.detail?' ('+td.detail+')':'');throw new Error(msg)}\n  markStep('stepTest','ok','Sent');\n  const token=td.verification_token;status.textContent='Test alert sent. Waiting for this device to confirm receipt...';\n  let verified=false;\n  for(let i=0;i<18;i++){await new Promise(r=>setTimeout(r,1000));const rr=await fetch('/api/push/verification/'+encodeURIComponent(token)+'?mode='+encodeURIComponent(pushDeviceMeta().display_mode)+'&platform='+encodeURIComponent(pushDeviceMeta().platform||''),{cache:'no-store'});if(rr.ok){const rd=await rr.json();if(rd.verified){verified=true;break}}}\n  if(!verified){markStep('stepConfirmed','bad','Not confirmed');throw new Error('The push service accepted the test, but this device did not receive/confirm it. Check OS/browser notification settings and retry.')}\n  markStep('stepConfirmed','ok','Confirmed');status.textContent='✓ Notifications are active for this '+(pushDeviceMeta().display_mode==='pwa'?'installed PWA':'browser')+'.';btn.style.display='none';document.getElementById('continueBtn').style.display='block';\n }catch(e){status.textContent=e.message||'Notification verification failed.';btn.disabled=false;btn.textContent='Retry Notification Setup'}\n}\nverifyNotifications();\n</script>{% endblock %}",
    "nurse.html": "{% extends 'base.html' %}{% block title %}My Rooms - Burjeel ED Call{% endblock %}{% block content %}<div id=\"nurse-live-root\" data-refresh-key=\"{{refresh_key}}\"><div class=\"page-head\"><div><div class=\"eyebrow\">Nurse workspace</div><h1>My Rooms</h1><p>Assigned rooms and live patient calls.</p></div><div class=\"push-controls\"><span id=\"pushHealthBadge\" class=\"push-health checking\">Checking notifications…</span><button class=\"btn secondary notification-verify\" data-notification-button data-notification-state=\"verify\" onclick=\"enablePush()\">🟠 Verify Notifications</button></div></div><div class=\"stats\"><div class=\"stat\"><b>{{ rooms|length }}</b><span>Assigned rooms</span></div><div class=\"stat red\"><b>{{ calls|length }}</b><span>Active calls</span></div><div class=\"stat amber\"><b>{{ pending_handovers|length }}</b><span>Pending handovers</span></div></div>{% if pending_handovers %}<section class=\"panel\"><h2>Handover requests</h2>{% for h in pending_handovers %}<div class=\"list-row\"><b>{{h.room.code}}</b><span>From {{h.from_nurse.name}}</span><button class=\"btn primary\" onclick=\"acceptHandover({{h.id}})\">Accept</button></div>{% endfor %}</section><br>{% endif %}<section class=\"grid rooms\">{% for room in rooms %}<article class=\"room-card {% if calls|selectattr('room_id','equalto',room.id)|list %}hot{% endif %}\"><div class=\"room-top\"><b>{{ room.code }}</b><span>{{ room.zone }}</span></div>{% set rcalls = calls|selectattr('room_id','equalto',room.id)|list %}{% if rcalls %}{% set c=rcalls[0] %}<div class=\"call-status {{ c.status }}\">{{ c.status.replace('_',' ')|upper }}</div>{% set c_elapsed=elapsed_for_call(c) %}<div class=\"timer\" data-created=\"{{utc_iso(c.created_at)}}\" data-elapsed=\"{{c_elapsed}}\" {% if c.arrived_at or c.resolved_at %}data-stop=\"{{utc_iso(c.arrived_at or c.resolved_at)}}\"{% endif %} data-call-id=\"{{c.id}}\">{{'%02d:%02d'|format(c_elapsed//60,c_elapsed%60)}}</div><p>{{ c.reason }}</p><div class=\"actions\">{% if c.status in ['new','escalated'] %}<button class=\"btn primary\" onclick=\"callAction({{c.id}},'ack')\">Acknowledge</button>{% endif %}{% if c.status in ['acknowledged','taken_over','escalated'] %}<button class=\"btn success\" onclick=\"callAction({{c.id}},'arrive')\">Arrived</button>{% endif %}</div>{% else %}<div class=\"ready\">● Ready</div><p class=\"muted\">No active patient call</p>{% endif %}{% if colleagues %}<div class=\"handover-box\"><select id=\"handover-{{room.id}}\"><option value=\"\">Hand over to…</option>{% for n in colleagues %}<option value=\"{{n.id}}\">{{n.name}}</option>{% endfor %}</select><button class=\"btn secondary\" onclick=\"handoverRoom({{room.id}})\">Handover</button></div>{% endif %}</article>{% endfor %}</section></div>{% endblock %}{% block scripts %}<script>startSilentRefresh('#nurse-live-root',3000);</script>{% endblock %}",
    "charge.html": "{% extends 'base.html' %}{% block title %}Charge Nurse Live Board{% endblock %}{% block body_class %}charge-board-page{% endblock %}\n{% block content %}\n<div id=\"charge-live-root\" class=\"charge-command\" data-refresh-key=\"{{refresh_key}}\">\n  <header class=\"charge-command-head\">\n    <div>\n      <div class=\"eyebrow\">CHARGE NURSE COMMAND</div>\n      <div class=\"command-title-row\"><h1>Live Board</h1><span class=\"live-pill\">● Live</span></div>\n      <p>All active calls, escalations, ownership and room assignment.</p>\n    </div>\n    <div class=\"command-actions\">\n      <button class=\"btn secondary notification-verify\" data-notification-button data-notification-state=\"verify\" onclick=\"enablePush()\">🟠 Verify Notifications</button>\n      <button class=\"btn primary\" onclick=\"toggleChargeFullscreen()\">⛶ Full screen</button>\n    </div>\n  </header>\n\n  <div class=\"charge-kpis\">\n    <div class=\"command-kpi\"><span>🛏</span><div><b>{{rooms|length}}</b><small>Rooms</small></div></div>\n    <div class=\"command-kpi danger\"><span>☎</span><div><b>{{calls|length}}</b><small>Open calls</small></div></div>\n    <div class=\"command-kpi amber\"><span>⚠</span><div><b>{{calls|selectattr('status','equalto','escalated')|list|length}}</b><small>Escalated</small></div></div>\n  </div>\n\n  <div class=\"charge-board-controls\">\n    <div class=\"charge-filter-group\">\n      <button class=\"charge-filter active\" data-filter=\"all\" onclick=\"setChargeFilter('all',this)\">All rooms ({{rooms|length}})</button>\n      <button class=\"charge-filter\" data-filter=\"active\" onclick=\"setChargeFilter('active',this)\">Active calls ({{calls|length}})</button>\n    </div>\n    <div class=\"charge-view-switch\">\n      <button class=\"charge-view-btn active\" data-view=\"cards\" onclick=\"setChargeView('cards',this)\">▤ Cards</button>\n      <button class=\"charge-view-btn\" data-view=\"compact\" onclick=\"setChargeView('compact',this)\">▦ Compact</button>\n      <button class=\"charge-view-btn\" data-view=\"list\" onclick=\"setChargeView('list',this)\">☷ List</button>\n    </div>\n  </div>\n\n  <section id=\"chargeRoomGrid\" class=\"charge-room-grid view-cards\">\n  {% for room in rooms %}{% set rcalls=calls|selectattr('room_id','equalto',room.id)|list %}{% set c=rcalls[0] if rcalls else none %}\n    <article class=\"command-room-card {% if c %}has-call status-{{c.status}}{% else %}ready-card{% endif %}\"\n      data-active=\"{{1 if c else 0}}\"\n      data-priority=\"{{0 if c and c.status=='escalated' else (1 if c else 9)}}\"\n      data-elapsed=\"{{elapsed_for_call(c) if c else 0}}\">\n      <div class=\"command-room-top\">\n        <div><b>{{room.code}}</b>{% if c %}<span class=\"room-state state-{{c.status}}\">{{'Patient calling' if c.status=='new' else c.status.replace('_',' ')|title}}</span>{% else %}<span class=\"room-state state-ready\">● Ready</span>{% endif %}</div>\n        <span>{{room.zone}}</span>\n      </div>\n\n      {% if c %}\n        {% set c_elapsed=elapsed_for_call(c) %}\n        <div class=\"command-call-body\">\n          <div class=\"command-reason\">{{c.reason}}</div>\n          <div class=\"command-call-meta\">\n            <div><b class=\"timer\" data-created=\"{{utc_iso(c.created_at)}}\" data-elapsed=\"{{c_elapsed}}\" {% if c.arrived_at or c.resolved_at %}data-stop=\"{{utc_iso(c.arrived_at or c.resolved_at)}}\"{% endif %} data-call-id=\"{{c.id}}\">{{'%02d:%02d'|format(c_elapsed//60,c_elapsed%60)}}</b><small>Elapsed</small></div>\n            <div class=\"assigned-nurse\"><small>Assigned nurse</small><b>👤 {{c.assigned_nurse.name if c.assigned_nurse else 'Unassigned'}}</b></div>\n          </div>\n          {% if c.escalation_reason %}<div class=\"command-escalation\">⚠ {{c.escalation_reason}}</div>{% endif %}\n          <div class=\"command-card-actions\">\n            {% if c.status in ['new','escalated'] %}<button class=\"btn danger\" onclick=\"callAction({{c.id}},'takeover')\">Take over</button>{% endif %}\n            {% if c.status in ['taken_over','acknowledged','escalated'] %}<button class=\"btn success\" onclick=\"callAction({{c.id}},'arrive')\">Arrived</button>{% endif %}\n            {% if c.status in ['new','acknowledged','escalated','taken_over'] %}<button class=\"btn secondary\" onclick=\"openReassignModal({{c.id}},'{{room.code|e}}','{{c.assigned_nurse.name|e if c.assigned_nurse else 'Unassigned'}}')\">⇄ Reassign nurse</button>{% endif %}\n          </div>\n        </div>\n      {% else %}\n        <div class=\"ready-room-body\">\n          <p>Room ready</p>\n          <div class=\"assigned-nurse\"><small>Assigned nurse</small><b>👤 {{room.assigned_nurse.name if room.assigned_nurse else 'Unassigned'}}</b></div>\n          <button class=\"btn secondary\" onclick=\"openRoomAssignModal({{room.id}},'{{room.code|e}}')\">⇄ Reassign nurse</button>\n        </div>\n      {% endif %}\n    </article>\n  {% endfor %}\n  </section>\n</div>\n\n<div class=\"admin-modal\" id=\"reassignModal\" aria-hidden=\"true\"><div class=\"admin-modal-card command-modal\">\n  <button class=\"modal-close\" onclick=\"closeCommandModal('reassignModal')\">×</button>\n  <div class=\"eyebrow\">CALL OWNERSHIP</div><h2 id=\"reassignModalTitle\">Reassign nurse</h2><p>The original call timer continues without resetting.</p>\n  <input type=\"hidden\" id=\"reassignCallId\">\n  <label>Nurse<select id=\"modalReassignNurse\"><option value=\"\">Select active nurse…</option>{% for n in nurses %}<option value=\"{{n.id}}\">{{n.name}}</option>{% endfor %}</select></label>\n  <label>Reason<select id=\"modalReassignReason\"><option value=\"\">Select reason…</option><option>Workload balancing</option><option>Break</option><option>Shift change</option><option>No response</option><option>Clinical priority</option><option>Other</option></select></label>\n  <button class=\"btn primary wide\" onclick=\"submitReassignModal()\">Reassign nurse</button>\n</div></div>\n\n<div class=\"admin-modal\" id=\"roomAssignModal\" aria-hidden=\"true\"><div class=\"admin-modal-card command-modal\">\n  <button class=\"modal-close\" onclick=\"closeCommandModal('roomAssignModal')\">×</button>\n  <div class=\"eyebrow\">ROOM ASSIGNMENT</div><h2 id=\"roomAssignModalTitle\">Reassign room</h2>\n  <input type=\"hidden\" id=\"roomAssignId\">\n  <label>Nurse<select id=\"roomAssignNurse\"><option value=\"\">Select active nurse…</option>{% for n in nurses %}<option value=\"{{n.id}}\">{{n.name}}</option>{% endfor %}</select></label>\n  <button class=\"btn primary wide\" onclick=\"submitRoomAssign()\">Assign nurse</button>\n</div></div>\n{% endblock %}\n{% block scripts %}<script>\nvar chargeView=localStorage.getItem('charge-command-view')||'cards',chargeFilter='all';\nfunction sortChargeCards(){var g=document.getElementById('chargeRoomGrid');if(!g)return;[...g.children].sort((a,b)=>(+a.dataset.priority)-(+b.dataset.priority)||(+b.dataset.elapsed)-(+a.dataset.elapsed)).forEach(x=>g.appendChild(x))}\nfunction applyChargeView(){var g=document.getElementById('chargeRoomGrid');if(!g)return;g.className='charge-room-grid view-'+chargeView;document.querySelectorAll('.charge-view-btn').forEach(b=>b.classList.toggle('active',b.dataset.view===chargeView));applyChargeFilter();sortChargeCards()}\nfunction setChargeView(v){chargeView=v;localStorage.setItem('charge-command-view',v);applyChargeView()}\nfunction setChargeFilter(v,el){chargeFilter=v;document.querySelectorAll('.charge-filter').forEach(b=>b.classList.remove('active'));if(el)el.classList.add('active');applyChargeFilter()}\nfunction applyChargeFilter(){document.querySelectorAll('.command-room-card').forEach(c=>c.style.display=(chargeFilter==='active'&&c.dataset.active!=='1')?'none':'')}\nfunction toggleChargeFullscreen(){if(!document.fullscreenElement){document.documentElement.requestFullscreen&&document.documentElement.requestFullscreen()}else document.exitFullscreen&&document.exitFullscreen()}\nfunction openReassignModal(id,room,current){document.getElementById('reassignCallId').value=id;document.getElementById('reassignModalTitle').textContent='Reassign nurse · '+room;document.getElementById('modalReassignNurse').value='';document.getElementById('modalReassignReason').value='';document.getElementById('reassignModal').classList.add('open')}\nfunction openRoomAssignModal(id,room){document.getElementById('roomAssignId').value=id;document.getElementById('roomAssignModalTitle').textContent='Reassign room · '+room;document.getElementById('roomAssignNurse').value='';document.getElementById('roomAssignModal').classList.add('open')}\nfunction closeCommandModal(id){document.getElementById(id).classList.remove('open')}\nasync function submitReassignModal(){const id=document.getElementById('reassignCallId').value,n=document.getElementById('modalReassignNurse').value,r=document.getElementById('modalReassignReason').value;if(!n||!r)return;const res=await fetch('/api/call/'+id+'/reassign',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({nurse_id:n,reason:r})});if(res.ok){closeCommandModal('reassignModal');location.reload()}}\nasync function submitRoomAssign(){const id=document.getElementById('roomAssignId').value,n=document.getElementById('roomAssignNurse').value;if(!n)return;const res=await fetch('/api/room/'+id+'/assign',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({nurse_id:n})});if(res.ok){closeCommandModal('roomAssignModal');location.reload()}}\napplyChargeView();startSilentRefresh('#charge-live-root',3000,function(){applyChargeView()});\n</script>{% endblock %}",
    "manager.html": "{% extends 'base.html' %}{% block title %}KPI Management{% endblock %}\n{% block content %}\n<section class=\"management-console\">\n  <header class=\"management-head\">\n    <div><div class=\"eyebrow\">OPERATIONAL OVERSIGHT & ANALYTICS</div><h1>KPI Management</h1><p>Call-bell performance, SLA compliance, nurse workload and room performance.</p></div>\n    <details class=\"export-menu\"><summary>▣ Export ▾</summary><div><a href=\"/manager/export.xlsx?{{filter_qs}}\">Excel report</a><a href=\"/manager/export.pdf?{{filter_qs}}\">PDF report</a></div></details>\n  </header>\n\n  <form class=\"management-period\" method=\"get\" action=\"/manager\">\n    <div class=\"period-buttons\">\n      <button name=\"period\" value=\"today\" class=\"filter-chip {{'active' if period=='today' else ''}}\">Today</button>\n      <button name=\"period\" value=\"week\" class=\"filter-chip {{'active' if period=='week' else ''}}\">Weekly</button>\n      <button name=\"period\" value=\"month\" class=\"filter-chip {{'active' if period=='month' else ''}}\">Monthly</button>\n      <button type=\"button\" class=\"filter-chip {{'active' if period=='custom' else ''}}\" onclick=\"document.getElementById('customDates').classList.toggle('show')\">▣ Custom Date</button>\n    </div>\n    <div id=\"customDates\" class=\"custom-dates {{'show' if period=='custom' else ''}}\"><label>From <input type=\"date\" name=\"from\" value=\"{{from_date}}\"></label><label>To <input type=\"date\" name=\"to\" value=\"{{to_date}}\"></label><button class=\"btn primary\" name=\"period\" value=\"custom\">Apply</button></div>\n  </form>\n\n  <div class=\"management-primary-kpis\">\n    <div class=\"mg-kpi\"><span>☎</span><div><b>{{kpi.total}}</b><small>Total calls</small></div></div>\n    <div class=\"mg-kpi\"><span>◷</span><div><b>{{kpi.avg_response}}</b><small>Avg response</small></div></div>\n    <div class=\"mg-kpi success\"><span>✓</span><div><b>{{kpi.sla_rate}}</b><small>SLA compliance</small></div></div>\n    <div class=\"mg-kpi danger\"><span>⚠</span><div><b>{{kpi.escalations}}</b><small>Escalations</small></div></div>\n  </div>\n\n  <div class=\"management-delay-strip\"><span>🟠 <b>{{kpi.over_2m}}</b> Over 2 min</span><span>🔴 <b>{{kpi.over_5m}}</b> Over 5 min</span><span>Median response <b>{{kpi.median_response}}</b></span><span>Charge takeovers <b>{{kpi.takeovers}}</b></span></div>\n\n  <div class=\"manager-grid modern-manager-grid\">\n    <section class=\"panel performance-panel\"><div class=\"section-head\"><div><h2>Nurse Performance</h2><p class=\"muted\">Call volume, average response and escalation exposure.</p></div></div>\n      <div class=\"table-wrap\"><table><thead><tr><th>Nurse</th><th>Calls</th><th>Avg response</th><th>Over SLA</th><th>Escalated</th></tr></thead><tbody>\n      {% for n in nurse_perf %}<tr><td><b>{{n.name}}</b></td><td>{{n.calls}}</td><td><div class=\"perf-value\"><span class=\"perf-bar\"><i style=\"width:{{ [100, (n.avg_seconds or 0)/3]|min }}%\"></i></span><b>{{n.avg_response}}</b></div></td><td>{{n.over_sla}}</td><td>{{n.escalated}}</td></tr>{% endfor %}\n      {% if not nurse_perf %}<tr><td colspan=\"5\" class=\"muted\">No calls in this period.</td></tr>{% endif %}\n      </tbody></table></div>\n    </section>\n    <section class=\"panel performance-panel\"><div class=\"section-head\"><div><h2>Room Performance</h2><p class=\"muted\">Demand and response by room.</p></div></div>\n      <div class=\"table-wrap\"><table><thead><tr><th>Room</th><th>Zone</th><th>Calls</th><th>Avg response</th><th>SLA breaches</th></tr></thead><tbody>\n      {% for r in room_perf %}<tr><td><b>{{r.room}}</b></td><td>{{r.zone}}</td><td>{{r.calls}}</td><td><div class=\"perf-value\"><span class=\"perf-bar\"><i style=\"width:{{ [100, (r.avg_seconds or 0)/3]|min }}%\"></i></span><b>{{r.avg_response}}</b></div></td><td>{{r.over_sla}}</td></tr>{% endfor %}\n      {% if not room_perf %}<tr><td colspan=\"5\" class=\"muted\">No calls in this period.</td></tr>{% endif %}\n      </tbody></table></div>\n    </section>\n  </div>\n\n  <details class=\"panel lifecycle-details\"><summary><div><h2>Call Lifecycle</h2><p>{{range_label}} · Created → Acknowledged → Arrived → Resolved</p></div><span>View details ▾</span></summary>\n    <div class=\"table-wrap\"><table><thead><tr><th>Room</th><th>Created</th><th>Primary nurse</th><th>Status</th><th>Response</th><th>Total duration</th><th>Escalated</th><th>Takeover</th><th>Reason</th></tr></thead><tbody>\n    {% for c in call_rows %}<tr><td><b>{{c.room}}</b><br><small>{{c.zone}}</small></td><td>{{c.created}}</td><td>{{c.nurse}}</td><td><span class=\"badge {{c.status}}\">{{c.status_label}}</span></td><td>{{c.response}}</td><td>{{c.duration}}</td><td>{{c.escalated}}</td><td>{{c.takeover}}</td><td>{{c.reason}}</td></tr>{% endfor %}\n    {% if not call_rows %}<tr><td colspan=\"9\" class=\"muted\">No calls in this period.</td></tr>{% endif %}\n    </tbody></table></div>\n  </details>\n</section>\n{% endblock %}",
    "admin.html": "{% extends 'base.html' %}{% block title %}Admin Control Panel{% endblock %}\n{% block content %}\n<section class=\"admin-console\">\n  <aside class=\"admin-side\">\n    <div class=\"admin-side-brand\">\n      <div class=\"admin-side-icon\">⚙</div>\n      <div><b>Admin Console</b><small>Burjeel ED Call</small></div>\n    </div>\n    <nav class=\"admin-side-nav\">\n      <button class=\"admin-nav-item active\" data-admin-tab=\"overview\" onclick=\"showAdminTab('overview',this)\">▦ <span>Overview</span></button>\n      <button class=\"admin-nav-item\" data-admin-tab=\"staff\" onclick=\"showAdminTab('staff',this)\">👥 <span>Staff</span></button>\n      <button class=\"admin-nav-item\" data-admin-tab=\"permissions\" onclick=\"showAdminTab('permissions',this)\">🔐 <span>User Permissions & Work Profiles</span></button>\n      <button class=\"admin-nav-item\" data-admin-tab=\"rooms\" onclick=\"showAdminTab('rooms',this)\">🚪 <span>Rooms & QR</span></button>\n      <button class=\"admin-nav-item\" data-admin-tab=\"notifications\" onclick=\"showAdminTab('notifications',this)\">🔔 <span>Notifications</span></button>\n      <button class=\"admin-nav-item\" data-admin-tab=\"audit\" onclick=\"showAdminTab('audit',this)\">🧾 <span>Audit</span></button>\n      <a class=\"admin-nav-item\" href=\"/admin/diagnostics\">⚙ <span>Diagnostics</span>{% if open_errors %}<em>{{open_errors}}</em>{% endif %}</a>\n      <button class=\"admin-nav-item\" data-admin-tab=\"system\" onclick=\"showAdminTab('system',this)\">🖥 <span>System</span></button>\n    </nav>\n  </aside>\n\n  <div class=\"admin-main\">\n    <header class=\"admin-console-head\">\n      <div><div class=\"eyebrow\">SYSTEM ADMINISTRATION</div><h1>Control Panel</h1><p>Manage people, rooms, notifications and system health from one workspace.</p></div>\n      <div class=\"admin-health\">\n        <span class=\"health-dot ok\"></span>\n        <div><b>System operational</b><small>{{'Push configured' if vapid_ready else 'Push needs attention'}}</small></div>\n      </div>\n    </header>\n\n    <section class=\"admin-tab-pane active\" id=\"admin-tab-overview\">\n      <div class=\"admin-kpis\">\n        <button class=\"admin-kpi\" onclick=\"showAdminTab('staff')\"><span>Staff</span><b>{{users|length}}</b><small>{{active_staff}} active accounts</small></button>\n        <button class=\"admin-kpi\" onclick=\"showAdminTab('rooms')\"><span>Rooms</span><b>{{rooms|length}}</b><small>{{active_rooms}} open rooms</small></button>\n        <button class=\"admin-kpi danger\" onclick=\"location.href='/charge'\"><span>Active calls</span><b>{{active_calls}}</b><small>Live patient requests</small></button>\n        <button class=\"admin-kpi amber\" onclick=\"showAdminTab('notifications')\"><span>Push devices</span><b>{{push_count}}</b><small>{{push_healthy}} healthy devices</small></button>\n      </div>\n\n      <div class=\"admin-overview-grid\">\n        <div class=\"admin-card\">\n          <div class=\"admin-card-head\"><div><h2>Quick actions</h2><p>Common administration tasks.</p></div></div>\n          <div class=\"quick-actions\">\n            <button class=\"quick-action\" onclick=\"openAdminModal('staffModal')\"><span>＋</span><b>Add staff</b><small>Create a new staff account</small></button>\n            <button class=\"quick-action\" onclick=\"openAdminModal('roomModal')\"><span>＋</span><b>Add room</b><small>Create room and QR access</small></button>\n            <button class=\"quick-action\" onclick=\"showAdminTab('notifications')\"><span>🔔</span><b>Push health</b><small>Review registered devices</small></button>\n            <a class=\"quick-action\" href=\"/admin/diagnostics\"><span>⚙</span><b>Diagnostics</b><small>Inspect runtime errors</small></a>\n          </div>\n        </div>\n\n        <div class=\"admin-card\">\n          <div class=\"admin-card-head\"><div><h2>System health</h2><p>Current operational status.</p></div></div>\n          <div class=\"health-list\">\n            <div><span>Database</span><b class=\"health-good\">● Connected</b></div>\n            <div><span>Web Push</span><b class=\"{{'health-good' if vapid_ready else 'health-warn'}}\">● {{'Configured' if vapid_ready else 'Needs attention'}}</b></div>\n            <div><span>Diagnostics</span><b class=\"{{'health-good' if open_errors==0 else 'health-bad'}}\">● {{'Clear' if open_errors==0 else open_errors|string + ' open'}}</b></div>\n            <div><span>SLA threshold</span><b>{{sla_seconds}} sec</b></div>\n          </div>\n        </div>\n\n        <div class=\"admin-card admin-span-2\">\n          <div class=\"admin-card-head\"><div><h2>Recent activity</h2><p>Latest workflow and administration events.</p></div><button class=\"btn secondary\" onclick=\"showAdminTab('audit')\">View audit</button></div>\n          <div class=\"activity-list\">\n            {% for a in audit_logs[:8] %}\n            <div class=\"activity-row\"><div class=\"activity-dot\"></div><div><b>{{a.action.replace('_',' ')|title}}</b><small>{{a.detail or '-'}}</small></div><time>{{a.created_at.strftime('%H:%M')}}</time></div>\n            {% else %}<div class=\"empty-state\">No recent activity.</div>{% endfor %}\n          </div>\n        </div>\n      </div>\n    </section>\n\n    <section class=\"admin-tab-pane\" id=\"admin-tab-staff\">\n      <div class=\"admin-pane-head\"><div><h2>Staff & Roles</h2><p>Create accounts, change access and reset passwords.</p></div><button class=\"btn primary\" onclick=\"openAdminModal('staffModal')\">＋ Add staff</button></div>\n      <div class=\"admin-toolbar\"><input id=\"staffSearch\" oninput=\"filterAdminRows('staffTable',this.value)\" placeholder=\"Search staff by name, email or role\"></div>\n      <div class=\"table-wrap compact-table\"><table id=\"staffTable\"><thead><tr><th>Name</th><th>Email</th><th>Role</th><th>Status</th><th></th></tr></thead><tbody>\n      {% for u in users %}<tr data-search=\"{{u.name}} {{u.email}} {{u.role}}\">\n        <td><div class=\"person-cell\"><span>{{u.name[:1]|upper}}</span><div><b>{{u.name}}</b><small>Staff ID {{u.id}}</small></div></div></td>\n        <td>{{u.email}}</td><td>{{role_labels.get(u.role,u.role.replace('_',' ')|title)}}</td>\n        <td><span class=\"badge {{'resolved' if u.active else 'new'}}\">{{'Active' if u.active else 'Disabled'}}</span></td>\n        <td class=\"admin-actions-cell\"><details class=\"action-menu\"><summary>⋮</summary><div>\n          <form method=\"post\" action=\"/admin/users/{{u.id}}/toggle\"><button {% if u.id==current_user.id %}disabled{% endif %}>{{'Disable account' if u.active else 'Enable account'}}</button></form>\n          <button type=\"button\" onclick=\"openPasswordReset({{u.id}},'{{u.name|e}}')\">Reset password</button>\n        </div></details></td>\n      </tr>{% endfor %}\n      </tbody></table></div>\n    </section>\n\n    <section class=\"admin-tab-pane\" id=\"admin-tab-permissions\">\n      <div class=\"admin-pane-head\"><div><h2>User Permissions & Work Profiles</h2><p>Assign a standard work profile or customize system access for each staff member.</p></div><span class=\"permission-note\">Role = job title · Work Profile = system access</span></div>\n      <div class=\"profile-legend\">\n        {% for p in profile_options %}<span><b>{{p.label}}</b></span>{% endfor %}\n      </div>\n      <div class=\"permission-user-list\">\n        {% for row in permission_rows %}\n        <form class=\"permission-user-card\" method=\"post\" action=\"/admin/users/{{row.user.id}}/permissions\">\n          <div class=\"permission-user-head\">\n            <div class=\"person-cell\"><span>{{row.user.name[:1]|upper}}</span><div><b>{{row.user.name}}</b><small>{{row.role_label}} · {{row.user.email}}</small></div></div>\n            <label>Work Profile\n              <select name=\"profile_name\" id=\"profile-{{row.user.id}}\" onchange=\"applyPermissionPreset({{row.user.id}},this.value)\" {% if row.user.id==current_user.id and row.user.role=='admin' %}disabled{% endif %}>\n                {% for p in profile_options %}<option value=\"{{p.key}}\" {% if row.profile_name==p.key %}selected{% endif %}>{{p.label}}</option>{% endfor %}\n              </select>\n            </label>\n          </div>\n          <div class=\"permission-grid\" id=\"permission-grid-{{row.user.id}}\">\n            {% for p in permission_defs %}\n            <label class=\"permission-item\">\n              <input type=\"checkbox\" name=\"permissions\" value=\"{{p.key}}\" {% if p.key in row.permissions %}checked{% endif %} {% if row.user.id==current_user.id and row.user.role=='admin' %}disabled{% endif %}>\n              <span><b>{{p.label}}</b><small>{{p.description}}</small></span>\n            </label>\n            {% endfor %}\n          </div>\n          <div class=\"permission-actions\">\n            {% if not (row.user.id==current_user.id and row.user.role=='admin') %}<button class=\"btn primary\">Save permissions</button>{% endif %}\n            <button type=\"submit\" class=\"btn secondary\" formaction=\"/admin/users/{{row.user.id}}/permissions/reset\" formmethod=\"post\" {% if row.user.id==current_user.id and row.user.role=='admin' %}disabled{% endif %}>Restore role defaults</button>\n          </div>\n        </form>\n        {% endfor %}\n      </div>\n    </section>\n\n    <section class=\"admin-tab-pane\" id=\"admin-tab-rooms\">\n      <div class=\"admin-pane-head\"><div><h2>Rooms & QR</h2><p>Manage room identity, assignment, status and patient QR codes.</p></div><button class=\"btn primary\" onclick=\"openAdminModal('roomModal')\">＋ Add room</button></div>\n      {% if room_msg %}<div class=\"admin-room-msg {{'error' if room_msg_type=='error' else 'success'}}\">{{room_msg}}</div>{% endif %}\n      <div class=\"admin-toolbar\"><input id=\"roomSearch\" oninput=\"filterAdminRows('roomTable',this.value)\" placeholder=\"Search room, zone or assigned nurse\"></div>\n      <div class=\"table-wrap compact-table\"><table id=\"roomTable\"><thead><tr><th>Room</th><th>Assigned nurse</th><th>Status</th><th>QR</th><th></th></tr></thead><tbody>\n      {% for r in rooms %}<tr data-search=\"{{r.code}} {{r.zone}} {{r.assigned_nurse.name if r.assigned_nurse else 'Unassigned'}}\">\n        <td><div class=\"room-name-cell\"><b>{{r.code}}</b><small>{{r.zone}}</small></div></td>\n        <td><form class=\"inline-form compact\" method=\"post\" action=\"/admin/rooms/{{r.id}}/assign\"><select name=\"nurse_id\"><option value=\"\">Unassigned</option>{% for n in nurses %}<option value=\"{{n.id}}\" {% if r.assigned_nurse_id==n.id %}selected{% endif %}>{{n.name}}</option>{% endfor %}</select><button class=\"btn secondary\">Save</button></form></td>\n        <td><span class=\"badge {{'resolved' if r.occupied else 'new'}}\">{{'Open' if r.occupied else 'Closed'}}</span></td>\n        <td><a class=\"btn secondary qr-small\" href=\"/admin/rooms/{{r.id}}/qr\" target=\"_blank\">View QR</a></td>\n        <td class=\"admin-actions-cell\"><details class=\"action-menu\"><summary>⋮</summary><div>\n          <button type=\"button\" onclick=\"openRoomEdit({{r.id}},'{{r.code|e}}','{{r.zone|e}}')\">Edit room</button>\n          <a href=\"/admin/rooms/{{r.id}}/qr?print=1\" target=\"_blank\">Print QR</a>\n          <form method=\"post\" action=\"/admin/rooms/{{r.id}}/toggle\"><button>{{'Close room' if r.occupied else 'Open room'}}</button></form>\n          <form method=\"post\" action=\"/admin/rooms/{{r.id}}/token\"><button>Regenerate QR</button></form>\n          <form method=\"post\" action=\"/admin/rooms/{{r.id}}/delete\" onsubmit=\"return confirm('Delete room {{r.code}}? This is only allowed when the room has no call or handover history.');\"><button class=\"danger-link\">Delete room</button></form>\n        </div></details></td>\n      </tr>{% endfor %}\n      </tbody></table></div>\n    </section>\n\n    <section class=\"admin-tab-pane\" id=\"admin-tab-notifications\">\n      <div class=\"admin-pane-head\"><div><h2>Notifications</h2><p>Browser and installed PWA subscriptions registered for staff devices.</p></div></div>\n      <div class=\"admin-kpis notification-kpis\">\n        <div class=\"admin-kpi\"><span>Subscriptions</span><b>{{push_count}}</b><small>Web Push endpoints</small></div>\n        <div class=\"admin-kpi\"><span>Tracked devices</span><b>{{push_devices|length}}</b><small>Browser / PWA contexts</small></div>\n        <div class=\"admin-kpi\"><span>Healthy</span><b>{{push_healthy}}</b><small>Active device registrations</small></div>\n        <div class=\"admin-kpi danger\"><span>Needs attention</span><b>{{push_attention}}</b><small>Expired / unhealthy</small></div>\n      </div>\n      <div class=\"table-wrap compact-table\"><table><thead><tr><th>Mode</th><th>Platform</th><th>Health</th><th>Verified</th><th>Last seen</th></tr></thead><tbody>\n        {% for d in push_devices %}<tr><td><span class=\"device-mode\">{{d.display_mode|upper}}</span></td><td>{{d.platform or '-'}}</td><td><span class=\"badge {{'resolved' if d.health_status in ['active','registered'] else 'new'}}\">{{d.health_status.replace('_',' ')|title}}</span></td><td>{{'Yes' if d.last_verified else 'No'}}</td><td>{{d.last_seen.strftime('%Y-%m-%d %H:%M') if d.last_seen else '-'}}</td></tr>\n        {% else %}<tr><td colspan=\"5\" class=\"empty-state\">No push devices registered yet.</td></tr>{% endfor %}\n      </tbody></table></div>\n    </section>\n\n    <section class=\"admin-tab-pane\" id=\"admin-tab-audit\">\n      <div class=\"admin-pane-head\"><div><h2>Audit Log</h2><p>Latest 100 workflow and administration actions.</p></div></div>\n      <div class=\"admin-toolbar\"><input oninput=\"filterAdminRows('auditTable',this.value)\" placeholder=\"Search action, user, room or detail\"></div>\n      <div class=\"table-wrap compact-table\"><table id=\"auditTable\"><thead><tr><th>Time</th><th>Action</th><th>User</th><th>Room</th><th>Detail</th></tr></thead><tbody>\n      {% for a in audit_logs %}<tr data-search=\"{{a.action}} {{a.user.name if a.user else ''}} {{a.room_id or ''}} {{a.detail or ''}}\"><td>{{a.created_at.strftime('%Y-%m-%d %H:%M:%S')}}</td><td><b>{{a.action.replace('_',' ')|title}}</b></td><td>{{a.user.name if a.user else '-'}}</td><td>{{a.room_id or '-'}}</td><td>{{a.detail or '-'}}</td></tr>{% endfor %}\n      </tbody></table></div>\n    </section>\n\n    <section class=\"admin-tab-pane\" id=\"admin-tab-system\">\n      <div class=\"admin-pane-head\"><div><h2>System</h2><p>Application configuration and runtime readiness.</p></div></div>\n      <div class=\"system-grid modern\">\n        <div><span>Database</span><b class=\"health-good\">Connected</b><small>Primary application store</small></div>\n        <div><span>SLA</span><b>{{sla_seconds}} sec</b><small>Escalation threshold</small></div>\n        <div><span>Web Push</span><b class=\"{{'health-good' if vapid_ready else 'health-warn'}}\">{{'Configured' if vapid_ready else 'Not configured'}}</b><small>Background notifications</small></div>\n        <div><span>Runtime</span><b>FastAPI</b><small>Single-file app.py deployment</small></div>\n        <div><span>Diagnostics</span><b class=\"{{'health-good' if open_errors==0 else 'health-bad'}}\">{{open_errors}} open</b><small><a href=\"/admin/diagnostics\">Open monitor</a></small></div>\n        <div><span>Push devices</span><b>{{push_count}}</b><small>{{push_healthy}} healthy</small></div>\n        <div class=\"live-screen-sound-card\"><span>Live Screen Sound</span><b>Repeat reminder</b><form method=\"post\" action=\"/admin/live-screen-sound\" class=\"live-screen-sound-form\"><small>Default is 02:00. New patient calls sound immediately; reminders continue until Arrived / Closed / Resolved.</small><div class=\"live-sound-inputs\"><label>Minutes<input name=\"sound_repeat_minutes\" type=\"number\" min=\"0\" max=\"60\" value=\"{{escalation_settings.sound_repeat // 60}}\" required></label><b>:</b><label>Seconds<input name=\"sound_repeat_seconds\" type=\"number\" min=\"0\" max=\"59\" value=\"{{escalation_settings.sound_repeat % 60}}\" required></label></div><button class=\"btn primary\" type=\"submit\">Save sound interval</button></form></div>\n      </div>\n    </section>\n  </div>\n</section>\n\n<div class=\"admin-modal\" id=\"staffModal\" aria-hidden=\"true\"><div class=\"admin-modal-card\"><button class=\"modal-close\" onclick=\"closeAdminModal('staffModal')\">×</button><div class=\"eyebrow\">STAFF MANAGEMENT</div><h2>Add staff member</h2><p>Create a secure staff account and assign a role.</p><form class=\"modal-form\" method=\"post\" action=\"/admin/users\"><label>Full name<input name=\"name\" required></label><label>Email<input name=\"email\" type=\"email\" required></label><label>Role<select name=\"role\" required><option value=\"nurse\">Nurse</option><option value=\"charge\">Nurse In Charge</option><option value=\"nurse_supervisor\">Nurse Supervisor</option><option value=\"manager\">Nurse Manager</option><option value=\"ed_manager\">ED Manager</option><option value=\"hod\">HOD</option><option value=\"admin\">System Admin</option></select></label><label>Temporary password<input name=\"password\" type=\"password\" minlength=\"8\" required></label><button class=\"btn primary wide\">Add staff</button></form></div></div>\n\n<div class=\"admin-modal\" id=\"roomModal\" aria-hidden=\"true\"><div class=\"admin-modal-card\"><button class=\"modal-close\" onclick=\"closeAdminModal('roomModal')\">×</button><div class=\"eyebrow\">ROOM MANAGEMENT</div><h2>Add room</h2><p>Create a room, zone and initial nurse assignment.</p><form class=\"modal-form\" method=\"post\" action=\"/admin/rooms\"><label>Room code<input name=\"code\" placeholder=\"ED-09\" required></label><label>Zone<input name=\"zone\" placeholder=\"ED Main\" required></label><label>Assigned nurse<select name=\"nurse_id\"><option value=\"\">Unassigned</option>{% for n in nurses %}<option value=\"{{n.id}}\">{{n.name}}</option>{% endfor %}</select></label><button class=\"btn primary wide\">Add room</button></form></div></div>\n\n<div class=\"admin-modal\" id=\"passwordModal\" aria-hidden=\"true\"><div class=\"admin-modal-card\"><button class=\"modal-close\" onclick=\"closeAdminModal('passwordModal')\">×</button><div class=\"eyebrow\">SECURITY</div><h2 id=\"passwordModalTitle\">Reset password</h2><form id=\"passwordResetForm\" class=\"modal-form\" method=\"post\"><label>New password<input name=\"password\" type=\"password\" minlength=\"8\" required></label><button class=\"btn primary wide\">Reset password</button></form></div></div>\n\n<div class=\"admin-modal\" id=\"roomEditModal\" aria-hidden=\"true\"><div class=\"admin-modal-card\"><button class=\"modal-close\" onclick=\"closeAdminModal('roomEditModal')\">×</button><div class=\"eyebrow\">ROOM MANAGEMENT</div><h2>Edit room</h2><form id=\"roomEditForm\" class=\"modal-form\" method=\"post\"><label>Room code<input id=\"editRoomCode\" name=\"code\" maxlength=\"40\" required></label><label>Zone<input id=\"editRoomZone\" name=\"zone\" maxlength=\"80\" required></label><button class=\"btn primary wide\">Save changes</button></form></div></div>\n{% endblock %}\n{% block scripts %}<script>\n(function(){\n  const initial=(location.hash||'').replace('#','');\n  const allowed=['overview','staff','permissions','rooms','notifications','audit','system'];\n  if(allowed.includes(initial))showAdminTab(initial);\n})();\nfunction showAdminTab(name,btn){\n  document.querySelectorAll('.admin-tab-pane').forEach(x=>x.classList.remove('active'));\n  document.querySelectorAll('.admin-nav-item[data-admin-tab]').forEach(x=>x.classList.remove('active'));\n  const pane=document.getElementById('admin-tab-'+name);if(pane)pane.classList.add('active');\n  const nav=btn||document.querySelector('.admin-nav-item[data-admin-tab=\"'+name+'\"]');if(nav)nav.classList.add('active');\n  history.replaceState(null,'','#'+name);\n  if(window.innerWidth<800)window.scrollTo({top:0,behavior:'smooth'});\n}\nfunction openAdminModal(id){const m=document.getElementById(id);if(m){m.classList.add('open');m.setAttribute('aria-hidden','false')}}\nfunction closeAdminModal(id){const m=document.getElementById(id);if(m){m.classList.remove('open');m.setAttribute('aria-hidden','true')}}\nfunction openPasswordReset(id,name){document.getElementById('passwordResetForm').action='/admin/users/'+id+'/reset';document.getElementById('passwordModalTitle').textContent='Reset password · '+name;openAdminModal('passwordModal')}\nfunction openRoomEdit(id,code,zone){document.getElementById('roomEditForm').action='/admin/rooms/'+id+'/edit';document.getElementById('editRoomCode').value=code;document.getElementById('editRoomZone').value=zone;openAdminModal('roomEditModal')}\nconst ADMIN_PERMISSION_PRESETS={{permission_presets_json|safe}};\nfunction applyPermissionPreset(uid,profile){\n  const grid=document.getElementById('permission-grid-'+uid),allowed=new Set(ADMIN_PERMISSION_PRESETS[profile]||[]);\n  if(!grid)return;\n  grid.querySelectorAll('input[type=\"checkbox\"][name=\"permissions\"]').forEach(cb=>{cb.checked=allowed.has(cb.value)});\n}\nfunction filterAdminRows(tableId,value){const q=(value||'').trim().toLowerCase();document.querySelectorAll('#'+tableId+' tbody tr').forEach(r=>{const hay=(r.dataset.search||r.textContent||'').toLowerCase();r.style.display=!q||hay.includes(q)?'':'none'})}\ndocument.addEventListener('click',e=>{if(e.target.classList.contains('admin-modal'))closeAdminModal(e.target.id)});\n</script>{% endblock %}",
    "diagnostics.html": "{% extends 'base.html' %}{% block title %}System Error Monitor{% endblock %}\n{% block content %}\n<section class=\"diag-shell\">\n  <div class=\"diag-head\">\n    <div><div class=\"eyebrow\">ADMIN DIAGNOSTICS</div><h1>System Error Monitor</h1><p>Unhandled runtime errors and automatic route-recovery events. Patient identifiers/signature payloads are redacted.</p></div>\n    <div class=\"diag-head-actions\"><a class=\"btn secondary\" href=\"/admin/diagnostics/application-log\">Application Log</a><a class=\"btn secondary\" href=\"/admin/diagnostics/runtime-report\">Runtime Report</a><a class=\"btn primary\" href=\"/admin/diagnostics\">Refresh</a></div>\n  </div>\n\n  <div class=\"diag-stats\">\n    <div class=\"diag-stat\"><span>Open Errors</span><b>{{open_errors}}</b><small class=\"diag-pill ok\">{{'Clear' if open_errors==0 else 'Needs review'}}</small></div>\n    <div class=\"diag-stat\"><span>Recovered Automatically</span><b>{{recovered}}</b><small class=\"diag-pill ok\">Self-healed</small></div>\n    <div class=\"diag-stat\"><span>Captured Events</span><b>{{captured}}</b><small class=\"diag-pill blue\">Latest {{rows_limit}}</small></div>\n    <div class=\"diag-stat\"><span>Schema Gate</span><b class=\"{{'diag-ready' if schema_ready else 'diag-bad'}}\">{{'READY' if schema_ready else 'DEGRADED'}}</b><small>{{schema_time}}</small></div>\n  </div>\n\n  <section class=\"diag-panel\">\n    <form class=\"diag-filter\" method=\"get\" action=\"/admin/diagnostics\">\n      <label>Search<input name=\"q\" value=\"{{q}}\" placeholder=\"Error ID / path / message / exception\"></label>\n      <label>Status<select name=\"status\"><option value=\"all\" {% if status=='all' %}selected{% endif %}>All</option><option value=\"open\" {% if status=='open' %}selected{% endif %}>Open</option><option value=\"recovered\" {% if status=='recovered' %}selected{% endif %}>Recovered</option></select></label>\n      <label>Rows<select name=\"rows\">{% for n in [50,100,250] %}<option value=\"{{n}}\" {% if rows_limit==n %}selected{% endif %}>{{n}}</option>{% endfor %}</select></label>\n      <button class=\"btn primary\">Apply</button>\n      <a class=\"btn secondary\" href=\"/admin/diagnostics/download?q={{q|urlencode}}&status={{status}}&rows={{rows_limit}}\">Download JSONL</a>\n    </form>\n    <form method=\"post\" action=\"/admin/diagnostics/clear\" onsubmit=\"return confirm('Clear all captured diagnostic events?');\"><button class=\"btn danger diag-clear\">Clear Monitor Events</button></form>\n\n    <div class=\"table-wrap diag-table\"><table><thead><tr><th>Time</th><th>Status</th><th>Count</th><th>Error ID</th><th>Source</th><th>Request</th><th>User</th><th>Exception</th><th>Message / Trace</th></tr></thead><tbody>\n    {% for e in events %}<tr>\n      <td>{{e.last_seen.strftime('%Y-%m-%d %H:%M:%S') if e.last_seen else '-'}}</td>\n      <td><span class=\"diag-status {{e.status}}\">{{e.status|title}}</span></td>\n      <td>{{e.count}}</td><td><code>{{e.error_id}}</code></td><td>{{e.source}}</td><td><code>{{e.request_path}}</code></td><td>{{e.user_ref}}</td><td>{{e.exception_type}}</td>\n      <td class=\"diag-message\"><b>{{e.message or '-'}}</b>{% if e.trace %}<details><summary>Trace</summary><pre>{{e.trace}}</pre></details>{% endif %}</td>\n    </tr>{% else %}<tr><td colspan=\"9\" class=\"diag-empty\">No captured runtime errors.</td></tr>{% endfor %}\n    </tbody></table></div>\n  </section>\n</section>\n{% endblock %}",
    "alerts.html": "{% extends 'base.html' %}{% block title %}Alerts - Burjeel ED Call{% endblock %}\n{% block content %}\n<section class=\"alerts-page\">\n  <div class=\"alerts-page-head\"><div><div class=\"eyebrow\">NOTIFICATION CENTER</div><h1>Alerts</h1><p>Patient calls, re-calls, assignments, handovers and management escalation notifications.</p></div><button class=\"btn secondary\" onclick=\"markAllAlertsRead()\">Mark all as read</button></div>\n  <div class=\"alerts-page-summary\"><b>{{unread}}</b><span>Unread alerts</span></div>\n  <div class=\"alerts-history\">\n    {% for a in alerts %}\n    <a class=\"alert-history-row {% if not a.read_at %}unread{% endif %}\" href=\"{{a.url or '/'}}\" onclick=\"markAlertRead({{a.id}})\">\n      <div class=\"alert-kind-icon\">🔔</div>\n      <div><div class=\"alert-history-title\"><b>{{a.title}}</b>{% if not a.read_at %}<span>NEW</span>{% endif %}</div><p>{{a.body}}</p><small>{{a.kind.replace('-',' ')|title}} · {{a.created_at.strftime('%Y-%m-%d %H:%M:%S')}}</small></div>\n    </a>\n    {% else %}<div class=\"alerts-empty-page\">No alerts yet.</div>{% endfor %}\n  </div>\n</section>\n{% endblock %}",
    "qr.html": "{% extends 'base.html' %}{% block title %}{{room.code}} QR Code{% endblock %}{% block content %}\n<section class=\"qr-shell\"><div class=\"qr-poster\"><img class=\"qr-logo\" src=\"/static/icons/icon-512.png\"><div class=\"eyebrow\">{{room.zone}}</div><h1>{{room.code}}</h1><h2>Scan to call your nurse</h2><p class=\"qr-ar\" dir=\"rtl\">امسح الكود لاستدعاء الممرضة</p><img class=\"qr-image\" src=\"data:image/svg+xml;base64,{{qr_data}}\" alt=\"QR code for {{room.code}}\"><p class=\"muted\">Point your phone camera at the QR code.</p><p class=\"qr-ar muted\" dir=\"rtl\">وجّه كاميرا الهاتف إلى رمز QR.</p><div class=\"qr-print-actions\"><button class=\"btn primary\" onclick=\"window.print()\">Print QR</button><a class=\"btn secondary\" href=\"/admin#rooms\">Back</a></div></div></section>\n{% if auto_print %}<script>window.addEventListener('load',()=>setTimeout(()=>window.print(),300));</script>{% endif %}\n{% endblock %}",
    "patient.html": "{% extends 'base.html' %}{% block title %}{{ room.code }} - {{ t.page_title }}{% endblock %}\n{% block content %}\n<section id=\"patient-live-root\" class=\"patient-shell patient-lang\" data-refresh-key=\"{{refresh_key}}\" dir=\"{{ 'rtl' if lang=='ar' else 'ltr' }}\"><div class=\"patient-card\">\n<div class=\"patient-lang-switch\" dir=\"ltr\"><a class=\"lang-chip {{'active' if lang=='en' else ''}}\" href=\"?lang=en\">EN</a><span>•</span><a class=\"lang-chip {{'active' if lang=='ar' else ''}}\" href=\"?lang=ar\">ع</a></div>\n<div class=\"patient-brand-block\">\n  <img class=\"hero-logo patient-main-logo\" src=\"/static/icons/icon-512.png\">\n  <div class=\"eyebrow patient-zone\">{{ room.zone }}</div>\n  <h1 class=\"patient-room-code\">{{ room.code }}</h1>\n</div>\n<h2 class=\"patient-question\">{{ t.help_title }}</h2><p class=\"muted patient-note\">{{ t.help_note }}</p>\n<div id=\"patient-state\">{% if active_call %}\n<div class=\"call-active\"><div class=\"pulse\"></div><h2>{{ t.call_active }}</h2><div class=\"timer\" data-created=\"{{active_created_iso}}\" data-elapsed=\"{{active_elapsed}}\" {% if active_stop_iso %}data-stop=\"{{active_stop_iso}}\"{% endif %} data-call-id=\"{{ active_call.id }}\">{{'%02d:%02d'|format(active_elapsed//60,active_elapsed%60)}}</div><p id=\"patientStatus\">{{ status_label }}</p><p class=\"muted patient-call-note\">{{ t.arrived_note if active_call.arrived_at or active_call.resolved_at else (t.ack_note if active_call.acknowledged_at else t.wait_note) }}</p>\n{% if not active_call.arrived_at and not active_call.resolved_at %}\n<div id=\"recallPanel\" class=\"recall-panel\" data-recall-at=\"{{recall.recall_available_at}}\" data-recall-count=\"{{recall.recall_count}}\" data-recall-limit=\"{{recall.recall_limit}}\" {% if not recall.can_recall %}hidden{% endif %}>\n  <p class=\"recall-warning\">⚠️ {{t.recall_message}}</p>\n  <form id=\"recallForm\" method=\"post\" action=\"/room/{{room.qr_token}}/recall/{{active_call.id}}?lang={{lang}}\" class=\"recall-form\">\n    <button id=\"recallBtn\" type=\"submit\" class=\"btn recall-btn\">{{t.recall_button}}</button>\n  </form>\n  <small id=\"recallStatus\">{{recall.recall_count}} / {{recall.recall_limit}}</small>\n</div>\n{% endif %}</div>\n{% else %}\n<form id=\"patientCallForm\" method=\"post\" action=\"/room/{{room.qr_token}}/call\">\n  <input type=\"hidden\" name=\"lang\" value=\"{{lang}}\">\n  <div class=\"reason-grid\">\n    <input class=\"reason-radio\" type=\"radio\" name=\"reason\" value=\"General assistance\" id=\"reason-general\" checked>\n    <label class=\"reason\" for=\"reason-general\"><span class=\"reason-icon\">🤝</span><span>{{t.general}}</span></label>\n    <input class=\"reason-radio\" type=\"radio\" name=\"reason\" value=\"Pain\" id=\"reason-pain\">\n    <label class=\"reason\" for=\"reason-pain\"><span class=\"reason-icon\">❤️‍🩹</span><span>{{t.pain}}</span></label>\n    <input class=\"reason-radio\" type=\"radio\" name=\"reason\" value=\"Toilet assistance\" id=\"reason-toilet\">\n    <label class=\"reason\" for=\"reason-toilet\"><span class=\"reason-icon\">🚻</span><span>{{t.toilet}}</span></label>\n    <input class=\"reason-radio\" type=\"radio\" name=\"reason\" value=\"IV / Medication\" id=\"reason-medication\">\n    <label class=\"reason\" for=\"reason-medication\"><span class=\"reason-icon\">💧</span><span>{{t.medication}}</span></label>\n  </div>\n  <div class=\"patient-call-action\">\n    <button id=\"callBtn\" type=\"submit\" class=\"call-btn call-btn-pro\" aria-label=\"{{t.call_nurse}}\">\n      <span class=\"call-btn-icon\">🔔</span>\n      <span class=\"call-btn-label\">{{ t.call_nurse }}</span>\n      <span class=\"call-btn-sub\">{{ t.call_hint if t.call_hint is defined else '' }}</span>\n    </button>\n    <p id=\"callMessage\" class=\"muted call-message\"></p>\n  </div>\n</form>\n{% endif %}</div>\n<div class=\"physical-note\">⚠️ {{ t.physical_note }}</div></div></section>\n{% endblock %}\n{% block scripts %}<script>\n(function(){\n  const patientLang='{{lang}}', qs=new URLSearchParams(location.search), stored=localStorage.getItem('edcall-lang');\n  if(!qs.has('lang')&&stored&&stored!==patientLang){const u=new URL(location.href);u.searchParams.set('lang',stored);location.replace(u.toString());return;}\n  localStorage.setItem('edcall-lang',patientLang);\n\n  window.rebindPatientControls=function(){\n    if(window.bindRecallForm)window.bindRecallForm();\n    const form=document.getElementById('patientCallForm'), btn=document.getElementById('callBtn');\n    if(!form||!btn)return;\n    form.onsubmit=async function(e){\n      if(!window.fetch)return;\n      e.preventDefault();\n      const fd=new FormData(form), reason=fd.get('reason')||'General assistance';\n      btn.disabled=true;btn.classList.add('sending');\n      const label=btn.querySelector('.call-btn-label'),sub=btn.querySelector('.call-btn-sub');\n      if(label)label.textContent={{ \"'جاري إرسال النداء...'\" if lang==\"ar\" else \"'Sending call...'\" }};\n      if(sub)sub.textContent={{ \"'يرجى الانتظار'\" if lang==\"ar\" else \"'Please wait'\" }};\n      try{\n        const r=await fetch('/api/room/{{room.qr_token}}/call',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({reason})});\n        const d=await r.json();\n        if(!r.ok||!d.ok)throw new Error('call_failed');\n        btn.classList.remove('sending');btn.classList.add('sent');\n        if(label)label.textContent={{ \"'تم إرسال النداء ✓'\" if lang==\"ar\" else \"'Call sent ✓'\" }};\n        if(sub)sub.textContent={{ \"'تم إشعار الممرضة'\" if lang==\"ar\" else \"'Nurse notified'\" }};\n        if(window.silentRefreshNow)await silentRefreshNow('#patient-live-root',window.rebindPatientControls);else location.reload();\n      }catch(err){\n        const m=document.getElementById('callMessage');if(m)m.textContent='{{t.error}}';\n        btn.disabled=false;btn.classList.remove('sending');\n        if(label)label.textContent='{{t.call_nurse}}';\n        if(sub)sub.textContent='{{ t.call_hint if t.call_hint is defined else \"\" }}';\n      }\n    };\n  };\n  window.bindRecallForm=function(){\n    const form=document.getElementById('recallForm'),btn=document.getElementById('recallBtn'),panel=document.getElementById('recallPanel'),st=document.getElementById('recallStatus');\n    if(!form||!btn)return;\n    form.onsubmit=async function(e){\n      if(!window.fetch)return;\n      e.preventDefault();\n      btn.disabled=true;btn.textContent={{ \"'جاري إعادة النداء...'\" if lang==\"ar\" else \"'Sending re-call...'\" }};\n      try{\n        const timer=document.querySelector('#patient-live-root .timer[data-call-id]');\n        const id=timer&&timer.dataset.callId;\n        if(!id)throw new Error('call_id_missing');\n        const r=await fetch('/api/call/'+id+'/recall',{method:'POST',cache:'no-store',headers:{'Accept':'application/json'}});\n        const d=await r.json();\n        if(!r.ok||!d.ok)throw new Error(d.error||'recall_failed');\n        if(st)st.textContent={{ \"'تمت إعادة النداء ✓'\" if lang==\"ar\" else \"'Re-call sent ✓'\" }};\n        if(panel){panel.dataset.recallAt=d.recall_available_at||'';panel.dataset.recallCount=d.recall_count||0;panel.hidden=true}\n        setTimeout(()=>{if(window.silentRefreshNow)silentRefreshNow('#patient-live-root',window.rebindPatientControls)},400);\n      }catch(err){\n        if(st)st.textContent=err.message||'{{t.error}}';\n        btn.disabled=false;btn.textContent='{{t.recall_button}}';\n      }\n    };\n  };\n  window.updateRecallAvailability=function(){\n    const p=document.getElementById('recallPanel');if(!p)return;\n    const limit=Number(p.dataset.recallLimit||3),count=Number(p.dataset.recallCount||0),at=Date.parse(p.dataset.recallAt||'');\n    if(count>=limit){p.hidden=true;return}\n    if(at&&Date.now()>=at){p.hidden=false;const b=document.getElementById('recallBtn');if(b)b.disabled=false}\n  };\n  window.bindRecallForm();\n  window.rebindPatientControls();\n  window.updateRecallAvailability();\n  if(!window.__recallUiTimer)window.__recallUiTimer=setInterval(window.updateRecallAvailability,1000);\n  if(window.initPatientLiveSync)window.initPatientLiveSync();\n})();\n</script>{% endblock %}",
"wallboard.html": "{% extends 'base.html' %}{% block title %}Call Bell Screen{% endblock %}{% block body_class %}wallboard-page{% endblock %}\n{% block content %}\n<section class=\"wallboard-shell\" data-lang=\"{{lang}}\">\n  <div class=\"wallboard-head\">\n    <div><div class=\"eyebrow\">{{ t.kicker }}</div><div class=\"command-title-row\"><h1>Live Screen</h1><span class=\"live-pill\">● {{t.connected}}</span></div><p>{{ t.subtitle }}</p></div>\n    <div class=\"wallboard-tools\">\n      <div class=\"wall-clock\" id=\"wallClock\">--:--:--</div>\n      <div class=\"view-switch\" title=\"Display mode\">\n        <button class=\"view-btn\" data-view=\"compact\" onclick=\"setWallView('compact',this)\">▦ {{t.compact}}</button>\n        <button class=\"view-btn\" data-view=\"cards\" onclick=\"setWallView('cards',this)\">▤ {{t.cards}}</button>\n        <button class=\"view-btn\" data-view=\"list\" onclick=\"setWallView('list',this)\">☷ {{t.list}}</button>\n        <button class=\"view-btn\" data-view=\"zones\" onclick=\"setWallView('zones',this)\">▥ {{t.zones}}</button>\n      </div>\n      <div class=\"lang-mini\" dir=\"ltr\"><a class=\"{{'active' if lang=='en' else ''}}\" href=\"/wallboard?lang=en\">EN</a><span>•</span><a class=\"{{'active' if lang=='ar' else ''}}\" href=\"/wallboard?lang=ar\">ع</a></div>\n      <button id=\"soundBtn\" class=\"btn primary\" onclick=\"enableWallboardSound()\">🔊 {{t.enable_sound}}</button>\n      <button class=\"btn secondary\" onclick=\"toggleWallFullscreen()\">⛶ {{t.fullscreen}}</button>\n    </div>\n  </div>\n\n  <div class=\"wall-stats\">\n    <div class=\"wall-stat danger\"><span>☎ {{t.active_calls}}</span><b id=\"statActive\">{{stats.active}}</b></div>\n    <div class=\"wall-stat overdue\"><span>◷ {{t.over_sla}}</span><b id=\"statOverSla\">{{stats.over_sla}}</b></div>\n    <div class=\"wall-stat\"><span>🛏 {{t.total_rooms}}</span><b id=\"statRooms\">{{stats.total_rooms}}</b><small><b id=\"statEscalated\">{{stats.escalated}}</b> {{t.escalated}}</small></div>\n    <div class=\"wall-stat wall-hidden-stat\"><span>{{t.avg_response}}</span><b id=\"statAvg\">{{stats.avg_response}}</b></div>\n    <div class=\"wall-stat wall-hidden-stat\"><span>{{t.unassigned}}</span><b id=\"statUnassigned\">{{stats.unassigned}}</b></div>\n  </div>\n\n  <div class=\"wall-filters\">\n    <button class=\"filter-chip active\" onclick=\"setWallFilter('all',this)\">{{t.all}}</button>\n    {% for z in zones %}<button class=\"filter-chip\" onclick=\"setWallFilter('{{z}}',this)\">{{z}}</button>{% endfor %}\n    <label class=\"show-idle\"><input id=\"idleToggle\" type=\"checkbox\" checked onchange=\"toggleIdle(this.checked)\"> {{t.show_idle}}</label>\n    <div class=\"wall-sort\"><label>Sort <select id=\"wallSort\" onchange=\"setWallSort(this.value)\"><option value=\"critical\">Critical first</option><option value=\"oldest\">Oldest call</option><option value=\"room\">Room order</option></select></label></div>\n    <span id=\"wallLiveDot\" class=\"wall-live ok\">● LIVE</span><span id=\"wallConnection\" class=\"wall-connection\">{{t.connected}}</span><span id=\"wallUpdated\" class=\"wall-updated\">Updated now</span>\n  </div>\n\n  <div id=\"wallRooms\" class=\"wall-room-grid view-compact\">\n    {% for r in rooms_view %}\n    <article class=\"wall-room-card status-{{r.status}} alert-{{r.alert_stage}} {% if r.over_sla %}sla-breach{% endif %}\" data-zone=\"{{r.zone}}\" data-status=\"{{r.status}}\" data-alert-stage=\"{{r.alert_stage}}\" data-pulse-stage=\"\" data-room=\"{{r.room}}\" data-created=\"{{r.created_at}}\" data-stop=\"{{r.timer_stop or ''}}\" data-elapsed=\"{{r.elapsed_seconds}}\">\n      <div class=\"wall-room-top\"><div><b>{{r.room}}</b><span>{{r.zone}}</span></div><span class=\"wall-badge\">{{r.status_label}}</span></div>\n      <div class=\"wall-room-body\">\n        <div class=\"room-reason\">{{r.reason_label}}</div>\n        <div class=\"room-meta\"><span>{{t.assigned_nurse}}</span><b>{{r.nurse}}</b></div>\n        <div class=\"room-elapsed-pro\"><span class=\"elapsed-caption\">{{t.elapsed}}</span><b class=\"wall-live-timer\">{{r.elapsed_label}}</b><div class=\"sla-track\"><i style=\"width:{{r.sla_progress}}%\"></i></div></div>\n      </div>\n      {% if r.takeover %}<div class=\"wall-takeover\">{{t.charge_takeover}}: {{r.takeover}}</div>{% endif %}\n    </article>\n    {% endfor %}\n  </div>\n\n  <div class=\"wall-bottom\">\n    <section class=\"wall-panel\"><h2>{{t.nurse_workload}}</h2><div id=\"wallWorkload\">{% for n in workload %}<div class=\"workload-row\"><b>{{n.name}}</b><span>{{n.rooms}} {{t.rooms}}</span><span>{{n.active}} {{t.calls}}</span></div>{% endfor %}</div></section>\n    <section class=\"wall-panel\"><h2>{{t.pending_handover}}</h2><div id=\"wallHandovers\">{% if not handovers %}<div class=\"muted\">{{t.none}}</div>{% endif %}{% for h in handovers %}<div class=\"handover-row\"><b>{{h.room}}</b><span>{{h.from_nurse}} → {{h.to_nurse}}</span></div>{% endfor %}</div></section>\n    <section class=\"wall-panel\"><h2>{{t.recent_resolved}}</h2><div id=\"wallRecent\">{% if not recent %}<div class=\"muted\">{{t.none}}</div>{% endif %}{% for r in recent %}<div class=\"recent-row\"><b>{{r.room}}</b><span>{{r.response}}</span><span>{{r.nurse}}</span></div>{% endfor %}</div></section>\n  </div>\n</section>\n{% endblock %}\n{% block scripts %}\n<script>\nwindow.WALLBOARD_LANG='{{lang}}';\nwindow.WALLBOARD_I18N={{ wall_i18n_json|safe }};\nvar wallFilter='all', wallSound=false, showIdle=true, wallView=localStorage.getItem('wallboard-view')||'compact', wallSort=localStorage.getItem('wallboard-sort')||'critical', seenCalls=new Set({{ call_ids_json|safe }}), seenStages={{ call_stages_json|safe }}, wallLastUpdate=Date.now();\nvar wallSoundRepeatSeconds=Math.max(1,Number({{sound_repeat_seconds|default(120)}})||120),wallSoundBuckets={};\n\nfunction setWallView(v,el){wallView=v;localStorage.setItem('wallboard-view',v);var box=document.getElementById('wallRooms');box.className='wall-room-grid view-'+v;document.querySelectorAll('.view-btn').forEach(function(x){x.classList.toggle('active',x.dataset.view===v)});renderZoneHeaders();applyWallFilter()}\nfunction renderZoneHeaders(){document.querySelectorAll('.zone-divider').forEach(function(x){x.remove()});if(wallView!=='zones')return;var box=document.getElementById('wallRooms'),cards=[...box.querySelectorAll('.wall-room-card')],seen={};cards.forEach(function(c){var z=c.dataset.zone||'Other';if(!seen[z]){var d=document.createElement('div');d.className='zone-divider';d.dataset.zone=z;d.innerHTML='<span>'+z+'</span>';box.insertBefore(d,c);seen[z]=true}})}\nfunction setWallFilter(v,el){wallFilter=v;document.querySelectorAll('.filter-chip').forEach(function(x){x.classList.remove('active')});el.classList.add('active');applyWallFilter()}\nfunction toggleIdle(v){showIdle=v;applyWallFilter()}\nfunction applyWallFilter(){document.querySelectorAll('.wall-room-card').forEach(function(c){var zoneOk=(wallFilter==='all'||c.dataset.zone===wallFilter),idleOk=(showIdle||c.dataset.status!=='ready');c.style.display=(zoneOk&&idleOk)?'':'none'});document.querySelectorAll('.zone-divider').forEach(function(d){var any=[...document.querySelectorAll('.wall-room-card[data-zone=\"'+d.dataset.zone+'\"]')].some(function(c){return c.style.display!=='none'});d.style.display=any?'':'none'})}\nfunction fmtSec(s){s=Math.max(0,Math.round(s||0));if(s>=3600)return String(Math.floor(s/3600)).padStart(2,'0')+':'+String(Math.floor((s%3600)/60)).padStart(2,'0')+':'+String(s%60).padStart(2,'0');return String(Math.floor(s/60)).padStart(2,'0')+':'+String(s%60).padStart(2,'0')}\nfunction setWallSort(v){wallSort=v;localStorage.setItem('wallboard-sort',v);refreshWallboard()}\nfunction stageRank(s){return ({critical:0,warning:1,fresh:2,takeover:3,acknowledged:4,arrived:5,idle:6})[s]??9}\nfunction sortRooms(items){return [...items].sort(function(a,b){if(wallSort==='room')return String(a.room).localeCompare(String(b.room),undefined,{numeric:true});if(wallSort==='oldest'){if(a.id&&!b.id)return-1;if(!a.id&&b.id)return 1;return (b.elapsed_seconds||0)-(a.elapsed_seconds||0)}var ar=stageRank(a.alert_stage),br=stageRank(b.alert_stage);if(ar!==br)return ar-br;return (b.elapsed_seconds||0)-(a.elapsed_seconds||0)})}\nfunction applyWallPulseStage(card,sec){\n  var status=card.dataset.status||'';\n  var stage='idle';\n\n  // Only Arrived / Ready / Closed stop visual escalation.\n  if(status==='arrived'||status==='ready'||status==='closed')stage=(status==='arrived'?'arrived':'idle');\n  else if(status==='taken_over')stage='takeover';\n  else if(status==='escalated')stage='critical';\n  else if(status==='new'||status==='acknowledged'){\n    if(sec<120)stage='fresh';\n    else if(sec<240)stage='warning';\n    else stage='critical';\n  }\n\n  // Do NOT remove/re-add the same animation class every second.\n  // Restarting it on each timer tick prevented the pulse from reaching its visible phase.\n  if(card.dataset.pulseStage===stage)return;\n  card.classList.remove('alert-fresh','alert-warning','alert-critical','alert-takeover','alert-arrived','alert-idle','alert-acknowledged');\n  card.classList.add('alert-'+stage);\n  card.dataset.alertStage=stage;\n  card.dataset.pulseStage=stage;\n}\nfunction updateWallTimers(){document.querySelectorAll('.wall-room-card').forEach(function(card){var sec=Number(card.dataset.elapsed||0);if(card.dataset.created){var start=new Date(card.dataset.created).getTime();if(start){var end=card.dataset.stop?new Date(card.dataset.stop).getTime():Date.now();sec=Math.max(0,Math.floor((end-start)/1000));card.dataset.elapsed=sec;var el=card.querySelector('.wall-live-timer');if(el)el.textContent=fmtSec(sec);var bar=card.querySelector('.sla-track i');if(bar)bar.style.width=Math.min(100,(sec/240)*100)+'%'}}applyWallPulseStage(card,sec)})}\nfunction updateWallFreshness(){var s=Math.max(0,Math.floor((Date.now()-wallLastUpdate)/1000)),el=document.getElementById('wallUpdated');if(el)el.textContent=s<2?'Updated now':'Updated '+s+'s ago'}\nfunction stageAlert(c){var old=seenStages[c.id];if(old&&old!==c.alert_stage){if(c.alert_stage==='warning')beep(660,.16,2);if(c.alert_stage==='critical'){beep(480,.20,3);speakCall(c,true)}if(c.alert_stage==='takeover')beep(740,.16,2)}seenStages[c.id]=c.alert_stage}\nfunction wallStatusLabel(s){return (window.WALLBOARD_I18N.status||{})[s]||s.replaceAll('_',' ')}\nfunction wallReasonLabel(s){return (window.WALLBOARD_I18N.reason||{})[s]||s}\nfunction beep(freq,dur,repeats){freq=freq||880;dur=dur||.18;repeats=repeats||2;try{var AC=window.AudioContext||window.webkitAudioContext;var ctx=window._wallAC||(window._wallAC=new AC());var t=ctx.currentTime;for(var i=0;i<repeats;i++){var o=ctx.createOscillator(),g=ctx.createGain();o.frequency.value=freq;o.connect(g);g.connect(ctx.destination);g.gain.setValueAtTime(.12,t);g.gain.exponentialRampToValueAtTime(.001,t+dur);o.start(t);o.stop(t+dur);t+=dur+.08}}catch(e){}}\nfunction speakCall(c,escalated){if(!wallSound||!('speechSynthesis'in window))return;var msg=window.WALLBOARD_LANG==='ar'?(escalated?('تم تصعيد النداء للغرفة '+c.room):('نداء جديد من الغرفة '+c.room)):(escalated?('Escalated call, room '+c.room):('New nurse call, room '+c.room));speechSynthesis.cancel();var u=new SpeechSynthesisUtterance(msg);u.lang=window.WALLBOARD_LANG==='ar'?'ar-AE':'en-US';u.rate=.9;speechSynthesis.speak(u)}\nfunction repeatWaitingSound(c){
  if(!wallSound||!c||!c.id)return;
  if(['arrived','ready','closed','resolved'].includes(c.status)){delete wallSoundBuckets[c.id];return}
  if(!['new','acknowledged','escalated','taken_over'].includes(c.status))return;
  var interval=Math.max(1,Number(wallSoundRepeatSeconds)||120);
  var elapsed=Math.max(0,Number(c.elapsed_seconds)||0);
  var bucket=Math.floor(elapsed/interval);
  if(bucket<1)return;
  if((wallSoundBuckets[c.id]||0)>=bucket)return;
  wallSoundBuckets[c.id]=bucket;
  beep(c.status==='escalated'?480:760,.18,c.status==='escalated'?3:2);
  if('speechSynthesis'in window){
    var msg=window.WALLBOARD_LANG==='ar'?('تذكير: المريض ما زال ينتظر في الغرفة '+c.room):('Reminder: patient is still waiting, room '+c.room);
    speechSynthesis.cancel();
    var u=new SpeechSynthesisUtterance(msg);
    u.lang=window.WALLBOARD_LANG==='ar'?'ar-AE':'en-US';
    u.rate=.9;
    speechSynthesis.speak(u)
  }
}
function enableWallboardSound(){wallSound=true;localStorage.setItem('wallboard-sound','1');beep(880,.12,1);document.getElementById('soundBtn').textContent='🔊 '+window.WALLBOARD_I18N.sound_on;refreshWallboard()}\nfunction toggleWallFullscreen(){if(!document.fullscreenElement){if(document.documentElement.requestFullscreen)document.documentElement.requestFullscreen()}else if(document.exitFullscreen)document.exitFullscreen()}\nfunction renderRooms(d){var box=document.getElementById('wallRooms'),h='';sortRooms(d.rooms_view).forEach(function(r){var stop=r.timer_stop||'',progress=Math.min(100,((r.elapsed_seconds||0)/240)*100);h+='<article class=\"wall-room-card status-'+r.status+' alert-'+r.alert_stage+' '+(r.over_sla?'sla-breach':'')+'\" data-zone=\"'+r.zone+'\" data-status=\"'+r.status+'\" data-alert-stage=\"'+r.alert_stage+'\" data-pulse-stage=\"\" data-room=\"'+r.room+'\" data-created=\"'+(r.created_at||'')+'\" data-stop=\"'+stop+'\" data-elapsed=\"'+(r.elapsed_seconds||0)+'\"><div class=\"wall-room-top\"><div><b>'+r.room+'</b><span>'+r.zone+'</span></div><span class=\"wall-badge\">'+wallStatusLabel(r.status)+'</span></div><div class=\"wall-room-body\"><div class=\"room-reason\">'+wallReasonLabel(r.reason)+'</div><div class=\"room-meta\"><span>'+window.WALLBOARD_I18N.assigned_nurse+'</span><b>'+r.nurse+'</b></div><div class=\"room-elapsed-pro\"><span class=\"elapsed-caption\">'+window.WALLBOARD_I18N.elapsed+'</span><b class=\"wall-live-timer\">'+fmtSec(r.elapsed_seconds||0)+'</b><div class=\"sla-track\"><i style=\"width:'+progress+'%\"></i></div></div></div>'+(r.takeover?'<div class=\"wall-takeover\">'+window.WALLBOARD_I18N.charge_takeover+': '+r.takeover+'</div>':'')+'</article>'});box.innerHTML=h;box.className='wall-room-grid view-'+wallView;renderZoneHeaders();applyWallFilter();updateWallTimers()}\nfunction renderWorkload(items){var h='';items.forEach(function(n){var load=Math.min(100,n.rooms*18+n.active*30);h+='<div class=\"workload-card\"><div><b>'+n.name+'</b><small>'+n.rooms+' '+window.WALLBOARD_I18N.rooms+' · '+n.active+' '+window.WALLBOARD_I18N.calls+'</small></div><div class=\"workload-meter\"><i style=\"width:'+load+'%\"></i></div></div>'});document.getElementById('wallWorkload').innerHTML=h||'<div class=\"muted\">'+window.WALLBOARD_I18N.none+'</div>'}\nfunction renderHandovers(items){var h='';items.forEach(function(x){h+='<div class=\"handover-row\"><b>'+x.room+'</b><span>'+x.from_nurse+' → '+x.to_nurse+'</span></div>'});document.getElementById('wallHandovers').innerHTML=h||'<div class=\"muted\">'+window.WALLBOARD_I18N.none+'</div>'}\nfunction renderRecent(items){var h='';items.forEach(function(x){h+='<div class=\"recent-row\"><b>'+x.room+'</b><span>'+x.response+'</span><span>'+x.nurse+'</span></div>'});document.getElementById('wallRecent').innerHTML=h||'<div class=\"muted\">'+window.WALLBOARD_I18N.none+'</div>'}\nasync function refreshWallboard(){try{var r=await fetch('/api/wallboard',{cache:'no-store'});if(!r.ok)throw new Error('offline');var d=await r.json(),c=document.getElementById('wallConnection'),dot=document.getElementById('wallLiveDot');c.textContent=window.WALLBOARD_I18N.connected;c.classList.remove('bad');if(dot){dot.textContent='● LIVE';dot.classList.remove('bad');dot.classList.add('ok')}wallLastUpdate=Date.now();document.getElementById('statRooms').textContent=d.stats.total_rooms;document.getElementById('statActive').textContent=d.stats.active;document.getElementById('statEscalated').textContent=d.stats.escalated;document.getElementById('statOverSla').textContent=d.stats.over_sla;document.getElementById('statAvg').textContent=d.stats.avg_response;document.getElementById('statUnassigned').textContent=d.stats.unassigned;wallSoundRepeatSeconds=Math.max(1,Number(d.sound_repeat_seconds)||120);d.calls.forEach(function(x){if(!seenCalls.has(x.id)){seenCalls.add(x.id);if(wallSound){beep(x.status==='escalated'?520:880,.18,x.status==='escalated'?3:2);speakCall(x,x.status==='escalated')}}stageAlert(x);repeatWaitingSound(x)});var openSoundIds=new Set(d.calls.filter(function(x){return !['arrived','ready','closed','resolved'].includes(x.status)}).map(function(x){return String(x.id)}));Object.keys(wallSoundBuckets).forEach(function(id){if(!openSoundIds.has(String(id)))delete wallSoundBuckets[id]});renderRooms(d);renderWorkload(d.workload);renderHandovers(d.handovers);renderRecent(d.recent)}catch(e){var el=document.getElementById('wallConnection'),dot=document.getElementById('wallLiveDot');el.textContent=window.WALLBOARD_I18N.disconnected;el.classList.add('bad');if(dot){dot.textContent='● OFFLINE';dot.classList.remove('ok');dot.classList.add('bad')}}}\nsetInterval(refreshWallboard,3000);setInterval(function(){document.getElementById('wallClock').textContent=new Date().toLocaleTimeString([], {hour12:false});updateWallTimers();updateWallFreshness()},1000);\ndocument.getElementById('idleToggle').checked=true;document.getElementById('wallSort').value=wallSort;setWallView(wallView,document.querySelector('.view-btn[data-view=\"'+wallView+'\"]'));updateWallTimers();if(localStorage.getItem('wallboard-sound')==='1'){wallSound=true;document.getElementById('soundBtn').textContent='🔊 '+window.WALLBOARD_I18N.sound_on}\n</script>\n{% endblock %}"
}
APP_CSS = ":root{--blue:#0059ad;--deep:#023d7b;--ink:#172033;--muted:#667085;--bg:#f3f6fa;--red:#d92d20;--amber:#f79009;--green:#039855;--line:#dfe5ec}*{box-sizing:border-box}body{margin:0;font-family:Inter,Segoe UI,Arial,sans-serif;background:var(--bg);color:var(--ink)}.topbar{height:76px;background:#fff;border-bottom:1px solid var(--line);display:flex;align-items:center;justify-content:space-between;padding:10px 28px;position:sticky;top:0;z-index:10}.brand{display:flex;align-items:center;gap:12px}.brand img{width:46px;height:46px}.brand b,.brand span{display:block}.brand span{font-size:12px;color:var(--muted)}.userbox{text-align:right}.userbox span,.userbox small{display:block}.userbox a{font-size:12px;color:var(--blue)}.nav{background:#fff;padding:10px 28px;display:flex;gap:18px;border-bottom:1px solid var(--line)}.nav a{text-decoration:none;color:var(--deep);font-weight:650}.container{max-width:1280px;margin:0 auto;padding:28px}.page-head{display:flex;justify-content:space-between;align-items:center;margin-bottom:22px}.page-head h1{margin:4px 0 6px;font-size:32px}.page-head p,.muted{color:var(--muted)}.eyebrow{text-transform:uppercase;letter-spacing:.12em;font-size:12px;font-weight:800;color:var(--blue)}.stats{display:grid;grid-template-columns:repeat(3,1fr);gap:14px;margin:18px 0 26px}.stats.four{grid-template-columns:repeat(4,1fr)}.stat{background:#fff;border:1px solid var(--line);border-radius:18px;padding:18px}.stat b{font-size:30px;display:block}.stat span{color:var(--muted)}.stat.red{border-left:5px solid var(--red)}.stat.amber{border-left:5px solid var(--amber)}.grid.rooms{display:grid;grid-template-columns:repeat(auto-fit,minmax(250px,1fr));gap:16px}.room-card{background:#fff;border:1px solid var(--line);border-radius:20px;padding:18px;min-height:220px;box-shadow:0 8px 24px rgba(16,24,40,.04)}.room-card.hot{border:2px solid #f1a9a5}.room-top{display:flex;justify-content:space-between;align-items:center}.room-top b{font-size:25px}.room-top span,.assignment{font-size:13px;color:var(--muted)}.assignment{margin:8px 0 16px}.call-status{display:inline-block;margin-top:18px;padding:7px 10px;border-radius:999px;background:#eef2f6;font-weight:800;font-size:12px}.call-status.new,.badge.new{background:#fee4e2;color:#b42318}.call-status.escalated,.badge.escalated{background:#fef0c7;color:#b54708}.call-status.acknowledged,.badge.acknowledged{background:#e0f2fe;color:#026aa2}.call-status.taken_over,.badge.taken_over{background:#f3e8ff;color:#7f56d9}.call-status.arrived,.badge.arrived,.badge.resolved{background:#dcfae6;color:#067647}.timer{font-size:36px;font-variant-numeric:tabular-nums;font-weight:800;margin:9px 0}.ready{margin-top:34px;color:var(--green);font-weight:800}.actions{display:flex;gap:8px;flex-wrap:wrap;margin-top:14px}.btn{border:0;border-radius:12px;padding:11px 15px;font-weight:750;cursor:pointer}.btn.primary{background:var(--blue);color:#fff}.btn.secondary{background:#eaf2fb;color:var(--deep)}.btn.success{background:var(--green);color:#fff}.btn.danger{background:var(--red);color:#fff}.btn.wide{width:100%;font-size:16px}.alert{padding:12px;border-radius:12px;margin:12px 0}.alert.danger{background:#fee4e2;color:#b42318}.alert.small{font-size:12px}.login-shell,.patient-shell{display:grid;place-items:center;min-height:calc(100vh - 150px)}.login-card,.patient-card{width:min(480px,100%);background:#fff;border:1px solid var(--line);border-radius:28px;padding:32px;text-align:center;box-shadow:0 18px 50px rgba(16,24,40,.08)}.hero-logo{width:104px;height:104px;object-fit:contain}.login-card label{text-align:left;display:block;font-weight:700;margin:15px 0}.login-card input,.room-card select{width:100%;margin-top:7px;border:1px solid #cfd7e2;border-radius:12px;padding:12px;font-size:15px}.demo{margin-top:22px;padding:14px;background:#f7f9fc;border-radius:14px;font-size:12px}.demo span{display:block;color:var(--muted);margin-top:4px}.patient-card h1{font-size:44px;margin:5px}.reason-grid{display:grid;grid-template-columns:1fr 1fr;gap:10px;margin:22px 0}.reason{padding:12px;border:1px solid var(--line);border-radius:12px;background:#fff;cursor:pointer}.reason.selected{border:2px solid var(--blue);background:#eef6ff}.call-btn{width:220px;height:220px;border-radius:50%;border:10px solid #d7e8fb;background:var(--blue);color:#fff;font-size:46px;box-shadow:0 15px 30px rgba(0,89,173,.25);cursor:pointer}.call-btn span{font-size:20px;display:block;margin-top:8px}.call-active{padding:28px}.pulse{width:20px;height:20px;background:var(--red);border-radius:50%;margin:0 auto;box-shadow:0 0 0 0 rgba(217,45,32,.5);animation:pulse 1.5s infinite}@keyframes pulse{70%{box-shadow:0 0 0 20px rgba(217,45,32,0)}}.table-wrap{background:#fff;border:1px solid var(--line);border-radius:18px;overflow:auto}table{width:100%;border-collapse:collapse}th,td{text-align:left;padding:14px;border-bottom:1px solid #edf0f4;font-size:14px}th{background:#f8fafc}.badge{padding:6px 9px;border-radius:999px;font-size:12px}.two-col{display:grid;grid-template-columns:1fr 1fr;gap:16px}.panel{background:#fff;border:1px solid var(--line);border-radius:18px;padding:18px}.list-row{display:grid;grid-template-columns:1fr 1fr 1fr;gap:8px;padding:12px 0;border-bottom:1px solid #edf0f4;font-size:14px}@media(max-width:700px){.container{padding:16px}.topbar{padding:8px 14px}.brand b{font-size:14px}.nav{padding:8px 14px;overflow:auto}.stats,.stats.four,.two-col{grid-template-columns:1fr 1fr}.page-head{align-items:flex-start;gap:12px}.page-head h1{font-size:26px}.call-btn{width:190px;height:190px}.userbox span{display:none}}.handover-box{display:flex;gap:8px;margin-top:18px;border-top:1px solid #edf0f4;padding-top:14px}.handover-box select{flex:1;border:1px solid #cfd7e2;border-radius:10px;padding:9px;background:#fff}\n.admin-tabs{display:flex;gap:10px;flex-wrap:wrap;margin:0 0 18px}.admin-tabs a{background:#fff;border:1px solid var(--line);padding:9px 13px;border-radius:999px;text-decoration:none;color:var(--deep);font-weight:750}.admin-section{margin-bottom:18px;scroll-margin-top:110px}.section-head{display:flex;justify-content:space-between;align-items:center;margin-bottom:14px}.admin-form{margin:12px 0 18px}.grid-form{display:grid;grid-template-columns:repeat(5,minmax(130px,1fr));gap:10px}.admin-form input,.admin-form select,.inline-form select,.row-actions input{border:1px solid #cfd7e2;border-radius:10px;padding:10px;background:#fff;min-width:0}.row-actions,.inline-form{display:flex;gap:6px;align-items:center;flex-wrap:wrap}.row-actions form{display:flex;gap:6px}.row-actions input{width:145px}.status-pill{padding:9px 12px;border-radius:999px;font-weight:800;font-size:12px}.status-pill.ok{background:#dcfae6;color:#067647}.status-pill.warn{background:#fef0c7;color:#b54708}.system-grid{display:grid;grid-template-columns:repeat(4,1fr);gap:12px}.system-grid>div{background:#f8fafc;border:1px solid var(--line);border-radius:14px;padding:14px}.system-grid span,.system-grid b{display:block}.system-grid span{color:var(--muted);font-size:12px;margin-bottom:5px}.mono-link{font-family:ui-monospace,SFMono-Regular,Consolas,monospace;font-size:12px;color:var(--blue);word-break:break-all}.btn:disabled{opacity:.45;cursor:not-allowed}@media(max-width:900px){.grid-form{grid-template-columns:1fr 1fr}.system-grid{grid-template-columns:1fr 1fr}}@media(max-width:600px){.grid-form,.system-grid{grid-template-columns:1fr}.row-actions input{width:120px}}\n"
APP_CSS += "\n/* QR + bilingual patient journey */\n.qr-action{text-decoration:none;display:inline-block}.qr-shell{display:grid;place-items:center;min-height:calc(100vh - 150px)}.qr-poster{width:min(520px,100%);background:#fff;border:1px solid var(--line);border-radius:28px;padding:28px;text-align:center;box-shadow:0 18px 50px rgba(16,24,40,.08)}.qr-logo{width:86px;height:86px;object-fit:contain}.qr-poster h1{font-size:42px;margin:6px 0}.qr-poster h2{margin:8px 0}.qr-ar{font-family:Tahoma,Arial,sans-serif}.qr-image{width:min(330px,85vw);height:auto;margin:16px auto;display:block}.qr-print-actions{display:flex;gap:10px;justify-content:center;margin-top:18px}.patient-lang-switch{display:flex;align-items:center;justify-content:flex-end;gap:6px;margin-bottom:8px}.lang-chip{border:1px solid #d0d7e2;background:#fff;border-radius:999px;padding:5px 9px;font-size:12px;font-weight:800;color:var(--deep);cursor:pointer}.lang-chip.active{background:#eaf2fb;border-color:#8eb8e5}.patient-question{font-size:22px;margin:14px 0 4px}.patient-note{margin-top:0}.reason{display:flex;align-items:center;justify-content:center;gap:7px;min-height:58px}.reason-icon{font-size:20px}.physical-note{margin-top:24px;padding:12px 14px;background:#fff7e8;border:1px solid #f4d49b;border-radius:14px;font-size:13px;line-height:1.5}.patient-lang[dir=\"rtl\"] .patient-card{text-align:right}.patient-lang[dir=\"rtl\"] .eyebrow,.patient-lang[dir=\"rtl\"] h1,.patient-lang[dir=\"rtl\"] .patient-question,.patient-lang[dir=\"rtl\"] .patient-note,.patient-lang[dir=\"rtl\"] .call-active{text-align:center}.patient-lang[dir=\"rtl\"] .reason{font-family:Tahoma,Arial,sans-serif}.patient-lang[dir=\"rtl\"] .call-btn span{font-family:Tahoma,Arial,sans-serif}@media print{.topbar,.nav,.qr-print-actions{display:none!important}.container{padding:0}.qr-shell{min-height:auto}.qr-poster{box-shadow:none;border:none;width:100%;padding:10mm}.qr-image{width:95mm}.qr-logo{width:25mm;height:25mm}}"
APP_CSS += ".wallboard-shell{max-width:1600px;margin:0 auto}.wallboard-head{display:flex;justify-content:space-between;align-items:flex-start;gap:24px;margin-bottom:18px}.wallboard-head h1{font-size:38px;margin:2px 0 4px}.wallboard-head p{margin:0;color:var(--muted)}.wallboard-tools{display:flex;align-items:center;gap:10px;flex-wrap:wrap;justify-content:flex-end}.wall-clock{font-size:28px;font-weight:850;font-variant-numeric:tabular-nums;background:#fff;border:1px solid var(--line);padding:8px 13px;border-radius:14px}.lang-mini{display:flex;gap:5px;align-items:center}.lang-mini a{padding:5px 8px;border-radius:999px;text-decoration:none;color:var(--deep);font-weight:800;font-size:12px;border:1px solid var(--line);background:#fff}.lang-mini a.active{background:#eaf2fb}.wall-stats{display:grid;grid-template-columns:repeat(5,1fr);gap:12px;margin-bottom:16px}.wall-stat{background:#fff;border:1px solid var(--line);border-radius:16px;padding:14px 16px}.wall-stat span{display:block;color:var(--muted);font-size:12px}.wall-stat b{font-size:28px}.wall-stat.danger{border-left:5px solid var(--red)}.wall-stat.overdue{border-left:5px solid var(--amber)}.wall-filters{display:flex;align-items:center;gap:8px;flex-wrap:wrap;margin-bottom:16px}.filter-chip{border:1px solid var(--line);background:#fff;padding:7px 11px;border-radius:999px;font-weight:750;color:var(--deep);cursor:pointer}.filter-chip.active{background:var(--deep);color:#fff}.wall-live{margin-left:auto;color:var(--green);font-weight:900}.wall-connection{font-size:12px;color:var(--green);font-weight:800}.wall-connection.bad{color:var(--red)}.wall-call-grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(280px,1fr));gap:14px;min-height:250px}.wall-call-card{background:#fff;border:2px solid var(--line);border-radius:20px;padding:17px;box-shadow:0 8px 20px rgba(16,24,40,.04)}.wall-call-card.status-new{border-color:#f0a6a1}.wall-call-card.status-escalated,.wall-call-card.status-taken_over{border-color:#f79009}.wall-call-card.status-arrived{border-color:#66c98b}.wall-call-card.sla-breach{animation:wallPulse 1.1s infinite;background:#fff7f6}@keyframes wallPulse{0%,100%{box-shadow:0 0 0 0 rgba(217,45,32,.12)}50%{box-shadow:0 0 0 8px rgba(217,45,32,.13)}}.wall-call-top{display:flex;justify-content:space-between;gap:10px;align-items:flex-start}.wall-call-top b{font-size:30px;display:block}.wall-call-top span{font-size:12px;color:var(--muted)}.wall-badge{background:#eef2f6!important;color:var(--deep)!important;padding:6px 9px;border-radius:999px;font-weight:850}.status-escalated .wall-badge{background:#fef0c7!important;color:#b54708!important}.status-new .wall-badge{background:#fee4e2!important;color:#b42318!important}.status-arrived .wall-badge{background:#dcfae6!important;color:#067647!important}.wall-reason{font-size:21px;font-weight:850;margin:18px 0}.wall-call-mid{display:grid;grid-template-columns:1fr auto;gap:18px}.wall-call-mid span{display:block;font-size:11px;color:var(--muted)}.wall-call-mid b{font-size:18px}.wall-takeover{margin-top:12px;padding:9px;background:#f3e8ff;border-radius:10px;font-size:13px;color:#6941c6}.wall-empty{grid-column:1/-1;display:grid;place-items:center;background:#fff;border:1px dashed var(--line);border-radius:20px;min-height:220px;color:var(--muted);font-size:20px}.wall-bottom{display:grid;grid-template-columns:1.15fr 1fr 1fr;gap:14px;margin-top:16px}.wall-panel{background:#fff;border:1px solid var(--line);border-radius:18px;padding:16px}.wall-panel h2{font-size:17px;margin:0 0 10px}.workload-row,.handover-row,.recent-row{display:grid;grid-template-columns:1.2fr .8fr .8fr;gap:8px;padding:9px 0;border-bottom:1px solid #edf0f4;font-size:13px}.handover-row{grid-template-columns:.5fr 1.5fr}.recent-row{grid-template-columns:.5fr .7fr 1fr}.wallboard-shell[data-lang=\"ar\"]{direction:rtl}.wallboard-shell[data-lang=\"ar\"] .wall-live{margin-left:0;margin-right:auto}@media(max-width:1000px){.wall-stats{grid-template-columns:repeat(3,1fr)}.wall-bottom{grid-template-columns:1fr}.wallboard-head{flex-direction:column}.wallboard-tools{justify-content:flex-start}}@media(max-width:650px){.wall-stats{grid-template-columns:1fr 1fr}.wall-call-grid{grid-template-columns:1fr}.wall-call-top b{font-size:26px}}"
APP_CSS += ".view-switch{display:flex;gap:4px;background:#eef3f8;padding:4px;border-radius:12px}.view-btn{border:0;background:transparent;padding:7px 9px;border-radius:9px;font-size:12px;font-weight:800;color:var(--deep);cursor:pointer}.view-btn.active{background:#fff;box-shadow:0 1px 4px rgba(16,24,40,.12)}.show-idle{display:flex;align-items:center;gap:6px;font-size:12px;color:var(--muted);font-weight:700}.wall-stats{grid-template-columns:repeat(6,1fr)}.wall-room-grid{display:grid;gap:10px;align-items:stretch}.wall-room-grid.view-compact{grid-template-columns:repeat(auto-fill,minmax(180px,1fr))}.wall-room-grid.view-cards{grid-template-columns:repeat(auto-fill,minmax(280px,1fr));gap:14px}.wall-room-grid.view-list{grid-template-columns:1fr;gap:6px}.wall-room-grid.view-zones{grid-template-columns:repeat(auto-fill,minmax(210px,1fr));gap:10px}.wall-room-card{background:#fff;border:1px solid var(--line);border-radius:16px;padding:12px;min-height:126px;box-shadow:0 4px 14px rgba(16,24,40,.035);position:relative;overflow:hidden}.wall-room-card.status-ready{border-color:#dce6df;background:#fbfefc}.wall-room-card.status-new{border-color:#f0a6a1}.wall-room-card.status-escalated,.wall-room-card.status-taken_over{border-color:#f79009}.wall-room-card.status-arrived{border-color:#66c98b}.wall-room-card.sla-breach{animation:wallPulse 1.1s infinite;background:#fff7f6}.wall-room-top{display:flex;justify-content:space-between;gap:8px}.wall-room-top b{font-size:23px;display:block;line-height:1}.wall-room-top span{font-size:10px;color:var(--muted)}.wall-room-body{margin-top:9px}.room-reason{font-size:14px;font-weight:850;min-height:20px}.room-meta{margin-top:8px}.room-meta span{display:block;color:var(--muted);font-size:9px;text-transform:uppercase;letter-spacing:.04em}.room-meta b{font-size:13px}.room-elapsed{position:absolute;right:12px;bottom:10px;text-align:right}.wallboard-shell[data-lang=\"ar\"] .room-elapsed{right:auto;left:12px;text-align:left}.wall-room-card.status-ready .room-reason{color:var(--green)}.view-compact .wall-room-card{padding:10px;min-height:108px}.view-compact .wall-room-top b{font-size:20px}.view-compact .room-meta{margin-top:5px}.view-compact .room-meta span{display:none}.view-compact .room-meta b{font-size:12px}.view-compact .room-elapsed{bottom:8px}.view-cards .wall-room-card{min-height:170px;padding:16px}.view-cards .wall-room-top b{font-size:30px}.view-cards .room-reason{font-size:19px;margin-top:15px}.view-cards .room-meta b{font-size:16px}.view-list .wall-room-card{display:grid;grid-template-columns:130px 1fr 180px 110px;align-items:center;min-height:58px;padding:8px 12px}.view-list .wall-room-top{display:block}.view-list .wall-room-top b{font-size:20px}.view-list .wall-room-body{display:contents}.view-list .room-reason{font-size:14px;min-height:0}.view-list .room-meta{margin:0}.view-list .room-elapsed{position:static;text-align:right}.view-list .wall-takeover{grid-column:1/-1;margin-top:4px}.zone-divider{grid-column:1/-1;font-size:13px;font-weight:900;color:var(--deep);padding:8px 2px 3px;border-bottom:2px solid #dbe6f1;letter-spacing:.04em}.wall-badge{padding:5px 7px;border-radius:999px;font-weight:850;font-size:9px!important}.status-ready .wall-badge{background:#dcfae6!important;color:#067647!important}.wall-bottom{margin-top:12px}.wall-panel{padding:12px}.wall-panel h2{font-size:14px}.workload-row,.handover-row,.recent-row{padding:6px 0;font-size:11px}@media(max-width:1200px){.wall-stats{grid-template-columns:repeat(3,1fr)}.wall-room-grid.view-compact{grid-template-columns:repeat(auto-fill,minmax(165px,1fr))}}@media(max-width:700px){.view-switch{width:100%;overflow:auto}.wall-stats{grid-template-columns:repeat(2,1fr)}.wall-room-grid.view-compact,.wall-room-grid.view-cards,.wall-room-grid.view-zones{grid-template-columns:repeat(2,minmax(0,1fr))}.view-list .wall-room-card{grid-template-columns:100px 1fr}.view-list .room-meta,.view-list .room-elapsed{display:none}}"
APP_CSS += ".alert-legend{display:flex;gap:14px;align-items:center;flex-wrap:wrap;margin:-4px 0 12px;font-size:11px;font-weight:800;color:#475467}.legend-item{display:flex;align-items:center;gap:6px}.legend-item i{width:11px;height:11px;border-radius:50%;display:inline-block}.legend-green i{background:#12b76a;box-shadow:0 0 0 4px rgba(18,183,106,.12)}.legend-amber i{background:#f79009;box-shadow:0 0 0 4px rgba(247,144,9,.12)}.legend-red i{background:#d92d20;box-shadow:0 0 0 4px rgba(217,45,32,.12)}.legend-blue i{background:#2970ff;box-shadow:0 0 0 4px rgba(41,112,255,.12)}.wall-room-card.alert-fresh{border:2px solid #12b76a;animation:fullCardGreen 1.25s ease-in-out infinite}.wall-room-card.alert-warning{border:3px solid #f79009;animation:fullCardAmber .95s ease-in-out infinite}.wall-room-card.alert-critical{border:3px solid #d92d20;animation:fullCardRed .72s ease-in-out infinite}.wall-room-card.alert-takeover{border:3px solid #2970ff;animation:fullCardBlue .95s ease-in-out infinite}.wall-room-card.alert-arrived{border:2px solid #12b76a;background:#ecfdf3;animation:none}.wall-room-card.alert-idle{animation:none}.wall-room-card.alert-fresh .wall-badge{background:#d1fadf!important;color:#05603a!important}.wall-room-card.alert-warning .wall-badge{background:#fef0c7!important;color:#b54708!important}.wall-room-card.alert-critical .wall-badge{background:#fee4e2!important;color:#b42318!important}.wall-room-card.alert-takeover .wall-badge{background:#dbe9ff!important;color:#1849a9!important}@keyframes fullCardGreen{0%,100%{background:#f6fef9;box-shadow:0 0 0 0 rgba(18,183,106,.18),0 6px 16px rgba(18,183,106,.06)}50%{background:#d1fadf;box-shadow:0 0 0 7px rgba(18,183,106,.22),0 10px 30px rgba(18,183,106,.28)}}@keyframes fullCardAmber{0%,100%{background:#fffcf5;box-shadow:0 0 0 0 rgba(247,144,9,.18),0 6px 18px rgba(247,144,9,.08)}50%{background:#fef0c7;box-shadow:0 0 0 9px rgba(247,144,9,.25),0 12px 34px rgba(247,144,9,.32)}}@keyframes fullCardRed{0%,100%{background:#fff7f6;box-shadow:0 0 0 0 rgba(217,45,32,.22),0 8px 20px rgba(217,45,32,.10)}50%{background:#fecaca;box-shadow:0 0 0 11px rgba(217,45,32,.30),0 14px 40px rgba(217,45,32,.38)}}@keyframes fullCardBlue{0%,100%{background:#f5f8ff;box-shadow:0 0 0 0 rgba(41,112,255,.16),0 6px 18px rgba(41,112,255,.08)}50%{background:#dbe9ff;box-shadow:0 0 0 8px rgba(41,112,255,.22),0 12px 32px rgba(41,112,255,.30)}}@media (prefers-reduced-motion:reduce){.wall-room-card.alert-fresh,.wall-room-card.alert-warning,.wall-room-card.alert-critical,.wall-room-card.alert-takeover{animation-duration:2.4s}}"
APP_CSS += ".reassign-box{margin-top:14px;padding:12px;background:#f8fafc;border:1px solid #d9e2ec;border-radius:12px;display:grid;grid-template-columns:1fr 1fr auto;gap:8px}.reassign-title{grid-column:1/-1;display:flex;justify-content:space-between;gap:10px;align-items:center}.reassign-title b{font-size:13px}.reassign-title span{font-size:10px;color:var(--muted)}.reassign-box select{min-width:0;width:100%;border:1px solid #cfd7e2;border-radius:9px;padding:9px;background:#fff;color:var(--deep)}@media(max-width:700px){.reassign-box{grid-template-columns:1fr}.reassign-title{grid-column:auto;display:block}.reassign-title span{display:block;margin-top:3px}}"
APP_CSS += ".room-edit-form{display:grid;grid-template-columns:minmax(85px,110px) minmax(100px,150px) auto;gap:6px;align-items:center}.room-edit-form input{border:1px solid #cfd7e2;border-radius:9px;padding:9px;min-width:0;background:#fff}.admin-room-msg{margin:0 0 12px;padding:11px 13px;border-radius:10px;font-weight:750;font-size:13px}.admin-room-msg.success{background:#dcfae6;color:#067647;border:1px solid #abefc6}.admin-room-msg.error{background:#fee4e2;color:#b42318;border:1px solid #fecdca}.room-delete-btn{background:#d92d20!important;color:#fff!important;border-color:#d92d20!important}@media(max-width:900px){.room-edit-form{grid-template-columns:1fr}.room-edit-form .btn{width:100%}}"
APP_CSS += ".manager-export-actions{display:flex;gap:8px;flex-wrap:wrap}.manager-export-actions a{text-decoration:none}.manager-filter{margin-bottom:14px;padding:12px 14px}.period-buttons{display:flex;gap:7px;flex-wrap:wrap}.custom-dates{display:none;gap:10px;align-items:end;flex-wrap:wrap;margin-top:12px}.custom-dates.show{display:flex}.custom-dates label{display:grid;gap:4px;font-size:12px;font-weight:700;color:var(--muted)}.custom-dates input{border:1px solid #cfd7e2;border-radius:9px;padding:9px;background:#fff}.kpi-stats{grid-template-columns:repeat(8,minmax(110px,1fr))}.manager-grid{display:grid;grid-template-columns:1fr 1fr;gap:14px;margin:14px 0}.management-call-table{margin-top:14px}.management-call-table small{color:var(--muted)}@media(max-width:1200px){.kpi-stats{grid-template-columns:repeat(4,1fr)}}@media(max-width:900px){.manager-grid{grid-template-columns:1fr}}@media(max-width:650px){.kpi-stats{grid-template-columns:repeat(2,1fr)}}"
APP_CSS += ".wall-room-card.alert-acknowledged{border:2px solid #2e90fa;background:#eff8ff;animation:none;box-shadow:0 6px 18px rgba(46,144,250,.12)}.wall-room-card.alert-acknowledged .wall-badge{background:#d1e9ff!important;color:#175cd3!important}"
APP_CSS += ".password-wrap{position:relative}.password-wrap input{padding-right:46px!important;width:100%}.password-eye{position:absolute;right:7px;top:50%;transform:translateY(-50%);border:0;background:transparent;cursor:pointer;font-size:18px;padding:7px;border-radius:8px}.password-eye:hover{background:#eef4fb}.login-security-note{margin-top:14px;padding:10px 12px;border-radius:10px;background:#eff8ff;color:#175cd3;font-size:12px}.notify-setup-shell{min-height:72vh;display:grid;place-items:center;padding:24px}.notify-setup-card{width:min(620px,100%);background:#fff;border:1px solid #d8e2ed;border-radius:18px;padding:28px;box-shadow:0 18px 50px rgba(16,42,67,.10)}.notify-setup-card h1{margin:6px 0}.notify-steps{display:grid;gap:9px;margin:16px 0}.notify-step{display:grid;grid-template-columns:34px 1fr auto;gap:10px;align-items:center;padding:11px;border:1px solid #d8e2ed;border-radius:12px;background:#f8fafc}.notify-step .step-dot{width:28px;height:28px;border-radius:50%;display:grid;place-items:center;background:#e4e7ec;font-weight:900}.notify-step div{display:grid}.notify-step small{color:#667085}.notify-step .step-state{font-size:11px;font-weight:850;color:#667085}.notify-step.working{border-color:#84adff;background:#f5f8ff}.notify-step.ok{border-color:#75e0a7;background:#ecfdf3}.notify-step.ok .step-dot{background:#12b76a;color:#fff}.notify-step.bad{border-color:#fda29b;background:#fef3f2}.notify-step.bad .step-dot{background:#d92d20;color:#fff}.notify-status{padding:11px 12px;margin:12px 0;border-radius:10px;background:#f2f4f7;font-weight:700;font-size:13px}.notify-device{display:grid;grid-template-columns:110px 1fr;gap:5px 10px;margin-top:16px;padding:12px;border-top:1px solid #eaecf0;font-size:12px}.notify-device span{color:#667085}.notify-help{font-size:11px;margin-top:12px}.alert.success{background:#dcfae6;color:#067647;border:1px solid #abefc6}"
APP_CSS += ".notify-setup-card{width:min(720px,100%)}.notify-progress{display:grid;grid-template-columns:repeat(4,1fr);gap:7px;margin:14px 0}.notify-chip{min-height:38px;border:1px solid #cfd8e3;border-radius:999px;padding:7px 9px;display:grid;grid-template-columns:10px 1fr;grid-template-areas:\"dot title\" \"state state\";align-items:center;column-gap:6px;background:#fff;font-size:10px;text-align:center}.notify-chip .chip-dot{grid-area:dot;width:8px;height:8px;border-radius:50%;background:#98a2b3}.notify-chip b{grid-area:title;font-size:10px;white-space:nowrap}.notify-chip .step-state{grid-area:state;font-size:9px;color:#667085;margin-top:2px}.notify-chip.working{border-color:#84adff;background:#f5f8ff}.notify-chip.working .chip-dot{background:#2e90fa}.notify-chip.ok{border-color:#75e0a7;background:#ecfdf3}.notify-chip.ok .chip-dot{background:#12b76a}.notify-chip.bad{border-color:#fda29b;background:#fef3f2}.notify-chip.bad .chip-dot{background:#d92d20}@media(max-width:650px){.notify-progress{grid-template-columns:repeat(2,1fr)}}"
APP_CSS += ".notify-mini-shell{width:min(560px,calc(100% - 24px));margin:28px auto}.notify-mini-card{border:1.5px solid #7db4ff;border-radius:12px;background:#f8fbff;padding:10px 12px}.notify-mini-title{font-size:12px;color:#0b2a4a;margin-bottom:3px}.notify-mini-sub{font-size:9px;font-weight:700;color:#475467;margin-bottom:9px}.notify-mini-steps{display:grid;grid-template-columns:repeat(4,1fr);gap:6px}.notify-mini-chip{min-width:0;border:1px solid #cfd8e3;border-radius:999px;background:#fff;padding:5px 7px;display:grid;grid-template-columns:8px 1fr;grid-template-areas:\"dot title\" \"dot state\";column-gap:5px;align-items:center;text-align:center}.notify-mini-chip .chip-dot{grid-area:dot;width:7px;height:7px;border-radius:50%;background:#98a2b3}.notify-mini-chip b{grid-area:title;font-size:8px;white-space:nowrap}.notify-mini-chip small{grid-area:state;font-size:7px;color:#667085;white-space:nowrap}.notify-mini-chip.working{border-color:#84adff;background:#f5f8ff}.notify-mini-chip.working .chip-dot{background:#2e90fa}.notify-mini-chip.ok{border-color:#75e0a7;background:#ecfdf3}.notify-mini-chip.ok .chip-dot{background:#12b76a}.notify-mini-chip.bad{border-color:#fda29b;background:#fff1f0}.notify-mini-chip.bad .chip-dot{background:#d92d20}.notify-mini-status{margin:8px 0 7px;padding:7px 9px;border-radius:8px;background:#f2f4f7;font-size:9px;font-weight:750;color:#344054}.notify-verify-btn{width:100%;border:0;border-radius:8px;padding:9px 12px;background:linear-gradient(90deg,#2468e8,#0795ad);color:white;font-size:11px;font-weight:850;cursor:pointer}.notify-verify-btn.ready{background:linear-gradient(90deg,#079455,#12b76a)}.notify-platform-hint{margin:7px 0;padding:7px 8px;border-radius:8px;background:#fff6ed;border:1px solid #fedf89;color:#93370d;font-size:8px;font-weight:700}@media(max-width:560px){.notify-mini-shell{margin:16px auto}.notify-mini-steps{grid-template-columns:repeat(2,1fr)}.notify-mini-title{font-size:11px}}"
APP_CSS += "/* Responsive clinical UI + command-center wallboard */\nhtml{-webkit-text-size-adjust:100%}body{overflow-x:hidden}.container{width:100%}.topbar,.nav,.container{min-width:0}.nav{overflow-x:auto;white-space:nowrap;-webkit-overflow-scrolling:touch;scrollbar-width:none}.nav::-webkit-scrollbar{display:none}.btn,button,select,input{min-height:44px}.table-wrap{-webkit-overflow-scrolling:touch}.table-wrap table{min-width:720px}\n.wallboard-shell{max-width:1500px;margin:0 auto}.wallboard-head{align-items:flex-start}.wallboard-head h1{margin-bottom:4px}.wallboard-tools{flex-wrap:wrap;justify-content:flex-end}.wall-stats{gap:10px}.wall-stat{padding:14px 16px;border-radius:16px;box-shadow:0 4px 14px rgba(16,24,40,.035)}.wall-stat b{font-size:28px}.wall-filters{gap:9px;align-items:center}.wall-sort{margin-left:auto}.wall-sort label{display:flex;align-items:center;gap:6px;font-size:11px;font-weight:800;color:#667085}.wall-sort select{min-height:34px;border:1px solid #d0d5dd;border-radius:9px;background:#fff;padding:5px 26px 5px 8px;font-weight:750;color:#344054}.wall-live.ok{color:#079455}.wall-live.bad,.wall-connection.bad{color:#d92d20}.wall-updated{font-size:10px;color:#98a2b3;min-width:72px}\n.wall-room-grid.view-compact{grid-template-columns:repeat(auto-fit,minmax(190px,1fr));gap:10px}.wall-room-card{position:relative;overflow:hidden;transition:transform .18s ease,box-shadow .18s ease,border-color .18s ease}.wall-room-card:hover{transform:translateY(-1px)}.wall-room-top b{font-size:20px}.wall-room-top span{font-size:10px}.wall-badge{font-size:9px;padding:6px 8px}.room-reason{font-weight:800;min-height:34px}.room-meta{margin-top:8px}.room-elapsed-pro{display:grid;grid-template-columns:1fr auto;align-items:end;gap:6px;margin-top:10px}.elapsed-caption{font-size:9px;text-transform:uppercase;letter-spacing:.08em;color:#667085;font-weight:800}.wall-live-timer{font-size:28px;line-height:1;font-variant-numeric:tabular-nums;font-weight:900;color:#101828}.sla-track{grid-column:1/-1;height:4px;border-radius:999px;background:#eaecf0;overflow:hidden;margin-top:3px}.sla-track i{display:block;height:100%;background:linear-gradient(90deg,#12b76a 0 50%,#f79009 70%,#d92d20 100%);border-radius:inherit;transition:width .3s linear}.alert-critical .wall-live-timer{color:#b42318}.alert-warning .wall-live-timer{color:#b54708}.alert-fresh .wall-live-timer{color:#067647}.alert-acknowledged .wall-live-timer,.alert-takeover .wall-live-timer{color:#175cd3}\n.wall-bottom{gap:12px}.wall-panel{border-radius:16px;box-shadow:0 5px 18px rgba(16,24,40,.035)}.workload-card{display:grid;grid-template-columns:1fr 90px;gap:10px;align-items:center;padding:9px 0;border-bottom:1px solid #eef2f6}.workload-card:last-child{border-bottom:0}.workload-card div:first-child{display:grid}.workload-card small{color:#667085;margin-top:2px}.workload-meter{height:6px;border-radius:999px;background:#eaecf0;overflow:hidden}.workload-meter i{display:block;height:100%;border-radius:inherit;background:linear-gradient(90deg,#12b76a,#f79009,#d92d20)}\n@media(max-width:1024px){.container{padding:18px}.topbar{padding:8px 16px;height:auto;min-height:68px}.nav{padding:9px 16px;gap:12px}.page-head{align-items:flex-start;gap:14px}.page-head h1{font-size:28px}.stats,.stats.four,.kpi-stats{grid-template-columns:repeat(2,minmax(0,1fr))}.grid.rooms{grid-template-columns:repeat(2,minmax(0,1fr))}.two-col,.manager-grid{grid-template-columns:1fr}.wallboard-head{display:grid;grid-template-columns:1fr}.wallboard-tools{justify-content:flex-start}.wall-stats{grid-template-columns:repeat(3,1fr)}.wall-bottom{grid-template-columns:1fr 1fr}.wall-bottom .wall-panel:last-child{grid-column:1/-1}.wall-room-grid.view-compact{grid-template-columns:repeat(3,minmax(0,1fr))}.wall-room-grid.view-cards{grid-template-columns:repeat(2,minmax(0,1fr))}.room-card{min-height:auto}.reassign-box{grid-template-columns:1fr 1fr}.manager-export-actions{width:100%}}\n@media(max-width:767px){body{background:#f7f9fc}.topbar{gap:8px}.brand{gap:8px}.brand img{width:38px;height:38px}.brand b{font-size:14px}.brand span{font-size:10px}.userbox span{font-size:12px}.container{padding:12px}.nav a{font-size:12px}.page-head{display:grid;margin-bottom:14px}.page-head h1{font-size:24px}.page-head p{font-size:13px;margin-top:4px}.page-head>.btn,.page-head .manager-export-actions{width:100%}.page-head .manager-export-actions{display:grid;grid-template-columns:1fr 1fr}.stats,.stats.four,.kpi-stats{grid-template-columns:repeat(2,minmax(0,1fr));gap:8px;margin:12px 0 16px}.stat{border-radius:14px;padding:12px}.stat b{font-size:24px}.stat span{font-size:11px}.grid.rooms{grid-template-columns:1fr;gap:10px}.room-card{border-radius:16px;padding:14px}.room-top b{font-size:21px}.timer{font-size:32px}.actions{display:grid;grid-template-columns:repeat(2,1fr)}.actions .btn:only-child{grid-column:1/-1}.handover-box{display:grid;grid-template-columns:1fr auto;gap:7px}.login-shell,.patient-shell{min-height:auto;padding:8px}.login-card,.patient-card{width:100%;border-radius:20px;padding:20px}.hero-logo{width:82px;height:82px}.patient-card h1{font-size:34px}.reason-grid{gap:7px}.call-btn{width:180px;height:180px;font-size:38px}.call-btn span{font-size:17px}.panel{padding:14px!important;border-radius:14px}.inline-form{display:grid!important;grid-template-columns:1fr auto}.room-edit-form{grid-template-columns:1fr!important}.reassign-box{grid-template-columns:1fr!important}.reassign-title{display:block!important}.manager-filter{overflow:hidden}.period-buttons{display:grid;grid-template-columns:repeat(2,1fr)}.filter-chip{min-height:38px}.custom-dates{display:none;grid-template-columns:1fr}.custom-dates.show{display:grid}.custom-dates label,.custom-dates input{width:100%}.wallboard-shell{padding:0}.wallboard-head h1{font-size:24px}.wallboard-tools{display:grid;grid-template-columns:1fr 1fr;width:100%;gap:7px}.wall-clock,.view-switch,.lang-mini{grid-column:1/-1}.view-switch{display:grid;grid-template-columns:repeat(4,1fr);width:100%}.view-btn{padding:8px 4px;font-size:10px}.wall-stats{grid-template-columns:repeat(2,minmax(0,1fr));gap:7px}.wall-stat{padding:10px 12px}.wall-stat b{font-size:22px}.wall-filters{display:flex;overflow-x:auto;flex-wrap:nowrap;padding-bottom:5px}.wall-filters>*{flex:0 0 auto}.wall-sort{margin-left:0}.wall-updated{display:none}.wall-room-grid.view-compact,.wall-room-grid.view-cards,.wall-room-grid.view-zones{grid-template-columns:1fr 1fr!important}.wall-room-grid.view-list{display:grid!important;grid-template-columns:1fr!important}.wall-room-card{min-width:0;padding:11px!important}.wall-room-top b{font-size:18px}.room-reason{font-size:12px;min-height:28px}.room-meta{font-size:10px}.wall-live-timer{font-size:24px}.wall-bottom{grid-template-columns:1fr}.wall-bottom .wall-panel:last-child{grid-column:auto}.wall-panel{padding:12px}.workload-card{grid-template-columns:1fr 70px}.notify-mini-shell{width:calc(100% - 16px)}}\n@media(max-width:430px){.container{padding:9px}.stats,.stats.four,.kpi-stats{grid-template-columns:1fr 1fr}.wall-room-grid.view-compact,.wall-room-grid.view-cards,.wall-room-grid.view-zones{grid-template-columns:1fr!important}.wallboard-tools{grid-template-columns:1fr}.wallboard-tools>*{grid-column:1!important}.view-switch{grid-template-columns:repeat(2,1fr)}.reason-grid{grid-template-columns:1fr 1fr}.call-btn{width:160px;height:160px}.table-wrap{border-radius:12px}.notify-mini-steps{grid-template-columns:1fr 1fr}.notify-mini-chip b{font-size:7.5px}.manager-export-actions{grid-template-columns:1fr!important}}\n@media(orientation:landscape) and (max-height:600px){.topbar{position:static}.login-shell,.patient-shell{min-height:auto}.call-btn{width:145px;height:145px}.patient-card{padding:14px}.wallboard-head{margin-bottom:8px}.wall-stats{margin:8px 0}}\n"
APP_CSS += "/* Patient Arabic/mobile refinement */\n.patient-lang[dir=\"rtl\"] .patient-card{direction:rtl;text-align:center}\n.patient-lang[dir=\"rtl\"] .patient-lang-switch{direction:ltr;justify-content:flex-start}\n.patient-lang[dir=\"rtl\"] .hero-logo{width:76px;height:76px;margin:4px auto 6px}\n.patient-lang[dir=\"rtl\"] .eyebrow{font-size:12px;letter-spacing:.14em;margin-top:2px}\n.patient-lang[dir=\"rtl\"] .patient-card>h1{font-size:42px;line-height:1;margin:7px 0 18px}\n.patient-lang[dir=\"rtl\"] .patient-question{font-size:26px;line-height:1.45;margin:8px 0 8px}\n.patient-lang[dir=\"rtl\"] .patient-note{font-size:16px;line-height:1.7;margin:0 auto 14px;max-width:430px}\n.patient-lang[dir=\"rtl\"] .reason-grid{direction:rtl;margin:14px 0 20px;gap:9px}\n.patient-lang[dir=\"rtl\"] .reason{min-height:72px;display:flex;align-items:center;justify-content:center;gap:8px;font-size:17px;line-height:1.35;padding:10px}\n.patient-lang[dir=\"rtl\"] .reason-icon{font-size:25px}\n.patient-lang[dir=\"rtl\"] .call-btn{width:190px;height:190px;font-size:40px;border-width:8px}\n.patient-lang[dir=\"rtl\"] .call-btn span{font-size:18px;line-height:1.45;margin-top:4px}\n.patient-lang[dir=\"rtl\"] .physical-note{font-size:12px;line-height:1.55}\n.patient-lang-switch{display:flex;align-items:center;justify-content:flex-end;gap:7px;margin-bottom:2px}\n.lang-chip{min-width:44px;min-height:38px;padding:6px 10px;border-radius:999px;font-size:14px;font-weight:800}\n@media(max-width:600px){\n .patient-shell{padding:0!important;display:block}\n .patient-card{max-width:100%;border-radius:0;box-shadow:none;border-left:0;border-right:0;padding:14px 16px 20px}\n .patient-lang[dir=\"rtl\"] .hero-logo{width:68px;height:68px;margin:0 auto 4px}\n .patient-lang[dir=\"rtl\"] .patient-card>h1{font-size:38px;margin:5px 0 14px}\n .patient-lang[dir=\"rtl\"] .patient-question{font-size:23px;margin-top:4px}\n .patient-lang[dir=\"rtl\"] .patient-note{font-size:15px;line-height:1.6;margin-bottom:12px}\n .patient-lang[dir=\"rtl\"] .reason-grid{grid-template-columns:1fr 1fr;gap:8px;margin:12px 0 18px}\n .patient-lang[dir=\"rtl\"] .reason{min-height:66px;font-size:15px;padding:8px 7px}\n .patient-lang[dir=\"rtl\"] .reason-icon{font-size:22px}\n .patient-lang[dir=\"rtl\"] .call-btn{width:160px;height:160px;font-size:34px}\n .patient-lang[dir=\"rtl\"] .call-btn span{font-size:16px}\n .patient-lang-switch{margin-bottom:0}\n .lang-chip{min-width:42px;min-height:36px;font-size:13px}\n}\n@media(max-width:380px){\n .patient-card{padding:10px 12px 16px}\n .patient-lang[dir=\"rtl\"] .hero-logo{width:58px;height:58px}\n .patient-lang[dir=\"rtl\"] .patient-question{font-size:21px}\n .patient-lang[dir=\"rtl\"] .reason{font-size:14px;min-height:62px}\n .patient-lang[dir=\"rtl\"] .call-btn{width:148px;height:148px}\n}"
APP_CSS += "/* Premium patient call UI */\n.patient-card{position:relative;max-width:640px;margin:auto;background:#fff;border:1px solid #e4e7ec;border-radius:28px;padding:22px 24px 26px;box-shadow:0 20px 60px rgba(16,24,40,.08)}\n.patient-lang-switch{position:relative;z-index:2;margin-bottom:4px}\n.patient-brand-block{display:grid;justify-items:center;text-align:center;margin:0 auto 8px}\n.patient-main-logo{width:62px!important;height:62px!important;object-fit:contain;margin:0 auto 6px!important}\n.patient-zone{font-size:12px;letter-spacing:.18em;font-weight:900;color:#005eb8;margin:0}\n.patient-room-code{font-size:48px!important;line-height:1!important;margin:6px 0 14px!important;color:#101828}\n.patient-question{font-size:25px!important;line-height:1.35;margin:8px 0 6px!important;color:#101828}\n.patient-note{font-size:15px;line-height:1.6;max-width:480px;margin:0 auto 14px!important;color:#667085}\n.reason-grid{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:10px;margin:16px 0 22px}\n.reason{min-height:70px;border:1px solid #d0d5dd;border-radius:16px;background:#fff;display:flex;align-items:center;justify-content:center;gap:8px;padding:10px 12px;font-size:16px;font-weight:750;color:#344054;transition:.18s ease}\n.reason:hover{border-color:#84adff;background:#f8fbff}\n.reason.selected{border:2px solid #005eb8;background:#eaf4ff;box-shadow:0 4px 14px rgba(0,94,184,.08)}\n.reason-icon{font-size:25px}\n.patient-call-action{display:grid;justify-items:center;margin-top:2px}\n.call-btn-pro{position:relative;width:176px!important;height:176px!important;border-radius:50%!important;border:9px solid #d9ecff!important;background:linear-gradient(145deg,#005eb8,#003b71)!important;box-shadow:0 16px 36px rgba(0,79,158,.24),inset 0 1px 0 rgba(255,255,255,.18)!important;color:#fff;display:flex!important;flex-direction:column;align-items:center;justify-content:center;gap:3px;padding:12px;transition:transform .14s ease,box-shadow .18s ease,filter .18s ease}\n.call-btn-pro:active{transform:scale(.97)}\n.call-btn-pro:hover{filter:brightness(1.04);box-shadow:0 18px 42px rgba(0,79,158,.30),inset 0 1px 0 rgba(255,255,255,.2)!important}\n.call-btn-icon{font-size:43px;line-height:1}\n.call-btn-label{font-size:17px!important;font-weight:900;line-height:1.3;text-align:center}\n.call-btn-sub{font-size:10px;opacity:.78;line-height:1.2;text-align:center}\n.call-btn-pro.sending{animation:patientCallPulse 1.1s ease-in-out infinite}\n.call-btn-pro.sent{background:linear-gradient(145deg,#079455,#067647)!important;border-color:#d1fadf!important}\n.call-message{min-height:18px;margin-top:7px}\n@keyframes patientCallPulse{0%,100%{box-shadow:0 0 0 0 rgba(0,94,184,.18),0 16px 36px rgba(0,79,158,.20)}50%{box-shadow:0 0 0 12px rgba(0,94,184,.10),0 18px 42px rgba(0,79,158,.30)}}\n.call-active{border:1px solid #b2ddff;background:#f5fbff;border-radius:20px;padding:22px 18px;margin-top:14px;text-align:center}\n.call-active .timer{font-size:44px!important;font-variant-numeric:tabular-nums;margin:8px 0 4px}\n.call-active h2{margin:0 0 4px;font-size:22px}\n.call-active p{margin:6px 0}\n.physical-note{margin-top:20px;border-radius:12px;padding:10px 12px;background:#fffaeb;border:1px solid #fedf89;color:#7a2e0e;font-size:11px;line-height:1.5}\n.patient-lang[dir=\"rtl\"] .reason{flex-direction:row-reverse}\n.patient-lang[dir=\"rtl\"] .call-btn-label,.patient-lang[dir=\"rtl\"] .call-btn-sub{direction:rtl}\n@media(max-width:600px){\n .patient-card{padding:14px 14px 18px!important;border-radius:0!important;box-shadow:none!important;border-left:0!important;border-right:0!important}\n .patient-main-logo{width:56px!important;height:56px!important}\n .patient-room-code{font-size:42px!important;margin-bottom:12px!important}\n .patient-question{font-size:23px!important}\n .patient-note{font-size:14px!important}\n .reason-grid{gap:8px;margin:12px 0 18px}\n .reason{min-height:64px;font-size:15px;padding:8px}\n .reason-icon{font-size:22px}\n .call-btn-pro{width:162px!important;height:162px!important;border-width:8px!important}\n .call-btn-icon{font-size:38px}\n .call-btn-label{font-size:16px!important}\n .physical-note{margin-top:16px}\n}\n@media(max-width:380px){\n .patient-room-code{font-size:38px!important}\n .patient-question{font-size:21px!important}\n .reason{min-height:60px;font-size:14px}\n .call-btn-pro{width:150px!important;height:150px!important}\n .call-btn-icon{font-size:34px}\n .call-btn-label{font-size:15px!important}\n}\n@media(prefers-reduced-motion:reduce){.call-btn-pro.sending{animation:none}}"
APP_CSS += ".reason-radio{position:absolute;opacity:0;pointer-events:none}.reason-radio:checked+.reason{border:2px solid #005eb8;background:#eaf4ff;box-shadow:0 4px 14px rgba(0,94,184,.08)}.patient-lang-switch .lang-chip{text-decoration:none;display:inline-flex;align-items:center;justify-content:center;color:#344054;border:1px solid #d0d5dd;background:#fff}.patient-lang-switch .lang-chip.active{background:#eaf4ff;border-color:#84adff;color:#005eb8}.reason{cursor:pointer;-webkit-tap-highlight-color:transparent;touch-action:manipulation}.call-btn-pro{touch-action:manipulation;-webkit-tap-highlight-color:transparent}"
APP_CSS += ".push-controls{display:flex;align-items:center;gap:8px;flex-wrap:wrap}.push-health{display:inline-flex;align-items:center;min-height:32px;padding:6px 10px;border-radius:999px;font-size:11px;font-weight:850;border:1px solid #d0d5dd;background:#f8fafc;color:#475467}.push-health.ok{background:#ecfdf3;border-color:#75e0a7;color:#067647}.push-health.warn{background:#fffaeb;border-color:#fedf89;color:#b54708}.push-health.bad{background:#fef3f2;border-color:#fda29b;color:#b42318}.push-health.checking{background:#eff8ff;border-color:#b2ddff;color:#175cd3}@media(max-width:767px){.push-controls{width:100%;display:grid;grid-template-columns:1fr}.push-controls .btn{width:100%}.push-health{justify-content:center}}"
APP_CSS += ".recall-panel{margin-top:16px;padding:14px;border-radius:14px;border:1px solid #fedf89;background:#fffaeb;text-align:center}.recall-panel[hidden]{display:none!important}.recall-warning{margin:0 0 10px;color:#93370d;font-weight:750;line-height:1.45}.recall-btn{background:#f79009!important;color:#fff!important;border-color:#f79009!important;min-width:160px}.recall-btn:hover{filter:brightness(.97)}.recall-panel small{display:block;margin-top:7px;color:#667085;font-weight:700}.patient-lang[dir=\"rtl\"] .recall-warning{direction:rtl}@media(max-width:600px){.recall-panel{padding:12px}.recall-btn{width:100%}}"
APP_CSS += ".recall-form{margin:0;display:block}.recall-form .recall-btn{width:auto}@media(max-width:600px){.recall-form .recall-btn{width:100%}}"
APP_CSS += ".diag-shell{max-width:1600px;margin:0 auto}.diag-head{display:flex;justify-content:space-between;align-items:center;gap:20px;background:#fff;border:1px solid var(--line);border-radius:16px;padding:16px;margin-bottom:14px}.diag-head h1{margin:3px 0 3px;font-size:28px}.diag-head p{margin:0;color:var(--muted);font-size:12px}.diag-head-actions{display:flex;gap:8px;flex-wrap:wrap}.diag-stats{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin-bottom:14px}.diag-stat{background:#fff;border:1px solid var(--line);border-radius:14px;padding:15px;min-height:92px}.diag-stat span,.diag-stat small{display:block}.diag-stat span{font-size:11px;color:#344054}.diag-stat b{font-size:26px;display:block;margin:8px 0}.diag-stat small{font-size:10px;color:#667085}.diag-pill{display:inline-block!important;width:max-content;padding:4px 7px;border-radius:999px;font-weight:800}.diag-pill.ok{background:#dcfae6;color:#067647}.diag-pill.blue{background:#dbeafe;color:#175cd3}.diag-ready{color:#067647!important;font-size:16px!important}.diag-bad{color:#b42318!important;font-size:16px!important}.diag-panel{background:#fff;border:1px solid var(--line);border-radius:16px;padding:14px}.diag-filter{display:flex;align-items:end;gap:8px;flex-wrap:wrap}.diag-filter label{display:grid;gap:5px;font-size:10px;font-weight:700}.diag-filter input,.diag-filter select{height:42px;border:1px solid #cfd7e2;border-radius:9px;padding:8px;background:#fff;min-width:140px}.diag-filter input{min-width:260px}.diag-clear{margin:10px 0 12px;padding:8px 11px;font-size:11px}.diag-table table{min-width:1250px}.diag-table th,.diag-table td{font-size:11px;padding:10px;vertical-align:top}.diag-status{display:inline-block;padding:4px 7px;border-radius:999px;font-weight:800}.diag-status.open{background:#fee4e2;color:#b42318}.diag-status.recovered{background:#dcfae6;color:#067647}.diag-message{max-width:430px;white-space:normal}.diag-message details{margin-top:6px}.diag-message pre{white-space:pre-wrap;max-height:220px;overflow:auto;background:#101828;color:#f2f4f7;border-radius:8px;padding:8px;font-size:10px}.diag-empty{text-align:left;color:#667085}.diag-table code{font-size:10px}@media(max-width:900px){.diag-head{align-items:flex-start;display:grid}.diag-head-actions{display:grid;grid-template-columns:1fr 1fr}.diag-stats{grid-template-columns:1fr 1fr}.diag-filter{display:grid;grid-template-columns:1fr 1fr}.diag-filter label,.diag-filter input,.diag-filter select,.diag-filter .btn{width:100%;min-width:0}}@media(max-width:520px){.diag-stats{grid-template-columns:1fr}.diag-head-actions,.diag-filter{grid-template-columns:1fr}.diag-head h1{font-size:23px}}"
APP_CSS += ".live-screen-sound-card{grid-column:span 2}.live-screen-sound-form{display:grid;gap:9px;margin-top:8px}.live-sound-inputs{display:grid;grid-template-columns:1fr auto 1fr;gap:8px;align-items:end}.live-sound-inputs label{display:grid;gap:4px;font-size:11px;color:#667085}.live-sound-inputs input{width:100%;border:1px solid #d0d5dd;border-radius:9px;padding:8px;background:#fff}.live-screen-sound-form .btn{justify-self:start}@media(max-width:700px){.live-screen-sound-card{grid-column:auto}}"\nAPP_CSS += ".admin-console{display:grid;grid-template-columns:210px minmax(0,1fr);gap:20px;align-items:start}.admin-side{position:sticky;top:118px;background:#0b2a4a;color:#fff;border-radius:20px;padding:14px;min-height:620px;box-shadow:0 18px 40px rgba(11,42,74,.12)}.admin-side-brand{display:flex;align-items:center;gap:10px;padding:8px 8px 18px;border-bottom:1px solid rgba(255,255,255,.12)}.admin-side-icon{width:36px;height:36px;border-radius:10px;background:rgba(255,255,255,.12);display:grid;place-items:center;font-size:18px}.admin-side-brand b,.admin-side-brand small{display:block}.admin-side-brand small{font-size:10px;opacity:.7;margin-top:2px}.admin-side-nav{display:grid;gap:5px;margin-top:12px}.admin-nav-item{border:0;background:transparent;color:#dbe7f3;text-decoration:none;border-radius:11px;padding:10px 11px;display:flex;align-items:center;gap:9px;font:inherit;font-weight:700;font-size:12px;cursor:pointer;text-align:left;width:100%}.admin-nav-item:hover,.admin-nav-item.active{background:#fff;color:#0b2a4a}.admin-nav-item em{margin-left:auto;background:#d92d20;color:#fff;border-radius:999px;font-style:normal;font-size:9px;padding:2px 6px}.admin-main{min-width:0}.admin-console-head{display:flex;align-items:center;justify-content:space-between;gap:16px;margin-bottom:16px}.admin-console-head h1{margin:3px 0 4px;font-size:30px}.admin-console-head p{margin:0;color:var(--muted)}.admin-health{display:flex;align-items:center;gap:9px;background:#fff;border:1px solid var(--line);border-radius:14px;padding:10px 13px;min-width:190px}.admin-health b,.admin-health small{display:block}.admin-health small{font-size:10px;color:var(--muted);margin-top:2px}.health-dot{width:10px;height:10px;border-radius:50%}.health-dot.ok{background:#12b76a;box-shadow:0 0 0 5px #d1fadf}.admin-tab-pane{display:none}.admin-tab-pane.active{display:block}.admin-kpis{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin-bottom:14px}.admin-kpi{appearance:none;text-align:left;background:#fff;border:1px solid var(--line);border-radius:16px;padding:16px;color:var(--ink);cursor:pointer;min-height:108px}.admin-kpi:hover{transform:translateY(-1px);box-shadow:0 8px 24px rgba(16,24,40,.06)}.admin-kpi span,.admin-kpi small{display:block}.admin-kpi span{font-size:11px;color:#475467;font-weight:700}.admin-kpi b{display:block;font-size:30px;margin:6px 0}.admin-kpi small{font-size:10px;color:#667085}.admin-kpi.danger{border-left:4px solid var(--red)}.admin-kpi.amber{border-left:4px solid var(--amber)}.admin-overview-grid{display:grid;grid-template-columns:1fr 1fr;gap:14px}.admin-card{background:#fff;border:1px solid var(--line);border-radius:16px;padding:16px}.admin-span-2{grid-column:1/-1}.admin-card-head,.admin-pane-head{display:flex;justify-content:space-between;align-items:center;gap:12px;margin-bottom:14px}.admin-card-head h2,.admin-pane-head h2{margin:0 0 3px;font-size:18px}.admin-card-head p,.admin-pane-head p{margin:0;color:var(--muted);font-size:11px}.quick-actions{display:grid;grid-template-columns:1fr 1fr;gap:9px}.quick-action{text-decoration:none;border:1px solid #e4e7ec;background:#f8fafc;border-radius:13px;padding:12px;text-align:left;display:grid;grid-template-columns:34px 1fr;column-gap:9px;color:var(--ink);cursor:pointer}.quick-action>span{grid-row:1/3;width:34px;height:34px;border-radius:10px;background:#eaf2fb;display:grid;place-items:center;color:var(--blue);font-size:16px}.quick-action b,.quick-action small{display:block}.quick-action small{font-size:10px;color:var(--muted);margin-top:2px}.health-list{display:grid;gap:2px}.health-list>div{display:flex;justify-content:space-between;gap:16px;padding:11px 0;border-bottom:1px solid #f0f2f5;font-size:12px}.health-good{color:#067647}.health-warn{color:#b54708}.health-bad{color:#b42318}.activity-list{display:grid}.activity-row{display:grid;grid-template-columns:10px 1fr auto;gap:10px;align-items:start;padding:10px 0;border-bottom:1px solid #f0f2f5}.activity-dot{width:8px;height:8px;border-radius:50%;background:#2e90fa;margin-top:5px}.activity-row b,.activity-row small{display:block}.activity-row small{font-size:10px;color:var(--muted);margin-top:2px}.activity-row time{font-size:10px;color:var(--muted)}.admin-pane-head{background:#fff;border:1px solid var(--line);border-radius:16px;padding:15px}.admin-toolbar{display:flex;gap:8px;margin:10px 0}.admin-toolbar input{width:min(360px,100%);border:1px solid #cfd7e2;border-radius:11px;padding:11px 13px;background:#fff}.compact-table{border-radius:15px}.compact-table th,.compact-table td{padding:11px 12px;font-size:12px}.person-cell{display:flex;align-items:center;gap:9px}.person-cell>span{width:31px;height:31px;border-radius:50%;background:#eaf2fb;color:#0059ad;display:grid;place-items:center;font-weight:800}.person-cell b,.person-cell small,.room-name-cell b,.room-name-cell small{display:block}.person-cell small,.room-name-cell small{font-size:9px;color:var(--muted);margin-top:2px}.admin-actions-cell{width:50px;text-align:right}.action-menu{position:relative}.action-menu summary{list-style:none;cursor:pointer;font-size:20px;padding:2px 8px;border-radius:8px}.action-menu summary::-webkit-details-marker{display:none}.action-menu[open] summary{background:#eef2f6}.action-menu>div{position:absolute;right:0;top:34px;z-index:20;background:#fff;border:1px solid var(--line);box-shadow:0 12px 30px rgba(16,24,40,.14);border-radius:12px;padding:6px;min-width:170px;display:grid}.action-menu button,.action-menu a{border:0;background:transparent;text-decoration:none;color:#344054;text-align:left;padding:9px 10px;border-radius:8px;font:inherit;font-size:11px;cursor:pointer;width:100%}.action-menu button:hover,.action-menu a:hover{background:#f2f4f7}.action-menu form{margin:0}.action-menu .danger-link{color:#b42318}.inline-form.compact select{padding:8px 9px}.inline-form.compact .btn{padding:8px 10px}.qr-small{padding:8px 10px;font-size:11px}.device-mode{display:inline-block;padding:5px 7px;border-radius:7px;background:#eef4ff;color:#1849a9;font-size:10px;font-weight:800}.notification-kpis .admin-kpi{cursor:default}.system-grid.modern{grid-template-columns:repeat(3,1fr)}.system-grid.modern>div{background:#fff}.system-grid.modern small{display:block;color:var(--muted);font-size:10px;margin-top:4px}.admin-modal{position:fixed;inset:0;background:rgba(16,24,40,.48);display:none;place-items:center;padding:18px;z-index:100}.admin-modal.open{display:grid}.admin-modal-card{position:relative;width:min(440px,100%);background:#fff;border-radius:20px;padding:22px;box-shadow:0 30px 80px rgba(16,24,40,.28)}.admin-modal-card h2{margin:4px 0 4px}.admin-modal-card>p{margin:0 0 14px;color:var(--muted);font-size:12px}.modal-close{position:absolute;right:12px;top:10px;border:0;background:#f2f4f7;border-radius:50%;width:32px;height:32px;font-size:20px;cursor:pointer}.modal-form{display:grid;gap:11px}.modal-form label{display:grid;gap:5px;font-size:11px;font-weight:700}.modal-form input,.modal-form select{border:1px solid #cfd7e2;border-radius:10px;padding:11px;background:#fff}.empty-state{padding:20px;color:var(--muted);text-align:center}@media(max-width:1050px){.admin-console{grid-template-columns:170px minmax(0,1fr)}.admin-kpis{grid-template-columns:1fr 1fr}.system-grid.modern{grid-template-columns:1fr 1fr}}@media(max-width:800px){.admin-console{display:block}.admin-side{position:static;min-height:0;border-radius:16px;padding:8px;margin-bottom:14px}.admin-side-brand{display:none}.admin-side-nav{display:flex;overflow-x:auto;margin:0;gap:6px}.admin-nav-item{width:auto;white-space:nowrap;flex:0 0 auto}.admin-console-head{align-items:flex-start}.admin-overview-grid{grid-template-columns:1fr}.admin-span-2{grid-column:auto}.admin-health{display:none}.compact-table table{min-width:760px}}@media(max-width:560px){.admin-kpis{grid-template-columns:1fr 1fr;gap:8px}.admin-kpi{min-height:92px;padding:12px}.admin-kpi b{font-size:25px}.quick-actions{grid-template-columns:1fr}.admin-pane-head{align-items:flex-start;flex-direction:column}.admin-pane-head .btn{width:100%}.system-grid.modern{grid-template-columns:1fr}.admin-console-head h1{font-size:25px}}"
APP_CSS += "/* Dedicated fullscreen wallboard page */\nbody.wallboard-page{margin:0;overflow-x:hidden;background:#f3f6fa}\nbody.wallboard-page .topbar,\nbody.wallboard-page .nav{display:none!important}\nbody.wallboard-page .container{\n  max-width:none!important;\n  width:100%!important;\n  margin:0!important;\n  padding:14px 18px 18px!important;\n}\nbody.wallboard-page .wallboard-shell{\n  width:100%!important;\n  max-width:none!important;\n  min-height:calc(100vh - 32px);\n  margin:0!important;\n}\nbody.wallboard-page .wallboard-head{\n  margin-top:0!important;\n}\nbody.wallboard-page .wallboard-head h1{\n  font-size:clamp(30px,2.5vw,46px);\n}\nbody.wallboard-page .wallboard-head p{\n  margin-bottom:0;\n}\nbody.wallboard-page .wall-stats{\n  grid-template-columns:repeat(6,minmax(120px,1fr));\n}\nbody.wallboard-page .wall-room-grid{\n  width:100%;\n}\nbody.wallboard-page .wall-room-card{\n  width:100%;\n}\nbody.wallboard-page .wall-bottom{\n  width:100%;\n}\n@media(min-width:1400px){\n  body.wallboard-page .container{padding:16px 28px 22px!important}\n  body.wallboard-page .wall-room-card{min-height:72px}\n}\n@media(max-width:900px){\n  body.wallboard-page .container{padding:10px!important}\n  body.wallboard-page .wall-stats{grid-template-columns:repeat(3,1fr)}\n  body.wallboard-page .wallboard-head{align-items:flex-start}\n}\n@media(max-width:600px){\n  body.wallboard-page .wall-stats{grid-template-columns:repeat(2,1fr)}\n}"
APP_CSS += "/* Fullscreen Charge Nurse Live Board */\nbody.charge-board-page{margin:0;background:#f3f6fa;overflow-x:hidden}\nbody.charge-board-page .topbar,\nbody.charge-board-page .nav{display:none!important}\nbody.charge-board-page .container{\n  max-width:none!important;\n  width:100%!important;\n  margin:0!important;\n  padding:14px 18px 20px!important;\n}\n.charge-board-root{width:100%;min-height:calc(100vh - 34px)}\n.charge-board-head{display:flex;align-items:flex-start;justify-content:space-between;gap:18px;margin-bottom:14px}\n.charge-board-head h1{margin:3px 0 4px;font-size:clamp(28px,2.5vw,44px)}\n.charge-board-head p{margin:0;color:var(--muted)}\n.charge-board-tools{display:flex;align-items:center;gap:8px;flex-wrap:wrap;justify-content:flex-end}\n.charge-view-switch{display:flex;gap:4px;padding:4px;background:#e9eef5;border-radius:12px}\n.charge-view-btn{border:0;background:transparent;color:#344054;border-radius:9px;padding:9px 12px;font-weight:800;cursor:pointer}\n.charge-view-btn.active{background:#fff;color:var(--blue);box-shadow:0 2px 8px rgba(16,24,40,.08)}\n.charge-stats{display:grid;grid-template-columns:repeat(3,minmax(160px,1fr));gap:12px;margin-bottom:14px}\n.charge-stat{background:#fff;border:1px solid var(--line);border-radius:16px;padding:14px 16px}\n.charge-stat span{display:block;color:var(--muted);font-size:11px}\n.charge-stat b{display:block;font-size:30px;margin-top:5px}\n.charge-stat.danger{border-left:4px solid var(--red)}\n.charge-stat.amber{border-left:4px solid var(--amber)}\n.charge-room-grid{transition:.2s ease}\n.charge-room-card{min-height:0!important}\n.charge-call-summary{display:flex;justify-content:space-between;align-items:flex-end;gap:14px}\n.charge-call-reason{font-weight:750;color:#344054;padding-bottom:8px}\n.charge-ready-row{display:grid;gap:10px}\n.charge-ready-row .ready{margin-top:18px}\n.charge-ready-row select{width:100%;border:1px solid #cfd7e2;border-radius:10px;padding:9px;background:#fff}\n\n/* Compact: many rooms visible at once */\n.charge-board-root.view-compact .charge-room-grid{\n  grid-template-columns:repeat(auto-fit,minmax(220px,1fr));\n  gap:10px;\n}\n.charge-board-root.view-compact .charge-room-card{padding:13px;border-radius:15px}\n.charge-board-root.view-compact .room-top b{font-size:20px}\n.charge-board-root.view-compact .timer{font-size:28px}\n.charge-board-root.view-compact .reassign-box{padding-top:10px;margin-top:10px}\n.charge-board-root.view-compact .reassign-box select{font-size:11px;padding:7px}\n.charge-board-root.view-compact .reassign-box .btn{padding:8px 10px;font-size:11px}\n\n/* Cards: larger operational cards */\n.charge-board-root.view-cards .charge-room-grid{\n  grid-template-columns:repeat(auto-fit,minmax(310px,1fr));\n  gap:14px;\n}\n.charge-board-root.view-cards .charge-room-card{padding:18px;border-radius:20px}\n.charge-board-root.view-cards .room-top b{font-size:26px}\n.charge-board-root.view-cards .timer{font-size:38px}\n\n/* List: one room per row */\n.charge-board-root.view-list .charge-room-grid{display:grid;grid-template-columns:1fr;gap:8px}\n.charge-board-root.view-list .charge-room-card{padding:11px 14px;border-radius:14px}\n.charge-board-root.view-list .charge-room-main{\n  display:grid;\n  grid-template-columns:190px minmax(260px,1fr) auto;\n  column-gap:16px;\n  align-items:center;\n}\n.charge-board-root.view-list .room-top{grid-column:1;grid-row:1}\n.charge-board-root.view-list .assignment{grid-column:1;grid-row:2;margin:2px 0 0}\n.charge-board-root.view-list .charge-call-summary{grid-column:2;grid-row:1 / span 2;align-items:center}\n.charge-board-root.view-list .charge-call-summary .timer{font-size:30px;margin:3px 0}\n.charge-board-root.view-list .charge-call-reason{padding:0;font-size:15px}\n.charge-board-root.view-list .alert{grid-column:2;grid-row:3;margin:5px 0 0}\n.charge-board-root.view-list .actions{grid-column:3;grid-row:1;margin:0;justify-content:flex-end}\n.charge-board-root.view-list .reassign-box{grid-column:3;grid-row:2 / span 2;margin:0;padding:0;border:0;display:grid;grid-template-columns:1fr 1fr auto;gap:6px;min-width:390px}\n.charge-board-root.view-list .reassign-title{grid-column:1/-1;margin:0}\n.charge-board-root.view-list .charge-ready-row{grid-column:2 / span 2;grid-row:1 / span 2;display:flex;align-items:center;justify-content:space-between;gap:16px}\n.charge-board-root.view-list .charge-ready-row .ready{margin:0}\n.charge-board-root.view-list .charge-ready-row select{width:240px}\n\n@media(min-width:1500px){\n  body.charge-board-page .container{padding:16px 28px 22px!important}\n  .charge-board-root.view-compact .charge-room-grid{grid-template-columns:repeat(auto-fit,minmax(240px,1fr))}\n}\n@media(max-width:1050px){\n  .charge-board-root.view-list .charge-room-main{grid-template-columns:160px 1fr}\n  .charge-board-root.view-list .actions,\n  .charge-board-root.view-list .reassign-box,\n  .charge-board-root.view-list .charge-ready-row{grid-column:2;min-width:0}\n  .charge-board-root.view-list .actions{grid-row:2}\n  .charge-board-root.view-list .reassign-box{grid-row:3}\n  .charge-board-root.view-list .charge-ready-row{grid-row:1 / span 2}\n}\n@media(max-width:760px){\n  body.charge-board-page .container{padding:10px!important}\n  .charge-board-head{display:grid}\n  .charge-board-tools{justify-content:flex-start}\n  .charge-stats{grid-template-columns:repeat(3,1fr)}\n  .charge-board-root.view-list .charge-room-main{display:block}\n  .charge-board-root.view-list .reassign-box{display:grid;grid-template-columns:1fr}\n  .charge-board-root.view-list .reassign-title{grid-column:auto}\n  .charge-board-root.view-list .charge-ready-row{display:grid}\n}\n@media(max-width:520px){\n  .charge-stats{grid-template-columns:1fr 1fr}\n  .charge-stats .charge-stat:last-child{grid-column:1/-1}\n  .charge-view-switch{width:100%;display:grid;grid-template-columns:repeat(3,1fr)}\n  .charge-view-btn{padding:8px 5px;font-size:11px}\n  .charge-board-tools>.btn{flex:1}\n}"
APP_CSS += ".permission-note{font-size:11px;color:#475467;background:#f2f4f7;border-radius:999px;padding:7px 10px}.profile-legend{display:flex;gap:6px;flex-wrap:wrap;margin:10px 0 12px}.profile-legend span{background:#eef4ff;color:#1849a9;border-radius:999px;padding:6px 9px;font-size:10px}.permission-user-list{display:grid;gap:12px}.permission-user-card{background:#fff;border:1px solid var(--line);border-radius:16px;padding:15px}.permission-user-head{display:flex;align-items:center;justify-content:space-between;gap:14px;padding-bottom:12px;border-bottom:1px solid #eef1f4}.permission-user-head>label{display:grid;gap:5px;font-size:10px;font-weight:800;min-width:220px}.permission-user-head select{border:1px solid #cfd7e2;border-radius:10px;padding:9px;background:#fff}.permission-grid{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:8px;margin:12px 0}.permission-item{display:flex;align-items:flex-start;gap:8px;padding:10px;border:1px solid #e4e7ec;border-radius:11px;background:#fafbfc;cursor:pointer}.permission-item:has(input:checked){border-color:#84adff;background:#eff8ff}.permission-item input{margin-top:3px}.permission-item b,.permission-item small{display:block}.permission-item b{font-size:11px}.permission-item small{font-size:9px;color:#667085;margin-top:2px;line-height:1.35}.permission-actions{display:flex;gap:8px;justify-content:flex-end}.admin-side-nav .admin-nav-item span{overflow:hidden;text-overflow:ellipsis}@media(max-width:1050px){.permission-grid{grid-template-columns:repeat(2,1fr)}}@media(max-width:700px){.permission-user-head{display:grid}.permission-user-head>label{min-width:0}.permission-grid{grid-template-columns:1fr}.permission-actions{display:grid;grid-template-columns:1fr}.permission-actions .btn{width:100%}}"
APP_CSS += ".nav-alert-wrap{position:relative;margin-left:auto}.nav-alert-btn{border:0;background:#1570ef;color:#fff;border-radius:999px;padding:7px 10px;font-weight:850;font-size:11px;display:flex;align-items:center;gap:5px;cursor:pointer;white-space:nowrap}.nav-alert-btn b{background:#fff;color:#175cd3;border-radius:999px;min-width:18px;height:18px;padding:0 5px;display:grid;place-items:center;font-size:9px}.alerts-dropdown{position:absolute;right:0;top:calc(100% + 8px);z-index:90;width:min(390px,calc(100vw - 24px));background:#fff;border:1px solid #d0d5dd;border-radius:15px;box-shadow:0 18px 45px rgba(16,24,40,.18);overflow:hidden}.alerts-drop-head{display:flex;align-items:center;justify-content:space-between;gap:10px;padding:12px 13px;border-bottom:1px solid #eaecf0}.alerts-drop-head b,.alerts-drop-head small{display:block}.alerts-drop-head small{font-size:9px;color:#667085;margin-top:2px}.alerts-drop-head button{border:0;background:transparent;color:#175cd3;font-size:10px;font-weight:800;cursor:pointer}.alerts-drop-list{max-height:410px;overflow:auto}.alerts-drop-item{display:grid;grid-template-columns:28px 1fr;gap:8px;padding:10px 12px;text-decoration:none;color:#101828;border-bottom:1px solid #f2f4f7}.alerts-drop-item:hover{background:#f8fafc}.alerts-drop-item.unread{background:#eff8ff}.alert-mini-icon{width:28px;height:28px;border-radius:50%;background:#fff4e5;display:grid;place-items:center}.alerts-drop-item b,.alerts-drop-item small,.alerts-drop-item em{display:block}.alerts-drop-item b{font-size:11px}.alerts-drop-item small{font-size:9px;color:#475467;line-height:1.35;margin-top:2px}.alerts-drop-item em{font-size:8px;color:#98a2b3;font-style:normal;margin-top:3px}.alerts-view-all{display:block;text-align:center;padding:10px;color:#175cd3;text-decoration:none;font-size:10px;font-weight:850;background:#f8fafc}.alerts-empty{padding:24px;text-align:center;color:#667085;font-size:11px}.alerts-page{max-width:1100px;margin:0 auto}.alerts-page-head{display:flex;justify-content:space-between;align-items:center;gap:12px;margin-bottom:12px}.alerts-page-head h1{margin:4px 0}.alerts-page-head p{margin:0;color:#667085}.alerts-page-summary{display:flex;align-items:baseline;gap:8px;background:#fff;border:1px solid var(--line);border-radius:14px;padding:12px 15px;margin-bottom:12px}.alerts-page-summary b{font-size:28px}.alerts-page-summary span{color:#667085;font-size:11px}.alerts-history{display:grid;gap:8px}.alert-history-row{display:grid;grid-template-columns:38px 1fr;gap:10px;background:#fff;border:1px solid var(--line);border-radius:14px;padding:13px;text-decoration:none;color:#101828}.alert-history-row.unread{border-left:4px solid #1570ef;background:#f8fbff}.alert-kind-icon{width:38px;height:38px;border-radius:50%;background:#fff4e5;display:grid;place-items:center}.alert-history-title{display:flex;align-items:center;gap:7px}.alert-history-title span{background:#dbeafe;color:#175cd3;border-radius:999px;padding:2px 6px;font-size:8px;font-weight:900}.alert-history-row p{margin:4px 0;font-size:11px;color:#475467}.alert-history-row small{color:#98a2b3;font-size:9px}.alerts-empty-page{background:#fff;border:1px solid var(--line);border-radius:14px;padding:40px;text-align:center;color:#667085}@media(max-width:700px){.nav{align-items:center}.nav-alert-wrap{margin-left:0}.alerts-dropdown{position:fixed;left:12px;right:12px;top:auto;width:auto}.alerts-page-head{display:grid}.alerts-page-head .btn{width:100%}}"
APP_CSS += ".notification-active{background:#ecfdf3!important;border-color:#75e0a7!important;color:#067647!important}.notification-verify{background:#fffaeb!important;border-color:#fedf89!important;color:#b54708!important}.notification-blocked{background:#fef3f2!important;border-color:#fda29b!important;color:#b42318!important}.notification-checking{background:#eff8ff!important;border-color:#b2ddff!important;color:#175cd3!important;cursor:wait!important}.notification-active:hover{background:#dcfae6!important}.notification-blocked:hover{background:#fee4e2!important}"
APP_CSS += "/* 2026 operations UI refresh */\n.charge-command{max-width:1500px;margin:0 auto}.charge-command-head,.management-head{display:flex;align-items:flex-start;justify-content:space-between;gap:18px;margin-bottom:14px}.command-title-row{display:flex;align-items:center;gap:10px}.command-title-row h1,.management-head h1{margin:3px 0 4px;font-size:clamp(30px,3vw,44px)}.live-pill{display:inline-flex;align-items:center;gap:5px;background:#e8f8ef;color:#079455;border-radius:999px;padding:6px 10px;font-size:11px;font-weight:900}.command-actions{display:flex;gap:8px;flex-wrap:wrap}.charge-kpis,.management-primary-kpis{display:grid;grid-template-columns:repeat(3,1fr);gap:12px;margin-bottom:14px}.command-kpi,.mg-kpi{background:#fff;border:1px solid var(--line);border-radius:17px;padding:15px;display:flex;align-items:center;gap:12px;min-height:90px}.command-kpi>span,.mg-kpi>span{font-size:28px}.command-kpi b,.mg-kpi b{font-size:29px;display:block}.command-kpi small,.mg-kpi small{color:var(--muted);font-size:11px}.command-kpi.danger,.mg-kpi.danger{border-left:4px solid #e5242f}.command-kpi.amber{border-left:4px solid #f79009}.mg-kpi.success{border-left:4px solid #12b76a}.charge-board-controls{display:flex;align-items:center;justify-content:space-between;gap:12px;margin:12px 0}.charge-filter-group,.charge-view-switch{display:flex;gap:7px;align-items:center}.charge-filter,.charge-view-btn{border:1px solid #d0d5dd;background:#fff;color:#175cd3;border-radius:12px;padding:9px 14px;font-weight:850;cursor:pointer}.charge-filter.active,.charge-view-btn.active{background:#0967c8;color:#fff;border-color:#0967c8}.charge-room-grid{display:grid;gap:12px}.charge-room-grid.view-cards{grid-template-columns:repeat(auto-fit,minmax(300px,1fr))}.charge-room-grid.view-compact{grid-template-columns:repeat(auto-fit,minmax(230px,1fr))}.charge-room-grid.view-list{grid-template-columns:1fr}.command-room-card{background:#fff;border:1px solid var(--line);border-left:4px solid #12b76a;border-radius:17px;padding:15px;box-shadow:0 4px 14px rgba(16,24,40,.03)}.command-room-card.has-call{border-left-color:#e5242f}.command-room-card.status-escalated{background:linear-gradient(90deg,#fff5f5,#fff);border-color:#f97066;border-left:5px solid #e5242f}.command-room-top{display:flex;justify-content:space-between;gap:12px;align-items:center}.command-room-top>div{display:flex;align-items:center;gap:10px}.command-room-top b{font-size:22px}.command-room-top>span{font-size:10px;color:var(--muted)}.room-state{font-size:10px;font-weight:900}.state-new,.state-escalated{color:#d92d20}.state-ready{color:#079455}.command-call-body{display:grid;gap:11px;margin-top:12px}.command-reason{font-size:13px;color:#475467}.command-call-meta{display:grid;grid-template-columns:1fr 1fr;gap:10px;align-items:end}.command-call-meta .timer{font-size:34px;color:#d92d20;font-weight:900}.command-call-meta small,.assigned-nurse small{display:block;font-size:9px;color:var(--muted);margin-bottom:3px}.assigned-nurse b{font-size:12px}.command-escalation{background:#fff1f0;color:#b42318;border-radius:10px;padding:8px;font-size:10px}.command-card-actions{display:flex;gap:7px;flex-wrap:wrap}.command-card-actions .btn{flex:1;min-width:110px}.ready-room-body{display:grid;gap:10px;margin-top:16px}.ready-room-body p{margin:0;color:#079455;font-weight:800}.command-modal label{display:grid;gap:5px;font-size:11px;font-weight:800;margin-top:10px}.command-modal select{border:1px solid #cfd7e2;border-radius:10px;padding:11px;background:#fff}.charge-room-grid.view-list .command-room-card{padding:11px 14px}.charge-room-grid.view-list .command-call-body{grid-template-columns:1fr 1.2fr auto;align-items:center}.charge-room-grid.view-list .command-card-actions{justify-content:flex-end}.charge-room-grid.view-list .command-escalation{grid-column:1/-1}\n\nbody.wallboard-page .wall-stats{grid-template-columns:repeat(3,1fr)!important}.wall-hidden-stat{display:none!important}.wall-stat small{display:block;font-size:9px;color:var(--muted);margin-top:3px}.wall-room-card{border-radius:17px!important}.wall-room-card.status-escalated,.wall-room-card.alert-critical{background:linear-gradient(90deg,#fff5f5,#fff)!important}.wall-room-top b{font-size:22px!important}.wall-live-timer{font-size:26px!important}.wall-filters{background:#fff;border:1px solid var(--line);border-radius:14px;padding:9px 10px}.wall-bottom{margin-top:12px}\n\n.management-console{max-width:1300px;margin:0 auto}.management-period{margin:10px 0 14px}.management-period .period-buttons{display:flex;gap:7px;flex-wrap:wrap}.export-menu{position:relative}.export-menu summary{list-style:none;border:1px solid #b2c8df;background:#fff;color:#0967c8;border-radius:13px;padding:10px 14px;font-weight:900;cursor:pointer}.export-menu summary::-webkit-details-marker{display:none}.export-menu>div{position:absolute;right:0;top:46px;z-index:20;background:#fff;border:1px solid var(--line);border-radius:12px;box-shadow:0 12px 28px rgba(16,24,40,.12);padding:6px;min-width:150px}.export-menu a{display:block;text-decoration:none;color:#344054;padding:9px 10px;border-radius:8px;font-size:11px}.export-menu a:hover{background:#f2f4f7}.management-primary-kpis{grid-template-columns:repeat(4,1fr)}.management-delay-strip{display:flex;gap:24px;align-items:center;flex-wrap:wrap;background:#edf6ff;border-radius:13px;padding:10px 14px;margin-bottom:14px;font-size:11px}.management-delay-strip span{display:flex;align-items:center;gap:5px}.modern-manager-grid{display:grid;grid-template-columns:1fr 1fr;gap:14px}.performance-panel{border-radius:17px}.performance-panel table{font-size:11px}.perf-value{display:flex;align-items:center;gap:9px}.perf-bar{display:block;width:110px;height:8px;background:#edf1f5;border-radius:999px;overflow:hidden}.perf-bar i{display:block;height:100%;background:linear-gradient(90deg,#2e90fa,#1570ef);border-radius:999px}.lifecycle-details{margin-top:14px}.lifecycle-details>summary{list-style:none;cursor:pointer;display:flex;justify-content:space-between;align-items:center;gap:12px}.lifecycle-details>summary::-webkit-details-marker{display:none}.lifecycle-details>summary h2{margin:0}.lifecycle-details>summary p{margin:3px 0 0;color:var(--muted);font-size:10px}.lifecycle-details[open]>summary{margin-bottom:12px}\n\n@media(max-width:900px){.charge-command-head,.management-head{display:grid}.command-actions{justify-content:flex-start}.charge-kpis{grid-template-columns:repeat(3,1fr)}.charge-board-controls{align-items:flex-start;flex-direction:column}.management-primary-kpis{grid-template-columns:1fr 1fr}.modern-manager-grid{grid-template-columns:1fr}.wallboard-tools{justify-content:flex-start}.charge-room-grid.view-list .command-call-body{grid-template-columns:1fr}}\n@media(max-width:560px){.charge-kpis{grid-template-columns:1fr 1fr}.charge-kpis .command-kpi:last-child{grid-column:1/-1}.charge-filter-group,.charge-view-switch{width:100%;overflow-x:auto}.charge-filter,.charge-view-btn{white-space:nowrap}.management-primary-kpis{grid-template-columns:1fr 1fr}.command-call-meta{grid-template-columns:1fr 1fr}.management-delay-strip{gap:10px}.perf-bar{width:70px}.wallboard-tools{gap:6px}.wall-clock{font-size:20px!important}}"
APP_CSS += "/* Restore wallboard full-card pulse after UI refresh.\n   Use an overlay so later background !important rules cannot suppress the animation. */\n.wall-room-card{position:relative;isolation:isolate;overflow:visible}\n.wall-room-card::after{\n  content:\"\";position:absolute;inset:-3px;border-radius:inherit;pointer-events:none;\n  opacity:0;z-index:2\n}\n.wall-room-card.alert-fresh::after{\n  opacity:1;border:2px solid #12b76a;\n  animation:wallPulseFreshFix 1.55s ease-in-out infinite\n}\n.wall-room-card.alert-warning::after{\n  opacity:1;border:3px solid #f79009;\n  animation:wallPulseWarningFix 1.2s ease-in-out infinite\n}\n.wall-room-card.alert-critical::after,\n.wall-room-card.status-escalated::after{\n  opacity:1;border:3px solid #d92d20;\n  animation:wallPulseCriticalFix .9s ease-in-out infinite\n}\n.wall-room-card.alert-takeover::after{\n  opacity:1;border:3px solid #2970ff;\n  animation:wallPulseTakeoverFix 1.2s ease-in-out infinite\n}\n.wall-room-card.alert-arrived::after,\n.wall-room-card.alert-idle::after{display:none}\n@keyframes wallPulseFreshFix{\n  0%,100%{box-shadow:0 0 0 0 rgba(18,183,106,.08),0 0 10px rgba(18,183,106,.06);opacity:.55}\n  50%{box-shadow:0 0 0 8px rgba(18,183,106,.18),0 0 28px rgba(18,183,106,.30);opacity:1}\n}\n@keyframes wallPulseWarningFix{\n  0%,100%{box-shadow:0 0 0 0 rgba(247,144,9,.10),0 0 12px rgba(247,144,9,.10);opacity:.65}\n  50%{box-shadow:0 0 0 10px rgba(247,144,9,.22),0 0 34px rgba(247,144,9,.38);opacity:1}\n}\n@keyframes wallPulseCriticalFix{\n  0%,100%{box-shadow:0 0 0 0 rgba(217,45,32,.12),0 0 14px rgba(217,45,32,.14);opacity:.68}\n  50%{box-shadow:0 0 0 12px rgba(217,45,32,.28),0 0 42px rgba(217,45,32,.52);opacity:1}\n}\n@keyframes wallPulseTakeoverFix{\n  0%,100%{box-shadow:0 0 0 0 rgba(41,112,255,.10),0 0 12px rgba(41,112,255,.10);opacity:.65}\n  50%{box-shadow:0 0 0 9px rgba(41,112,255,.22),0 0 32px rgba(41,112,255,.38);opacity:1}\n}\n@media (prefers-reduced-motion:reduce){\n  .wall-room-card.alert-fresh::after,\n  .wall-room-card.alert-warning::after,\n  .wall-room-card.alert-critical::after,\n  .wall-room-card.status-escalated::after,\n  .wall-room-card.alert-takeover::after{animation-duration:2.2s}\n}"
APP_CSS += "/* Definitive wallboard pulse: visible over the whole card */\n.wall-room-card{position:relative!important;isolation:isolate!important;overflow:visible!important}\n.wall-room-card>*{position:relative;z-index:2}\n.wall-room-card::before{\n  content:\"\";position:absolute;inset:0;border-radius:inherit;pointer-events:none;z-index:1;\n  opacity:0;background:transparent\n}\n.wall-room-card.alert-fresh::before{\n  opacity:1;animation:wbFreshFill 1.55s ease-in-out infinite\n}\n.wall-room-card.alert-warning::before{\n  opacity:1;animation:wbWarningFill 1.2s ease-in-out infinite\n}\n.wall-room-card.alert-critical::before,\n.wall-room-card.status-escalated::before{\n  opacity:1;animation:wbCriticalFill .9s ease-in-out infinite\n}\n.wall-room-card.alert-takeover::before{\n  opacity:1;animation:wbTakeoverFill 1.2s ease-in-out infinite\n}\n.wall-room-card.alert-arrived::before,\n.wall-room-card.alert-idle::before,\n.wall-room-card.status-ready::before,\n.wall-room-card.status-closed::before{display:none!important}\n@keyframes wbFreshFill{\n  0%,100%{background:rgba(18,183,106,.07);box-shadow:inset 0 0 0 2px #12b76a,0 0 0 0 rgba(18,183,106,.10),0 0 12px rgba(18,183,106,.08)}\n  50%{background:rgba(18,183,106,.28);box-shadow:inset 0 0 0 4px #12b76a,0 0 0 8px rgba(18,183,106,.18),0 0 32px rgba(18,183,106,.42)}\n}\n@keyframes wbWarningFill{\n  0%,100%{background:rgba(247,144,9,.08);box-shadow:inset 0 0 0 3px #f79009,0 0 0 0 rgba(247,144,9,.10),0 0 14px rgba(247,144,9,.10)}\n  50%{background:rgba(247,144,9,.30);box-shadow:inset 0 0 0 4px #f79009,0 0 0 10px rgba(247,144,9,.22),0 0 36px rgba(247,144,9,.48)}\n}\n@keyframes wbCriticalFill{\n  0%,100%{background:rgba(217,45,32,.09);box-shadow:inset 0 0 0 3px #d92d20,0 0 0 0 rgba(217,45,32,.12),0 0 16px rgba(217,45,32,.14)}\n  50%{background:rgba(217,45,32,.34);box-shadow:inset 0 0 0 5px #d92d20,0 0 0 12px rgba(217,45,32,.28),0 0 44px rgba(217,45,32,.58)}\n}\n@keyframes wbTakeoverFill{\n  0%,100%{background:rgba(41,112,255,.08);box-shadow:inset 0 0 0 3px #2970ff,0 0 0 0 rgba(41,112,255,.10),0 0 14px rgba(41,112,255,.10)}\n  50%{background:rgba(41,112,255,.28);box-shadow:inset 0 0 0 4px #2970ff,0 0 0 9px rgba(41,112,255,.22),0 0 34px rgba(41,112,255,.46)}\n}\n@media (prefers-reduced-motion:reduce){\n  .wall-room-card.alert-fresh::before,\n  .wall-room-card.alert-warning::before,\n  .wall-room-card.alert-critical::before,\n  .wall-room-card.status-escalated::before,\n  .wall-room-card.alert-takeover::before{animation-duration:2.4s!important}\n}"
APP_CSS += "/* Pulse visibility hardening: animation is intentionally obvious across the full card */\n.wall-room-card.alert-fresh::before{animation:wbFreshFill .95s ease-in-out infinite!important}\n.wall-room-card.alert-warning::before{animation:wbWarningFill .85s ease-in-out infinite!important}\n.wall-room-card.alert-critical::before,\n.wall-room-card.status-escalated::before{animation:wbCriticalFill .72s ease-in-out infinite!important}\n.wall-room-card.alert-takeover::before{animation:wbTakeoverFill .9s ease-in-out infinite!important}\n@keyframes wbFreshFill{\n  0%,100%{background:rgba(18,183,106,.06);box-shadow:inset 0 0 0 3px #12b76a,0 0 0 0 rgba(18,183,106,.10)}\n  50%{background:rgba(18,183,106,.42);box-shadow:inset 0 0 0 5px #12b76a,0 0 0 8px rgba(18,183,106,.24),0 0 34px rgba(18,183,106,.52)}\n}\n@keyframes wbWarningFill{\n  0%,100%{background:rgba(247,144,9,.07);box-shadow:inset 0 0 0 3px #f79009,0 0 0 0 rgba(247,144,9,.10)}\n  50%{background:rgba(247,144,9,.45);box-shadow:inset 0 0 0 5px #f79009,0 0 0 10px rgba(247,144,9,.26),0 0 38px rgba(247,144,9,.58)}\n}\n@keyframes wbCriticalFill{\n  0%,100%{background:rgba(217,45,32,.08);box-shadow:inset 0 0 0 3px #d92d20,0 0 0 0 rgba(217,45,32,.12)}\n  50%{background:rgba(217,45,32,.48);box-shadow:inset 0 0 0 6px #d92d20,0 0 0 12px rgba(217,45,32,.30),0 0 46px rgba(217,45,32,.66)}\n}\n@keyframes wbTakeoverFill{\n  0%,100%{background:rgba(41,112,255,.07);box-shadow:inset 0 0 0 3px #2970ff,0 0 0 0 rgba(41,112,255,.10)}\n  50%{background:rgba(41,112,255,.42);box-shadow:inset 0 0 0 5px #2970ff,0 0 0 9px rgba(41,112,255,.25),0 0 36px rgba(41,112,255,.56)}\n}"
APP_JS = "function pad(v){return String(v).padStart(2,'0')}\nfunction updateTimers(){document.querySelectorAll('.timer[data-created]').forEach(el=>{let sec=0;if(el.dataset.elapsed!==undefined&&el.dataset.elapsed!==''){const base=Math.max(0,Number(el.dataset.elapsed)||0);if(el.dataset.stop){sec=base}else{if(!el.dataset.clientAnchor)el.dataset.clientAnchor=String(Date.now());sec=base+Math.max(0,Math.floor((Date.now()-Number(el.dataset.clientAnchor))/1000))}}else{const s=new Date(el.dataset.created);const stop=el.dataset.stop?new Date(el.dataset.stop).getTime():Date.now();sec=Math.max(0,Math.floor((stop-s.getTime())/1000))}el.textContent=`${pad(Math.floor(sec/60))}:${pad(sec%60)}`})}\nsetInterval(updateTimers,1000);updateTimers();\nasync function callAction(id,action){const r=await fetch(`/api/call/${id}/${action}`,{method:'POST'});if(r.ok){const d=await r.json().catch(()=>({}));const stop=d.call&&(d.call.arrived_at||d.call.resolved_at||d.call.stop_at);if(stop){document.querySelectorAll('.timer[data-call-id=\"'+id+'\"]').forEach(el=>{el.dataset.stop=stop});updateTimers()}if(action==='arrive'){setTimeout(()=>location.reload(),150)}else location.reload()}else alert('Action could not be completed.');}\nasync function assignRoom(id,nurseId){if(!nurseId)return;const r=await fetch(`/api/room/${id}/assign`,{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({nurse_id:nurseId})});if(r.ok)location.reload();}\nasync function reassignCall(id){const n=document.getElementById(`reassign-nurse-${id}`),rs=document.getElementById(`reassign-reason-${id}`);if(!n||!n.value){alert('Select a nurse.');return}if(!rs||!rs.value){alert('Select a reason.');return}const r=await fetch(`/api/call/${id}/reassign`,{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({nurse_id:n.value,reason:rs.value})});const d=await r.json().catch(()=>({}));if(r.ok){alert('Call reassigned. Original timer preserved.');location.reload()}else alert(d.error||'Reassignment could not be completed.');}\nfunction urlBase64ToUint8Array(base64String){const padding='='.repeat((4-base64String.length%4)%4);const base64=(base64String+padding).replace(/-/g,'+').replace(/_/g,'/');const raw=atob(base64);return Uint8Array.from([...raw].map(c=>c.charCodeAt(0)))}\nfunction setNotificationUi(state){\n  const map={\n    active:{text:'🟢 Notifications Active',cls:'notification-active',disabled:false},\n    verify:{text:'🟠 Verify Notifications',cls:'notification-verify',disabled:false},\n    blocked:{text:'🔴 Notifications Blocked',cls:'notification-blocked',disabled:false},\n    checking:{text:'🔵 Checking Notifications…',cls:'notification-checking',disabled:true}\n  };\n  const cfg=map[state]||map.verify;\n  document.querySelectorAll('[data-notification-button]').forEach(btn=>{\n    btn.textContent=cfg.text;\n    btn.classList.remove('notification-active','notification-verify','notification-blocked','notification-checking');\n    btn.classList.add(cfg.cls);\n    btn.disabled=cfg.disabled;\n    btn.dataset.notificationState=state;\n  });\n}\nasync function syncNotificationUi(){\n  if(window.NATIVE_PUSH_SESSION){const b=document.querySelector('[data-notification-button]');if(b){b.textContent='Android notifications';b.title='Manage Android notifications';}return;}\n  if(!document.querySelector('.userbox')||!('Notification'in window))return;\n  if(Notification.permission==='denied'){setNotificationUi('blocked');return}\n  try{\n    const r=await fetch('/api/push/status',{cache:'no-store',headers:{'Cache-Control':'no-cache'}});\n    if(!r.ok){setNotificationUi('verify');return}\n    const d=await r.json();\n    setNotificationUi(Notification.permission==='granted'&&d.verified?'active':'verify');\n  }catch(_){setNotificationUi('verify')}\n}\nasync function enablePush(){\n  if(window.NATIVE_PUSH_SESSION){location.href='/mobile/setup';return;}\n  const current=document.querySelector('[data-notification-button][data-notification-state=\"active\"]');\n  if(current)return;\n  if(!('Notification'in window)){setNotificationUi('blocked');return}\n  if(Notification.permission==='denied'){setNotificationUi('blocked');return}\n  setNotificationUi('checking');\n  try{\n    const ok=await ensurePushHealthy(true,false);\n    if(!ok){\n      setNotificationUi(Notification.permission==='denied'?'blocked':'verify');\n      if(Notification.permission!=='denied')location.href='/notification-setup';\n      return;\n    }\n    const r=await fetch('/api/push/status',{cache:'no-store',headers:{'Cache-Control':'no-cache'}});\n    const d=r.ok?await r.json():{};\n    if(Notification.permission==='granted'&&d.verified){\n      setNotificationUi('active');\n      return;\n    }\n    setNotificationUi('verify');\n    location.href='/notification-setup';\n  }catch(_){\n    setNotificationUi('verify');\n    location.href='/notification-setup';\n  }\n}\nif('serviceWorker'in navigator){navigator.serviceWorker.register('/sw.js').catch(()=>{})}\nasync function handoverRoom(roomId){const el=document.getElementById(`handover-${roomId}`);if(!el||!el.value)return;const r=await fetch('/api/handover',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({room_id:roomId,to_nurse_id:el.value})});if(r.ok){alert('Handover request sent.');location.reload()}else alert('Handover could not be created.');}\nasync function acceptHandover(id){const r=await fetch(`/api/handover/${id}/accept`,{method:'POST'});if(r.ok)location.reload();else alert('Could not accept handover.');}\nasync function silentRefreshNow(selector,afterRefresh){\n  const current=document.querySelector(selector); if(!current)return false;\n  try{\n    const r=await fetch(location.href,{cache:'no-store',headers:{'X-Silent-Refresh':'1'}});\n    if(!r.ok)return false;\n    const html=await r.text(); const doc=new DOMParser().parseFromString(html,'text/html');\n    const next=doc.querySelector(selector); const live=document.querySelector(selector);\n    if(!next||!live)return false;\n    if((next.dataset.refreshKey||'')===(live.dataset.refreshKey||''))return false;\n    const y=window.scrollY; live.replaceWith(next); window.scrollTo(0,y); updateTimers();\n    if(typeof afterRefresh==='function')afterRefresh();\n    return true;\n  }catch(e){return false}\n}\nfunction startSilentRefresh(selector,ms,afterRefresh){\n  if(!selector)return; const key='silentRefresh:'+selector;\n  if(window[key])clearInterval(window[key]);\n  window[key]=setInterval(()=>silentRefreshNow(selector,afterRefresh),ms||3000);\n}\nfunction togglePassword(id,btn){const el=document.getElementById(id);if(!el)return;const show=el.type==='password';el.type=show?'text':'password';btn.textContent=show?'🙈':'👁';btn.setAttribute('aria-label',show?'Hide password':'Show password')}\nasync function syncPatientCallState(){\n  const timer=document.querySelector('#patient-live-root .timer[data-call-id]');\n  if(!timer)return;\n  const id=timer.dataset.callId;\n  try{\n    const r=await fetch('/api/call/'+encodeURIComponent(id)+'/status?_ts='+Date.now(),{cache:'no-store',headers:{'Cache-Control':'no-cache'}});\n    if(!r.ok)return;\n    const d=await r.json();\n    const stop=d.arrived_at||d.resolved_at||d.stop_at;\n    if(d.elapsed_seconds!=null){timer.dataset.elapsed=String(Math.max(0,Number(d.elapsed_seconds)||0));timer.dataset.clientAnchor=String(Date.now())}\n    if(d.created_at)timer.dataset.created=d.created_at;\n    if(stop){timer.dataset.stop=stop}else if(timer.dataset.stop){delete timer.dataset.stop}\n    updateTimers();\n    const rp=document.getElementById('recallPanel');\n    if(rp){\n      if(d.recall_available_at)rp.dataset.recallAt=d.recall_available_at;\n      if(d.recall_count!=null)rp.dataset.recallCount=d.recall_count;\n      if(window.updateRecallAvailability)window.updateRecallAvailability();\n    }\n    if(d.status==='resolved'){\n      if(window.silentRefreshNow){\n        const changed=await silentRefreshNow('#patient-live-root',window.rebindPatientControls);\n        if(!changed) location.reload();\n      }else{\n        location.reload();\n      }\n    }\n  }catch(e){}\n}\nfunction initPatientLiveSync(){\n  if(!document.getElementById('patient-live-root'))return;\n  if(!window.__patientStatusSync){\n    syncPatientCallState();\n    window.__patientStatusSync=setInterval(syncPatientCallState,1500);\n  }\n  if(window.startSilentRefresh&&!window.__patientSilentRefresh){\n    window.__patientSilentRefresh=true;\n    startSilentRefresh('#patient-live-root',3000,window.rebindPatientControls);\n  }\n}\nif(document.readyState==='loading')document.addEventListener('DOMContentLoaded',initPatientLiveSync);else initPatientLiveSync();\nfunction pushDeviceMeta(){const standalone=window.matchMedia('(display-mode: standalone)').matches||window.navigator.standalone===true;return{display_mode:standalone?'pwa':'browser',platform:navigator.platform||'',user_agent:navigator.userAgent||''}}\nfunction pushContextLabel(){return pushDeviceMeta().display_mode==='pwa'?'PWA':'Browser'}\nfunction setPushHealthBadge(state,text){\n  const el=document.getElementById('pushHealthBadge');if(!el)return;\n  el.className='push-health '+state;el.textContent=text;\n}\nlet pushHealthInFlight=null;\nfunction ensurePushHealthy(forcePermission=false,silent=true){\n  if(pushHealthInFlight)return pushHealthInFlight;\n  pushHealthInFlight=checkPushHealthy(forcePermission,silent).finally(()=>{pushHealthInFlight=null});\n  return pushHealthInFlight;\n}\nasync function checkPushHealthy(forcePermission=false,silent=true){\n  if(window.NATIVE_PUSH_SESSION)return true;\n  if(!document.querySelector('.userbox'))return false;\n  if(!('serviceWorker'in navigator)||!('PushManager'in window)||!('Notification'in window)){setPushHealthBadge('bad','Push unsupported');return false}\n  let permission=Notification.permission;\n  if(permission!=='granted'&&forcePermission)permission=await Notification.requestPermission();\n  if(permission!=='granted'){setPushHealthBadge('warn',permission==='denied'?'Notifications blocked':'Notifications not enabled');return false}\n  const key=''+(window.VAPID_PUBLIC_KEY||'');if(!key){setPushHealthBadge('bad','Push not configured');return false}\n  try{\n    setPushHealthBadge('checking','Checking notifications…');\n    const reg=await navigator.serviceWorker.register('/sw.js',{updateViaCache:'none'});\n    try{await reg.update()}catch(_){}\n    const ready=await navigator.serviceWorker.ready;\n    let sub=await ready.pushManager.getSubscription();\n    let mustReplace=false;\n    if(sub&&sub.options&&sub.options.applicationServerKey){\n      const expected=urlBase64ToUint8Array(key),current=new Uint8Array(sub.options.applicationServerKey);\n      mustReplace=current.length!==expected.length||current.some((v,i)=>v!==expected[i]);\n    }\n    if(sub&&!mustReplace){\n      const hr=await fetch('/api/push/health',{method:'POST',cache:'no-store',headers:{'Content-Type':'application/json','Cache-Control':'no-cache'},body:JSON.stringify({endpoint:sub.endpoint,device:pushDeviceMeta()})});\n      if(!hr.ok)throw new Error('health_failed');\n    }\n    if(sub&&mustReplace){try{await sub.unsubscribe()}catch(_){};sub=null}\n    if(!sub)sub=await ready.pushManager.subscribe({userVisibleOnly:true,applicationServerKey:urlBase64ToUint8Array(key)});\n    const sr=await fetch('/api/push/subscribe',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({subscription:sub.toJSON?sub.toJSON():sub,device:pushDeviceMeta()})});\n    if(!sr.ok)throw new Error('subscribe_failed');\n    setPushHealthBadge('ok','Push registered · '+pushContextLabel());\n    return true;\n  }catch(e){setPushHealthBadge('bad','Push needs attention');if(!silent)throw e;return false}\n}\nfunction initPushSelfHealing(){\n  if(!document.querySelector('.userbox'))return;\n  if(!('Notification'in window))return;\n  if(Notification.permission==='granted')ensurePushHealthy(false,true);\n  window.addEventListener('focus',()=>ensurePushHealthy(false,true));\n  window.addEventListener('online',()=>ensurePushHealthy(false,true));\n  window.addEventListener('pageshow',()=>ensurePushHealthy(false,true));\n  document.addEventListener('visibilitychange',()=>{if(document.visibilityState==='visible')ensurePushHealthy(false,true)});\n  if(!window.__pushHealthTimer)window.__pushHealthTimer=setInterval(()=>ensurePushHealthy(false,true),300000);\n}\nif(document.readyState==='loading')document.addEventListener('DOMContentLoaded',initPushSelfHealing);else initPushSelfHealing();\nfunction reportClientRuntimeError(payload){\n  try{fetch('/api/client-error',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify(payload),keepalive:true}).catch(()=>{})}catch(_){}\n}\nwindow.addEventListener('error',e=>{if(!e||!e.message)return;reportClientRuntimeError({type:'JavaScriptError',message:e.message,source:e.filename||'browser-js',path:location.pathname,stack:e.error&&e.error.stack||''})});\nwindow.addEventListener('unhandledrejection',e=>{const r=e&&e.reason;reportClientRuntimeError({type:'UnhandledPromiseRejection',message:r&&r.message||String(r||'Promise rejected'),source:'browser-promise',path:location.pathname,stack:r&&r.stack||''})});\nfunction escAlert(s){return String(s==null?'':s).replace(/[&<>\"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','\"':'&quot;',\"'\":'&#39;'}[c]))}\nfunction alertTimeLabel(v){try{const d=new Date(v),sec=Math.max(0,Math.floor((Date.now()-d.getTime())/1000));if(sec<60)return 'now';if(sec<3600)return Math.floor(sec/60)+'m';if(sec<86400)return Math.floor(sec/3600)+'h';return d.toLocaleDateString()}catch(_){return ''}}\nasync function refreshAlertsCenter(){\n  const badge=document.getElementById('alertsBadge'),list=document.getElementById('alertsDropList');if(!badge&&!list)return;\n  try{\n    const r=await fetch('/api/alerts',{cache:'no-store',headers:{'Cache-Control':'no-cache'}});if(!r.ok)return;\n    const d=await r.json(),n=Number(d.unread||0);\n    if(badge){badge.textContent=n>99?'99+':n;badge.hidden=n<1}\n    const st=document.getElementById('alertsDropStatus');if(st)st.textContent=n?n+' unread':'All caught up';\n    if(list){\n      const items=(d.items||[]).slice(0,8);\n      list.innerHTML=items.length?items.map(a=>'<a class=\"alerts-drop-item '+(a.read?'':'unread')+'\" href=\"'+escAlert(a.url||'/')+'\" onclick=\"markAlertRead('+Number(a.id)+')\"><span class=\"alert-mini-icon\">🔔</span><span><b>'+escAlert(a.title)+'</b><small>'+escAlert(a.body)+'</small><em>'+alertTimeLabel(a.created_at)+'</em></span></a>').join(''):'<div class=\"alerts-empty\">No alerts yet.</div>';\n    }\n  }catch(_){}\n}\nfunction toggleAlertsDropdown(e){if(e)e.stopPropagation();const d=document.getElementById('alertsDropdown');if(!d)return;d.hidden=!d.hidden;if(!d.hidden)refreshAlertsCenter()}\nasync function markAlertRead(id){try{await fetch('/api/alerts/'+id+'/read',{method:'POST',keepalive:true});}catch(_){}}\nasync function markAllAlertsRead(){try{await fetch('/api/alerts/read-all',{method:'POST'});await refreshAlertsCenter()}catch(_){}}\ndocument.addEventListener('click',e=>{const d=document.getElementById('alertsDropdown'),w=document.querySelector('.nav-alert-wrap');if(d&&!d.hidden&&w&&!w.contains(e.target))d.hidden=true});\nfunction initAlertsCenter(){if(!document.getElementById('alertsNavButton'))return;refreshAlertsCenter();if(!window.__alertsTimer)window.__alertsTimer=setInterval(refreshAlertsCenter,5000)}\nif(document.readyState==='loading')document.addEventListener('DOMContentLoaded',initAlertsCenter);else initAlertsCenter();\nfunction initNotificationUi(){\n  if(!document.querySelector('[data-notification-button]'))return;\n  syncNotificationUi();\n  window.addEventListener('focus',syncNotificationUi);\n  document.addEventListener('visibilitychange',()=>{if(document.visibilityState==='visible')syncNotificationUi()});\n}\nif(document.readyState==='loading')document.addEventListener('DOMContentLoaded',initNotificationUi);else initNotificationUi();"
MANIFEST_JSON = "{\"name\":\"Burjeel ED Smart Call\",\"short_name\":\"ED Call\",\"id\":\"/\",\"scope\":\"/\",\"start_url\":\"/\",\"display\":\"standalone\",\"background_color\":\"#f3f6fa\",\"theme_color\":\"#0059ad\",\"icons\":[{\"src\":\"/static/icons/icon-192.png\",\"sizes\":\"192x192\",\"type\":\"image/png\"},{\"src\":\"/static/icons/icon-512.png\",\"sizes\":\"512x512\",\"type\":\"image/png\"}]}"
SERVICE_WORKER_JS = "self.addEventListener('install',()=>self.skipWaiting());\nself.addEventListener('activate',e=>e.waitUntil(self.clients.claim()));\nself.addEventListener('push',event=>{\n  const receivedAt=Date.now();\n  let data={title:'Burjeel ED Call',body:'New call',url:'/nurse',kind:'call'};\n  try{data={...data,...event.data.json()}}catch(e){}\n  const task=(async()=>{\n    await self.registration.showNotification(data.title,{\n      body:data.body,\n      icon:'/static/icons/icon-192.png',\n      badge:'/static/icons/icon-192.png',\n      vibrate:[350,120,350,120,650],\n      requireInteraction:true,\n      silent:false,\n      tag:data.tag||('burjeel-ed-'+(data.kind||'call')),\n      renotify:true,\n      timestamp:Date.now(),\n      data:{url:data.url,event_id:data.event_id,received_at_ms:receivedAt}\n    });\n    const displayedAt=Date.now();\n    const reportReceipt=(async()=>{if(data.receipt_token){\n      // Best-effort telemetry runs only after showNotification resolves.\n      const controller=new AbortController();\n      const timeout=setTimeout(()=>controller.abort(),5000);\n      try{await fetch('/api/push/receipt',{method:'POST',cache:'no-store',credentials:'omit',headers:{'Content-Type':'application/json'},signal:controller.signal,body:JSON.stringify({token:data.receipt_token,received_at_ms:receivedAt,displayed_at_ms:displayedAt})})}catch(e){}finally{clearTimeout(timeout)}\n    }})();\n    // Display first: a sleeping server must never block the visible notification.\n    if(data.verification_token){\n      const controller=new AbortController();\n      const timeout=setTimeout(()=>controller.abort(),5000);\n      try{await fetch('/api/push/verify/'+encodeURIComponent(data.verification_token),{method:'POST',cache:'no-store',signal:controller.signal})}catch(e){}finally{clearTimeout(timeout)}\n    }\n    await reportReceipt;\n  })();\n  event.waitUntil(task);\n});\nself.addEventListener('pushsubscriptionchange',event=>{\n  event.waitUntil((async()=>{\n    try{\n      let key=null;\n      if(event.oldSubscription&&event.oldSubscription.options)key=event.oldSubscription.options.applicationServerKey;\n      if(!key)return;\n      const sub=await self.registration.pushManager.subscribe({userVisibleOnly:true,applicationServerKey:key});\n      await fetch('/api/push/subscribe',{method:'POST',credentials:'include',headers:{'Content-Type':'application/json'},body:JSON.stringify({subscription:sub.toJSON?sub.toJSON():sub,device:{display_mode:'service_worker',platform:'',user_agent:''}})});\n    }catch(e){}\n  })());\n});\nself.addEventListener('notificationclick',event=>{\n  event.notification.close();\n  event.waitUntil((async()=>{\n    const url=event.notification.data&&event.notification.data.url||'/nurse';\n    const wins=await clients.matchAll({type:'window',includeUncontrolled:true});\n    for(const c of wins){if('focus'in c){try{await c.navigate(url)}catch(e){};return c.focus();}}\n    return clients.openWindow(url);\n  })());\n});"
LOGO_PNG_B64 = "iVBORw0KGgoAAAANSUhEUgAAAMAAAADACAMAAABlApw1AAABgFBMVEUAH1gAR5monqKdnqChoeUAYL0AZcpoaGhyoaGgoaL/qqoAGmAANZMAYL8AYcIAYcIAX8B/f5J/rH9/////f3///38APY8Aqv9/v/+qf3+goKD/AAD/f//MzJnU1NQAAAAAW7e5ubmWl5gAIWnExMQAYcP///+qqqoAVqykpaYAHGN/f3+6uroAAP8AAFW5uroAW7cAAH4AXbsAXLkAXLm8vLwAW7m7u7sAYL0AXLq8vLy8vLwAf3+Wl5i8vLwAIGkAOImZmpsARJmZmpoAf/8APr2YmpqfoKHBwcGXmJien6AAIGkAGmoAIGiampsAIGgA//+ZmpsAIGgAOncAJGsAIWkATKQAVf8Af7/AwMAAYrTBwcFfn78AHmgAHmiKiowAX8DIyMgAAD8AYMFVVVVmmM3AwMAAGlcAYLsAHmcAMH0AZpkAAKoAS6sAVVVVqqrBwcEAZsxVqtR/f7l/f/9tkbbBwcH//wAAHmUAH2gAYL4AYL8AZsY/f79VVapff7+0OPRkAAAAgHRSTlMXHB6hBGAgBQbQA+sHqm2e0A0HAgIC1QMEBkkBAgUGAP39/f38/gEECfz9Ag0BAy8wAtCNrY1QUP1vrW8CD9DP+M/6bwIEUP7PL/2uEnCvUAGPjgUQLvoDBJATsAgwThD+CgTQAwVMDCdv9wUDDQMDcwUGBAIHKwGKqpjJRwQDCAmMjMYAABOWSURBVHja7V0Hd9tGEmZ6uZJccr0tCyoBiCTYJRaRYhEpiqKKVSNZsh3biWPHueSSXC53+eu3CxAgys4CoCBH9572PStPjizOt1N2ZnZmNoH+z1fiDsAdgDsAdwDuANwBuAPgXGquWi2Xz1qbKW4VCaiBSnxS10vFYqXyVLrlAK5y1ZOzVirFkYX/symiGerrSbJ4nk8mMY5KXrqtAAR0YhK+WCr6L5KSjoVR8D9grtxGAAJSU96VQ1+ifNK7TjFfbiEAFa1xHvq5x+hTVPHSz9dR/zYCEJurVABv8V4Apf3GLQSgoJxPgrg19A9U530ylI9RCxLxqUCZowD4JwUAX0RHtw6AIIgtPwfK6BNU9HNAl7RbB0BEVY4G4CcaAL6CpFvHAXGVo3OgRAFQkn4+AIIoKjQG5Pz0E1+CCoCoMQVBo9fTbhyAYFgcwQ9gNQoAvuS3Q70e+SrdLABBROrJao7orEeFaQxgcCCZ13qu39DH3+ZLP0jo+CYBYKqrm9jbOXtMvlGcDChTAZTxQVakAcCW1LHXDWJVKyXsKekVzIobA6AisWw4axhC9Yrsu2AzIJWiA/iEDgCzYDZ3iDSi0M8J+cbfF6VojkZ4AMIVym1a+8xxrZc5gkkUYAZAB5mDBZpENj//a523forX8+hIuwEAWHGrLl+ZS61WVQODSjVBDFdiboiOJLLXkr35cwTJCtIasQMQ0aMTn6/GcauPCQaqCSI/UCXeKB0ANkSE+h+J5Pu48yK8NUqEpV89o6op5sNLtUqnn8QDX1PigTmZP0p1CvXzgy60IiRCqu+3LYhIjEGWQQBfQgAO19fp1FuK0I8RgICFHKKfrDRZNBQq+tgdUiZr5Et3PZPJZrPdJLh4Ph9SihKh6K+yyDfpN9ZYdsLgNl8XNPTLh2SfawblyS6hnRBPvmaSjMVXtFlcABRISc0lp71LNleKpFW+QKXDLl7rhHJCOqHdXNl1ngGgFO5EC8UBQd3kwjDAi2MXK88++j5rr4x7ZQ9hADhm0GIBIB0bOrAE/Wl5gK3XCtrJQIslRHnCAFW8NgB8ouCdEP0ZB1iAFmuEpa+A7mVgBOs1QIDqWIeJt6VcE8AfUKVOfFwVcBZYDEinzU8/yDAQ0IUIOxp9vGtorYqEawE4QnmeLz1HkqCqdEVmkD98w/j0AtrKRBQivvRMwvTjII/LBSFgAniBnmI7wT/EpwqgyCwBkjdQ0wCwzWLBOoX+fz9Hmmq6joEIEkz5l3R+7l49ExAFAYt+ogJEB1dYSkCEqOY7hiX0X8t15DZFQVkWwLFkRVNYpwpvUhAw6U9bJoSlBHj59h/T/1vbanCroiAsB6DviEWwVmnvOuKBEPTjU8DcugL6F5MFbgSXH+bRf5xZMhwVLQdAcjnCfLGgvYlEFwK2AI0HcwCsk8CnBkR+/iO4XHfslItLAJh53Ei+1Ed/dEkRm/6FBCE0YcvQL7q2GvAfEvqbbtcdK7ISGYDWMxXY5aQXvkYOa8omfzy1OR8gQw5FJoFAQ/CGHlxLhdUgAQpQkaeEsSJqWtwN2H/Dj0ChZMhWZL74Gj5y/GcmVgMxIgCNFodgBPs4NDbtQzpoNW0OaEEyZCky9h8a1DOfcRpAHPilTvdQ+oJIckOpcRADdh2btoL2ghGQnJCE3qR6XUSIIgHoNejZHCOtfIUVQQ5kwAg9cgC4FyRD2fUPJMP/ocdOXFkRIwF4oYPJkAbZ2o10AAs6rhRwYaUdhOAzoncKmKDZFJUoAI7AZIh+2tBIjmg0HYZUYZMFF2zyt+5hrmP/WW2BCRoxmhKf6mCkJxlR5jsbaQaEzn33hhWeMFmwvY9DT5L8gBJMrauIVqgPJtT4XxsJZLwfoykoR8MNz9HDdEnx9mMekcAPCpu4NfAogzgwew2K9dYTVq4LDTpDiAHek0ebQCxo75jkYxJHUH4pdSVENaPHEAsy2fb+AzNZ8VukDDpjOgO8ErsCsKB9sYIKZvSuKJ10GmJA5JN4hp5SqK9h+rHEmjtmBKwYwtDHhs67/s+jsqB9cUDEa87RDWKcgQxfZABmmQyV/kzmnvWhZKOVwTTtPhaGA4rJoPgTWzsFZO0Ftmwj819TL3qiO3NI0nxFDrXuPLOzZQMwIKDRRsdhkoZTagxV0FyxcXv7HC3IJ7/HMsyy34YK4hLxgCZ5LGlt3U5MXTg+GZEjpjnY7dh8GFE3zJlfae/t7JPjzZn+Hths9CJoXQnLBDR9rz+6oD/Tnrhzj2SDDAxjmcTy9P1aQd9b1B+QbwuO//edInYcx2BoX5QFoOdySC3xt04eJwuMHSSf8c4Ig+j8RgE2rDBpY+LvkY9c8f5zU4PtFdIVZcfE2rHDIVr3JDYnlEsQQ3Du32cEsJMnpmfkT7++7vVFwniiAQAk+yioOcXH8r1o6XtFbAYk0gouyYEYQExBmEMgCMBchmqHmazPhB9AyeOgXCA1/d2kHIcLCRKXTWwZBYeU7Te1oBBbyYziZ4AtRgESxARgxsXdTJbqA0y0uIp+BOE3HcArD0wLsQFIeT5JJ59miK5Ra7QBRXgyPsXEZQFIM/T2Okh/ph2bDCloCsd3cg4tpQOapCGpzsP0Oz2i69I/YsV23JqKBDGiM6dhG/miovN8lwFgLyYZEtGuzALAtar4hIEg0ABofVLBYJ5iDADtSTwqjMQOgwNmZQmpyqBbIz+AHt59qa7zl2b8wlg7sbAAu3FDNgACYXNN9JVZ0QAYpUen9YUP8TlThgo3L0GL+h6sC0jwsSHhrftqVIqO6pfaOgNAJh4ZYkqQ8+I/Va4+mru+IAee1nVX8U6NpcXYDq3EwIBROhwAAqG19i3MAW3fU3oUqATbsQDYkINUwF2iVHW5604O5JM8LQvBsEOxGKHOODwAA0MOEiFqOpEFIHMQA/1NZo6VnmYEANAyQWwtvr4hxUY0tAosAgQAAK1Ela3F2zEA2BhGBXAGAOhT87mHLABb1wYgoKkcFUBLpQOQ6LlQdlL8+jHBiOGK0vN0KQBAPiIAkpWNZQ3ek6MByNEBVGqRAOxNrKzsNYVIQfffSMsRALjMkANA/RJIR1PJP0cxRmRotEuFMISS1VQAQJU2zZ0jVxKF+PpgSC51NKXc+AyhmnIqgBKVfspBsHXuSQzGEJXRuTCmA1ilA9DDAdgjux8v+YYcCQaE8TDIihI7GgmA6yTbMm4k4iffvKEhGWIsSXIAgNQmHUAyCMC/PpugWGWfpguouTElSW4WgNQSANp7OxN6XjZmSSIY7v95Fwc58jAyAJ5315OT7/nuL7b2Ls4LRk5cQze/BCNT3xwNdqcds7Wac8UD5HvgJK7Ui8WSbqsC6SN/q1J5zUiJowevhHoLg3nFc//1avVlebVlNIkbdzWt1fLJWjXHzEp8JZF1enrqSImvFNCrXoqj4U5Q50tkZiUa0pEk9b4RvkPou++wm6h805OkvvTqaXdkfUVVEIVvFh0133wsiu5WwoTzlkgUfuf9DYIqKj8bgALeP8nTD6T18D4zRKh5lbOWaudgCLRXTLu2srhGk07z1vI3kToAVNdOiL6Qel/rfrN18rL6+uhNUyaVVwaiYBK/n3gN2xXddcASy1Kv0AFw3qEKptWS053pxqBpnZU3T71RiTG52N5qZ7tew55MXuK/uaQDgBsEZPxnujEiID66YTbs4z9PzrfNS/1sF2gviArA7O9Jd3ZH/tRenItY6wfne3ZVSEQAm4EtGhjDRhOxb0yuQT6RnG1nTQsEQKcDAMq9XD465sMbgxuBQMjf2QrXoVKKBMCbM5DT04GCZmLc5D/Y2QrZYnMJAAC6ZPyB3njcGbyDYjzgCl+g/R1KPVT2c6jDhgoAqLgDquIGtNb6JelH6KIdoU2LrwMHGRepTWk6WqqsgEb++VbI7o75Ag6yb6P1Wcnp3Y9iYcBkL3R7il0+THfmoja6yZ3R9X0etNOO3GYGOXPUsl9m6nV6bSFaYXWZZagdu5c6lF4vRwawgdRruw7sLjOKHeV/gACsRQUw7xS7HgfOIwOoQwCoZiigz+f6dmgSvkXLbmIAAFyluEgAxtMY6I/casknXdOJXBFZ7qTFRWjWkzeQEgOA7Shtfkn9gzzS4JBSra662RCgAnEAOA9thvhkqSLBIeW8pCW35mQDC0DnT/E4cu1QWszzev0piZAlGMAcg1hdtSHIzEalOFwJZouTrQRk85/7qKcmtgRi2m1JulkjGmxIMzb5+EePZ8zEliOzh3c2Z3KBJUHNmLzRJwHOBCY/bxbBoXAAkFlalDsjKQmWDYonqFlh2aFsl0/+CpMvQUMmEnBaj4xzkm/YBgXaoWzmIRaeY3hGRoJ56SPKLEcurqiS4dBlM10JHbOmrLAKX0XxMQeaIWqfybIAgG7XbHYdOw7smXQsAPPGUrDTKr6poU+giLJmjphYDgCWIDNCkG+WAQALrPE9+ulsSQCO6X1hWt2sEIt1nwCmAQ7a1O2HpwKGAfC+I8KRw/lxK/vM6xCB5IdpmfoHPku6mBTgHkkXAYDwF1draeAhVjCJn7BqMZvzJLcfvLtNLps5dA7qWU6ExJk7wBnD3baW4Bzs7G1l2gfATbKgNDvp6a6RqRc8wuTWguznyVrY4Z4MAN4YWYbanc3CiXvznDhYFG6U6I7NTD3yZrkLhTY0b4UtQwwRUuHweOT0Qwm5B9uOs+iCWpL8V7vVbShjDE3kqkgv2C3T+OzypVG0JQBQu9v9Df8aJj/hSctiIaLemzoKRMfD9JTI0gLCF3OvemF9nPPCjqMDoE8GM0loWiaUkO9O6YPFjL4K3aF52WApg6ZNwCiYNS8sAVrsHJRsX7ihD6jkGz1aft+wSUvtGRAcegzkoxlznhOgBJ0A9x1Dy4szdj9sPekjen2onMYQhEeWT0cblxQw5xmKByiD1OcXl/J9s9t2Ba1st+HC/IJHIJtggnVDQY8MOZo9AQfPSeC49kRIG+ocR27ObWJmZT/zsIBVoy6/NzDvrCT0IzQ1D7akkAiBDCiTdChW3glz2oivv6bJ6HIYp3fvG8p8DBUeuhPSIQAo0HUNt/p30ZiKexEwqsOtBQEV0liOBkRsNemrEg/PnYsCAJqwkVIxuGdo//ugSSNtlyESvu4EzWLZbWLWFqD6Z16HJuYnAB2mc4DLod9pPZR/yCzK91XmBzQqmTW6mzlsaxtAAXRUDmAENBZwVaQ2Gqie5GuBCLbcDf+B43BSXGoN+3gStQ2D18GDAASQo5b8qj20bypaJgjBgW1JlYAuDatImjvBZwJ9KEoejJPgg8z/IMUa+puE7NlzQQgWaoxVeByCfmIiVPS+5EeABehoCWfO4wuRxt4+Fn8+TH8TkSG731hgdau6itQ5rAiiz5jOJ+pEdeY8wxUx/e8+c9JfOwxggdUdIbA7xdwTADCCpndAmS41GksFNM5pS6Q1/yuU54FBE/TTeN/4RV+zO8Xc9bmkyarpkaL8cgGNSw3IhKVnHvoDOpywHdJCSZDsaxPDCCqXwRY0bGJrbj+PKSaarcgHc/qb4RTAgUDt28/XXCOxRXz4ljnsHv9KiXrEhDjLAjrF0tRWPbE/7yvk9T77xRHmwFSsyBiBYdzI7NGoXWZ7xuAsdrstrb6b28QOi2mweV0KGCTPHllrjPwukxRL72364DaWEG3Nf8sbw2iNYgSBIkhGU0w+aPp0wNBgbIqqpLTpGHiJgi1EB8EqkIKaTFTU0FAleIB50NhmwYg14KFnTCEicw8CZhakwLFmqvGiSOBDBIGDs8lwJQmamxcgRJ/hsF9hnQLgVDNi9wTGxVIEAHg1ZhJrzLsvLW7PKd/DR5nK0mHWSHpVCHWDFQaA9KLEmPLuOM7m49U/X1/vdruHpDxSw/u4SnIZxgrPAKIGSmwAClq+xrNYkLVox4QnF7P6+Yd9pAmeOb2yPA7DAI45pzYiADI+VWcgODRpP0z6koKn6EtqJZvJDBb9rVzIMoBwT1hISIIQ8HxS/3036SfedMO+hAfnv9fioHclSMtwyBuskI+I2IGYtxY+qRefPn8bYk8FfYqq0CZ/K+bKmykaBu5ECX0DF/YZF3yq+AOlpF4nFZAzKJ3DEwCPOSjBpJi1MSlK6CeErmII/ZAOdklcinDJ60WD+n6vN8uDrwj8AxzBmSMNOphOIXfiYoMRlYWvwojwFtNXSPoV76h+OSUFGPM3uIARvXX0TyhBM59+ZxSWqNUzGwJ3pkYqhYzymBQ+2eumj5s0Nt8+JzVNglKan0BJ1sUsbGOQYq5sSBLHvYxY2x/pOa/GEREjk/zGkQsa9GjXT+BDX6LzNkshELBRItmtaCUkEV+EI0woYtnxdHc1tDwwoPcToC3B6ygQNqjlteitFVGftDMI78/8wIp8BAC0+Y9GpVj0Cp7oryJKfY3GmXwUDtCn9wnqEuUXcT3r2KM+3wgosfHyckwrLgDUkAFSYvaTCD8PACRRRm0DHMCeTnz9N7EBoEWdwEFmxou3DcCMosaQK5FDt5ADWAt8asy/RXPm4lTheAFUwrnTcapwnAA0ygUdLaDhUurt5IAZFrj6f3XyJEKz5WpPJo85xNnSGyMAIkb5erGkJ+cw+H/va9g3w0H9vNe6dVZey6mxfmK8AEwX6dlTMqGiVNIfFvEBraDyZutstbxWzeWaVrLv1nLgZ1h3AO4A3AG4A3AH4A7A//X6H4ejI2HjG/kAAAAAAElFTkSuQmCC"
import nurse_workspace
nurse_workspace.install(TEMPLATES)

template_env = Environment(loader=DictLoader(TEMPLATES), autoescape=select_autoescape(['html','xml']))

def render_template(name, context, status_code=200):
    html=template_env.get_template(name).render(**context)
    if name=='login.html': html=html.replace('</body>',"<script>if(window.BurjeelNative){const s=document.createElement('script');s.src='/static/native-login.js';document.body.appendChild(s);}</script></body>",1)
    if context.get('native_client'): html=html.replace('<head>','<head><script>window.NATIVE_PUSH_SESSION=true;</script>',1)
    return HTMLResponse(html, status_code=status_code)

@app.get('/static/nurse-workspace.css')
def nurse_workspace_css(): return Response(nurse_workspace.asset('workspace.css'),media_type='text/css')
@app.get('/static/nurse-workspace.js')
def nurse_workspace_js(): return Response(nurse_workspace.asset('workspace.js'),media_type='application/javascript')

@app.get('/static/app.css')
def inline_css(): return Response(APP_CSS, media_type='text/css')
@app.get('/static/app.js')
def inline_js(): return Response(APP_JS, media_type='application/javascript')
@app.get('/static/icons/icon-192.png')
@app.get('/static/icons/icon-512.png')
@app.get('/favicon.ico')
def inline_logo(): return Response(base64.b64decode(LOGO_PNG_B64), media_type='image/png')


def now_utc(): return datetime.now(timezone.utc)
def _as_utc(dt):
    if not dt:return None
    return dt.replace(tzinfo=timezone.utc) if dt.tzinfo is None else dt.astimezone(timezone.utc)
def _utc_iso(dt):
    dt=_as_utc(dt)
    return dt.isoformat().replace('+00:00','Z') if dt else None
def _elapsed_seconds(c):
    start=_as_utc(c.created_at)
    end=_as_utc(c.arrived_at or c.resolved_at) or now_utc()
    return max(0,int((end-start).total_seconds())) if start else 0
def pwd_hash(p):
    salt=secrets.token_hex(16); dk=hashlib.pbkdf2_hmac('sha256',p.encode(),salt.encode(),120000)
    return salt+'$'+dk.hex()
def pwd_check(p,h):
    salt,digest=h.split('$',1); dk=hashlib.pbkdf2_hmac('sha256',p.encode(),salt.encode(),120000)
    return secrets.compare_digest(dk.hex(),digest)

class User(Base):
    __tablename__='users'; id=Column(Integer,primary_key=True); name=Column(String(120),nullable=False); email=Column(String(180),unique=True,nullable=False); password_hash=Column(String(255),nullable=False); role=Column(String(40),nullable=False); active=Column(Boolean,default=True)
class UserAccessProfile(Base):
    __tablename__='user_access_profiles'; id=Column(Integer,primary_key=True); user_id=Column(Integer,ForeignKey('users.id'),unique=True,nullable=False,index=True); profile_name=Column(String(60),nullable=False); permissions_json=Column(Text,nullable=False,default='[]'); updated_at=Column(DateTime(timezone=True),default=now_utc,nullable=False); user=relationship('User')
class Room(Base):
    __tablename__='rooms'; id=Column(Integer,primary_key=True); code=Column(String(40),unique=True,nullable=False); zone=Column(String(80),default='ED Main'); qr_token=Column(String(120),unique=True,nullable=False); occupied=Column(Boolean,default=True); assigned_nurse_id=Column(Integer,ForeignKey('users.id')); assigned_nurse=relationship('User',foreign_keys=[assigned_nurse_id])
class Call(Base):
    __tablename__='calls'; id=Column(Integer,primary_key=True); room_id=Column(Integer,ForeignKey('rooms.id'),nullable=False); status=Column(String(40),default='new'); reason=Column(String(80),default='General assistance'); created_at=Column(DateTime(timezone=True),default=now_utc,nullable=False); acknowledged_at=Column(DateTime(timezone=True)); arrived_at=Column(DateTime(timezone=True)); resolved_at=Column(DateTime(timezone=True)); escalated_at=Column(DateTime(timezone=True)); assigned_nurse_id=Column(Integer,ForeignKey('users.id')); taken_over_by_id=Column(Integer,ForeignKey('users.id')); escalation_reason=Column(String(255)); room=relationship('Room'); assigned_nurse=relationship('User',foreign_keys=[assigned_nurse_id]); taken_over_by=relationship('User',foreign_keys=[taken_over_by_id])
class AuditLog(Base):
    __tablename__='audit_logs'; id=Column(Integer,primary_key=True); call_id=Column(Integer,ForeignKey('calls.id')); room_id=Column(Integer,ForeignKey('rooms.id')); user_id=Column(Integer,ForeignKey('users.id')); action=Column(String(80),nullable=False); detail=Column(String(500)); created_at=Column(DateTime(timezone=True),default=now_utc,nullable=False); user=relationship('User')
class RuntimeErrorEvent(Base):
    __tablename__='runtime_error_events'; id=Column(Integer,primary_key=True); error_id=Column(String(24),nullable=False,index=True); status=Column(String(20),default='open',nullable=False,index=True); count=Column(Integer,default=1,nullable=False); source=Column(String(120),default='http'); request_path=Column(String(240)); user_ref=Column(String(80),default='redacted'); exception_type=Column(String(120)); message=Column(Text); trace=Column(Text); first_seen=Column(DateTime(timezone=True),default=now_utc,nullable=False); last_seen=Column(DateTime(timezone=True),default=now_utc,nullable=False); recovered_at=Column(DateTime(timezone=True))
class Handover(Base):
    __tablename__='handovers'; id=Column(Integer,primary_key=True); room_id=Column(Integer,ForeignKey('rooms.id'),nullable=False); from_nurse_id=Column(Integer,ForeignKey('users.id'),nullable=False); to_nurse_id=Column(Integer,ForeignKey('users.id'),nullable=False); status=Column(String(30),default='pending'); created_at=Column(DateTime(timezone=True),default=now_utc); accepted_at=Column(DateTime(timezone=True)); room=relationship('Room'); from_nurse=relationship('User',foreign_keys=[from_nurse_id]); to_nurse=relationship('User',foreign_keys=[to_nurse_id])
class PushSubscription(Base):
    __tablename__='push_subscriptions'; id=Column(Integer,primary_key=True); user_id=Column(Integer,ForeignKey('users.id'),nullable=False); endpoint=Column(Text,unique=True,nullable=False); payload=Column(Text,nullable=False); created_at=Column(DateTime(timezone=True),default=now_utc)
class PushVerification(Base):
    __tablename__='push_verifications'; id=Column(Integer,primary_key=True); token=Column(String(120),unique=True,nullable=False); user_id=Column(Integer,ForeignKey('users.id'),nullable=False); endpoint=Column(Text,nullable=False); created_at=Column(DateTime(timezone=True),default=now_utc,nullable=False); confirmed_at=Column(DateTime(timezone=True))
class PushDevice(Base):
    __tablename__='push_devices'; id=Column(Integer,primary_key=True); user_id=Column(Integer,ForeignKey('users.id'),nullable=False); device_key=Column(String(64),unique=True,nullable=False); endpoint=Column(Text,nullable=False); display_mode=Column(String(30),default='browser'); platform=Column(String(80)); user_agent=Column(String(500)); health_status=Column(String(30),default='registered'); last_seen=Column(DateTime(timezone=True),default=now_utc,nullable=False); last_verified=Column(DateTime(timezone=True)); created_at=Column(DateTime(timezone=True),default=now_utc,nullable=False)
class InAppNotification(Base):
    __tablename__='in_app_notifications'; id=Column(Integer,primary_key=True); user_id=Column(Integer,ForeignKey('users.id'),nullable=False,index=True); kind=Column(String(60),default='alert',nullable=False); title=Column(String(220),nullable=False); body=Column(Text); url=Column(String(300),default='/'); created_at=Column(DateTime(timezone=True),default=now_utc,nullable=False,index=True); read_at=Column(DateTime(timezone=True)); user=relationship('User')

class EscalationSetting(Base):
    __tablename__='escalation_settings'; key=Column(String(80),primary_key=True); value=Column(Text,nullable=False)
class EscalationNotice(Base):
    __tablename__='escalation_notices'; call_id=Column(Integer,ForeignKey('calls.id'),primary_key=True); action=Column(String(80),primary_key=True); created_at=Column(DateTime(timezone=True),default=now_utc,nullable=False)

from native_push import registration as native_registration
from native_push import fcm as native_fcm
NativePushDevice, NativeEnrollmentChallenge = native_registration.configure(Base, User)
Base.metadata.create_all(engine)

def _is_transient_db_error(exc):
    s=str(exc).lower()
    return any(x in s for x in ['ssl error','bad record mac','unexpected eof','server closed the connection','connection reset','connection is closed'])

def _reset_db_pool():
    try: engine.dispose()
    except Exception: pass

def get_db():
    db=SessionLocal()
    try:
        yield db
    except Exception as exc:
        if _is_transient_db_error(exc): _reset_db_pool()
        raise
    finally:
        try: db.close()
        except Exception: pass

def _diag_redact(value):
    s=str(value or '')
    s=re.sub(r'(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b','<email>',s)
    s=re.sub(r'\b[0-9a-fA-F]{8}-[0-9a-fA-F-]{27,}\b','<id>',s)
    s=re.sub(r'\b\d{6,}\b','<id>',s)
    s=re.sub(r'(?<![A-Za-z0-9])[A-Za-z0-9_-]{24,}(?![A-Za-z0-9])','<token>',s)
    return s[:8000]

def _diag_user_ref(request):
    try:
        uid=request.session.get('user_id')
        return ('staff-'+hashlib.sha256(str(uid).encode()).hexdigest()[:8]) if uid else 'anonymous'
    except Exception:
        return 'redacted'

def _record_runtime_event(source,path,exc_type,message,trace_text='',user_ref='redacted',db=None):
    own=db is None; s=db or SessionLocal()
    try:
        safe_path=(path or '/')[:240]
        safe_message=_diag_redact(message)[:1800]
        safe_trace=_diag_redact(trace_text)[:7000]
        fingerprint=hashlib.sha256(f'{source}|{safe_path}|{exc_type}|{safe_message}'.encode()).hexdigest()[:12].upper()
        row=s.query(RuntimeErrorEvent).filter_by(error_id=fingerprint,status='open').order_by(RuntimeErrorEvent.last_seen.desc()).first()
        if row:
            row.count=(row.count or 0)+1; row.last_seen=now_utc(); row.message=safe_message; row.trace=safe_trace or row.trace; row.user_ref=user_ref or row.user_ref
        else:
            s.add(RuntimeErrorEvent(error_id=fingerprint,status='open',count=1,source=(source or 'runtime')[:120],request_path=safe_path,user_ref=(user_ref or 'redacted')[:80],exception_type=(exc_type or 'Error')[:120],message=safe_message,trace=safe_trace,first_seen=now_utc(),last_seen=now_utc()))
        s.commit()
    except Exception:
        try:s.rollback()
        except Exception:pass
    finally:
        if own:s.close()

def _mark_runtime_recovered(path):
    s=SessionLocal()
    try:
        rows=s.query(RuntimeErrorEvent).filter_by(request_path=path,status='open').all()
        if rows:
            stamp=now_utc()
            for row in rows: row.status='recovered'; row.recovered_at=stamp; row.last_seen=stamp
            s.commit()
    except Exception:
        try:s.rollback()
        except Exception:pass
    finally:s.close()

@app.middleware('http')
async def system_error_monitor(request:Request,call_next):
    path=request.url.path
    skip=path.startswith('/static/') or path in ['/favicon.ico','/manifest.webmanifest','/sw.js']
    try:
        response=await call_next(request)
    except Exception as exc:
        if not skip:
            source=getattr(request.scope.get('endpoint'),'__name__','http')
            _record_runtime_event(source,path,type(exc).__name__,str(exc),traceback.format_exc(limit=10),_diag_user_ref(request))
        raise
    if not skip:
        if response.status_code>=500:
            source=getattr(request.scope.get('endpoint'),'__name__','http')
            _record_runtime_event(source,path,f'HTTP_{response.status_code}',f'HTTP {response.status_code} response',user_ref=_diag_user_ref(request))
        elif response.status_code<400 and not path.startswith('/api/client-error'):
            _mark_runtime_recovered(path)
    return response

def current_user(request:Request,db:Session):
    uid=request.session.get('user_id'); return db.get(User,uid) if uid else None

ROLE_LABELS={
 'nurse':'Nurse','charge':'Nurse In Charge','nurse_supervisor':'Nurse Supervisor','manager':'Nurse Manager',
 'ed_manager':'ED Manager','hod':'HOD','admin':'System Admin'
}
ROLE_EQUIVALENT={'nurse_supervisor':'charge','hod':'ed_manager'}
PERMISSION_DEFS=[
 ('my_rooms','My Rooms','Assigned rooms, calls and handovers'),
 ('live_board','Charge Live Board','Operational room board and escalations'),
 ('wallboard','Call Bell Live Screen','Full-screen live call display'),
 ('management','KPI Management','Operational KPI dashboard and reports'),
 ('export_reports','Export Reports','Download Excel and PDF KPI reports'),
 ('assign_rooms','Assign Rooms','Assign or reassign rooms to nurses'),
 ('reassign_calls','Reassign Calls','Move an active call to another nurse'),
 ('takeover_calls','Take Over Calls','Take ownership of escalated/open calls'),
 ('admin_control','Admin Control Panel','Staff, rooms and system administration'),
 ('user_permissions','User Permissions','Manage work profiles and permissions'),
 ('diagnostics','Diagnostics','Access System Error Monitor')
]
WORK_PROFILE_PRESETS={
 'bedside_nurse':{'label':'Bedside Nurse','permissions':['my_rooms']},
 'charge_nurse':{'label':'Nurse In Charge','permissions':['my_rooms','live_board','wallboard','assign_rooms','reassign_calls','takeover_calls']},
 'nurse_supervisor':{'label':'Nurse Supervisor','permissions':['live_board','wallboard','management','assign_rooms','reassign_calls','takeover_calls']},
 'nurse_manager':{'label':'Nurse Manager','permissions':['live_board','wallboard','management','export_reports']},
 'ed_manager':{'label':'ED Manager','permissions':['live_board','wallboard','management','export_reports']},
 'hod':{'label':'HOD','permissions':['live_board','wallboard','management','export_reports']},
 'system_admin':{'label':'System Admin','permissions':[x[0] for x in PERMISSION_DEFS]}
}
ROLE_DEFAULT_PROFILE={'nurse':'bedside_nurse','charge':'charge_nurse','nurse_supervisor':'nurse_supervisor','manager':'nurse_manager','ed_manager':'ed_manager','hod':'hod','admin':'system_admin'}

def require_role(request,db,roles):
    user=current_user(request,db)
    if not user: raise HTTPException(401)
    effective=ROLE_EQUIVALENT.get(user.role,user.role)
    if roles and user.role not in roles and effective not in roles: raise HTTPException(403)
    return user

def get_user_access_profile(db,user):
    if not user:return None
    return db.query(UserAccessProfile).filter_by(user_id=user.id).first()

def get_user_permissions(db,user):
    if not user:return set()
    if user.role=='admin': return set(WORK_PROFILE_PRESETS['system_admin']['permissions'])
    rec=get_user_access_profile(db,user)
    if rec:
        try:
            raw=json.loads(rec.permissions_json or '[]')
            if isinstance(raw,list): return {p for p in raw if p in {x[0] for x in PERMISSION_DEFS}} - (OPERATIONAL_PERMISSIONS if user.role in MONITOR_ONLY_ROLES else set())
        except Exception: pass
    profile=ROLE_DEFAULT_PROFILE.get(user.role,'bedside_nurse')
    return set(WORK_PROFILE_PRESETS.get(profile,WORK_PROFILE_PRESETS['bedside_nurse'])['permissions']) - (OPERATIONAL_PERMISSIONS if user.role in MONITOR_ONLY_ROLES else set())

def require_permission(request,db,permission):
    user=current_user(request,db)
    if not user: raise HTTPException(401)
    if permission not in get_user_permissions(db,user): raise HTTPException(403)
    return user

def ctx(request,db,**kwargs):
    user=current_user(request,db)
    installation=request.session.get('native_installation')
    native_client=bool(user and installation and db.query(NativePushDevice).filter_by(user_id=user.id,installation_id=installation,active=True).first())
    return {'native_client':native_client,'request':request,'current_user':user,'current_permissions':sorted(get_user_permissions(db,user)) if user else [],'role_labels':ROLE_LABELS,'sla_seconds':SLA_SECONDS,'vapid_public_key':VAPID_PUBLIC_KEY,'utc_iso':_utc_iso,'elapsed_for_call':_elapsed_seconds,**kwargs}
def _push_device_key(endpoint,display_mode='browser',platform=''):
    raw=f'{endpoint}|{display_mode or "browser"}|{platform or ""}'
    return hashlib.sha256(raw.encode()).hexdigest()
def _upsert_push_device(db,user_id,endpoint,meta=None,verified=False,health='registered'):
    meta=meta or {}; mode=str(meta.get('display_mode') or 'browser')[:30]; platform=str(meta.get('platform') or '')[:80]; ua=str(meta.get('user_agent') or '')[:500]
    key=_push_device_key(endpoint,mode,platform)
    rec=db.query(PushDevice).filter_by(device_key=key).first()
    if not rec:
        rec=PushDevice(user_id=user_id,device_key=key,endpoint=endpoint,display_mode=mode,platform=platform,user_agent=ua,health_status=health,last_seen=now_utc(),last_verified=now_utc() if verified else None); db.add(rec)
    else:
        rec.user_id=user_id; rec.endpoint=endpoint; rec.display_mode=mode; rec.platform=platform; rec.user_agent=ua or rec.user_agent; rec.health_status=health; rec.last_seen=now_utc()
        if verified: rec.last_verified=now_utc()
    return rec
def log_action(db,action,detail='',call=None,room=None,user=None): db.add(AuditLog(action=action,detail=detail,call_id=call.id if call else None,room_id=(room.id if room else (call.room_id if call else None)),user_id=user.id if user else None))

def _store_in_app_alert(user_id,title,body,url='/',kind='alert'):
    s=SessionLocal()
    try:
        s.add(InAppNotification(user_id=user_id,title=str(title)[:220],body=str(body or '')[:1800],url=str(url or '/')[:300],kind=str(kind or 'alert')[:60]))
        s.commit()
    except Exception as exc:
        try:s.rollback()
        except Exception:pass
        if _is_transient_db_error(exc): _reset_db_pool()
        print(f'IN_APP_ALERT_STORE_ERROR user={user_id} error={type(exc).__name__}: {str(exc)[:220]}',flush=True)
    finally:
        try:s.close()
        except Exception:pass

def send_push_to_user(db,user_id,title,body,url='/nurse',extra=None,endpoint=None):
    extra=extra or {}
    if not extra.get('verification') and not extra.get('skip_in_app'):
        _store_in_app_alert(user_id,title,body,url,extra.get('kind','alert'))
    sent=0; errors=[]; stale=[]
    q=db.query(PushSubscription).filter_by(user_id=user_id)
    if endpoint: q=q.filter_by(endpoint=endpoint)
    data={'title':title,'body':body,'url':url}
    if extra: data.update(extra)
    event_id=secrets.token_hex(16)
    data['event_id']=event_id
    data['sent_at_ms']=int(now_utc().timestamp()*1000)
    native_result={'sent':0,'failed':0,'disabled':0}
    if not endpoint and not extra.get('verification'):
        try: native_result=native_fcm.dispatch_native(SessionLocal,user_id,dict(data))
        except Exception as exc: print(f'NATIVE_FCM_FAILURE event={event_id} type={type(exc).__name__}',flush=True)
    if not(webpush and VAPID_PRIVATE_KEY): return {'sent':0,'mode':'in-app-fallback','errors':['VAPID not configured'],'event_id':event_id,'native':native_result}
    for sub in q.all():
        try:
            endpoint_hash=hashlib.sha256(sub.endpoint.encode()).hexdigest()[:12]
            data['receipt_token']=push_receipt_signer.dumps({'event_id':event_id,'user_id':user_id,'endpoint_hash':endpoint_hash,'kind':data.get('kind','call'),'sent_at_ms':data['sent_at_ms']})
            push_headers={'Urgency':'high'}
            ep=(sub.endpoint or '').lower()
            if 'notify.windows.com' in ep:
                push_headers.update({'X-WNS-Type':'wns/raw','Content-Type':'application/octet-stream'})
            response=webpush(subscription_info=json.loads(sub.payload),data=json.dumps(data),vapid_private_key=VAPID_PRIVATE_KEY,vapid_claims={'sub':VAPID_SUBJECT},headers=push_headers,ttl=3600,timeout=10); sent+=1
            modes=sorted({d.display_mode for d in db.query(PushDevice).filter_by(user_id=user_id,endpoint=sub.endpoint).all()})
            print(f'WEBPUSH_ACCEPTED event={event_id} kind={data.get("kind","call")} user={user_id} endpoint_hash={endpoint_hash} contexts={modes} status={getattr(response,"status_code",None)}',flush=True)
        except WebPushException as exc:
            status=getattr(getattr(exc,'response',None),'status_code',None)
            detail=f'{status or "push_error"}: {str(exc)[:300]}'
            errors.append(detail)
            print(f'WEBPUSH_ERROR user={user_id} endpoint={sub.endpoint[:80]} status={status} detail={detail}',flush=True)
            if status in [404,410]:
                stale.append(sub)
                for dev in db.query(PushDevice).filter_by(user_id=user_id,endpoint=sub.endpoint).all():
                    dev.health_status='expired'; dev.last_seen=now_utc()
        except Exception as exc:
            errors.append(f'error: {str(exc)[:180]}')
    for sub in stale: db.delete(sub)
    if stale: db.commit()
    return {'sent':sent,'mode':'webpush','errors':errors,'event_id':event_id,'native':native_result}
@app.post('/api/push/receipt')
async def push_display_receipt(request:Request):
    # Worker callbacks must work without an open page or a logged-in session.
    # Signed, expiring tokens restrict this endpoint to events we sent.
    if len(await request.body())>4096:return Response(status_code=413)
    try:
        body=await request.json()
        payload=push_receipt_signer.loads(body.get('token',''),max_age=86400)
        received=body.get('received_at_ms'); displayed=body.get('displayed_at_ms')
        if not isinstance(received,int) or isinstance(received,bool) or not isinstance(displayed,int) or isinstance(displayed,bool) or displayed<received:
            return Response(status_code=400)
        print(f'WEBPUSH_DISPLAYED event={payload["event_id"]} kind={payload["kind"]} user={payload["user_id"]} endpoint_hash={payload["endpoint_hash"]} sent_at_ms={payload["sent_at_ms"]} received_at_ms={received} displayed_at_ms={displayed} worker=recall-receipts-v1',flush=True)
    except (BadSignature,SignatureExpired):return Response(status_code=403)
    except (ValueError,TypeError,KeyError,AttributeError):return Response(status_code=400)
    return Response(status_code=204)

MONITOR_ONLY_ROLES={'manager','ed_manager','hod'}
OPERATIONAL_PERMISSIONS={'my_rooms','assign_rooms','reassign_calls','takeover_calls','admin_control','user_permissions','diagnostics'}
ESCALATION_DEFAULTS={'nurse_supervisor':120,'manager':120,'ed_manager':120,'hod':120,'sound_repeat':120}

def escalation_settings(db):
    row=db.get(EscalationSetting,'timing')
    values=dict(ESCALATION_DEFAULTS)
    values['same_duration']=False
    if row:
        try:
            saved=json.loads(row.value)
            values['same_duration']=saved.get('same_duration') is True
            values.update({k:v for k,v in saved.items() if k in values and type(v) is int and v>0})
        except (ValueError,TypeError,AttributeError): pass
    return values

@app.post('/admin/escalation-settings')
async def save_escalation_settings(request:Request,db:Session=Depends(get_db)):
    admin=require_permission(request,db,'admin_control'); form=await request.form(); values={}
    try:
        for key in ESCALATION_DEFAULTS:
            source='uniform' if form.get('same_duration')=='on' and key!='sound_repeat' else key
            fallback='nurse_supervisor' if source=='uniform' else key
            minutes=int(str(form.get(source+'_minutes',form.get(fallback+'_minutes','0')))); seconds=int(str(form.get(source+'_seconds',form.get(fallback+'_seconds','0'))))
            if minutes<0 or seconds<0 or seconds>59 or minutes*60+seconds<=0: raise ValueError()
            values[key]=minutes*60+seconds
        if form.get('same_duration')=='on':
            for key in NURSING_DELAY_ACTIONS: values[key[1]]=values['nurse_supervisor']
    except (ValueError,TypeError):
        return JSONResponse({'error':'Enter a positive duration; seconds must be 0–59.'},400)
    row=db.get(EscalationSetting,'timing')
    if not row: row=EscalationSetting(key='timing',value='{}');db.add(row)
    values['same_duration']=form.get('same_duration')=='on'
    row.value=json.dumps(values);log_action(db,'ESCALATION_TIMING_UPDATED',json.dumps(values),user=admin);db.commit()
    request.session['escalation_saved']=True
    return RedirectResponse('/admin#escalation-settings',303)

@app.post('/admin/live-screen-sound')
async def save_live_screen_sound(request:Request,db:Session=Depends(get_db)):
    admin=require_permission(request,db,'admin_control')
    form=await request.form()
    try:
        minutes=int(str(form.get('sound_repeat_minutes','0')))
        seconds=int(str(form.get('sound_repeat_seconds','0')))
        total=minutes*60+seconds
        if minutes<0 or seconds<0 or seconds>59 or total<=0 or total>3600:
            raise ValueError()
    except (ValueError,TypeError):
        return JSONResponse({'error':'Enter a repeat duration from 00:01 to 60:00; seconds must be 0–59.'},400)
    values=escalation_settings(db)
    values['sound_repeat']=total
    row=db.get(EscalationSetting,'timing')
    if not row:
        row=EscalationSetting(key='timing',value='{}')
        db.add(row)
    row.value=json.dumps(values)
    log_action(db,'LIVE_SCREEN_SOUND_UPDATED',f'Repeat every {minutes:02d}:{seconds:02d}; until Arrived/Closed/Resolved',user=admin)
    db.commit()
    request.session['escalation_saved']=True
    return RedirectResponse('/admin#system',303)

NURSING_DELAY_STEP_SECONDS=120
NURSING_DELAY_ACTIONS=[
    ('NURSING_DELAY_SUPERVISOR_NOTIFIED','nurse_supervisor','Nurse Supervisor'),
    ('NURSING_DELAY_MANAGER_NOTIFIED','manager','Nurse Manager'),
    ('NURSING_DELAY_ED_MANAGER_NOTIFIED','ed_manager','ED Manager'),
    ('NURSING_DELAY_HOD_NOTIFIED','hod','HOD')
]

def _audit_for_call_action(db,call_id,action):
    return db.query(AuditLog).filter_by(call_id=call_id,action=action).order_by(AuditLog.created_at.desc()).first()

def _notify_nursing_delay_level(db,c,action,role,label,elapsed_seconds):
    recipients=db.query(User).filter_by(role=role,active=True).all()
    room=c.room.code if c.room else f'Room {c.room_id}'
    nurse=c.assigned_nurse.name if c.assigned_nurse else 'Unassigned'
    minutes=max(1,int(elapsed_seconds//60))
    title=f'Nursing response delay - {room}'
    elapsed_label=f'{int(elapsed_seconds)//60:02d}:{int(elapsed_seconds)%60:02d}'
    note='Monitoring alert only; patient call remains with nursing.' if role in MONITOR_ONLY_ROLES else 'Please coordinate nurse attendance.'
    body=f'No nurse arrival after {elapsed_label}. Primary nurse: {nurse}. {note}'
    sent=0
    for person in recipients:
        result=send_push_to_user(db,person.id,title,body,role_home(person.role),extra={'kind':'nursing-delay','tag':f'nursing-delay-{c.id}-{role}','call_id':c.id,'room':room,'escalation_role':role})
        sent+=int(result.get('sent',0))
        if result.get('errors') or not result.get('sent'):
            _record_runtime_event('escalation-push','/escalation','EscalationPushError',f'{label}: call {c.id}; user {person.id}; no push delivery accepted' if not result.get('sent') else f'{label}: partial push failure',db=db)
    if not recipients:
        _record_runtime_event('escalation-push','/escalation','EscalationRecipientMissing',f'No active {label} configured; call {c.id}',db=db)
    detail=f'{label} notified for delayed nursing response; room={room}; nurse={nurse}; elapsed={minutes}m; recipients={len(recipients)}; push_sent={sent}'
    log_action(db,action,detail,call=c,room=c.room)
    return len(recipients),sent

def enforce_escalations(db):
    now=now_utc(); settings=escalation_settings(db)
    active=db.query(Call).filter(Call.status.in_(['new','acknowledged','escalated','taken_over'])).order_by(Call.created_at).all()
    for c in active:
        if c.arrived_at or c.resolved_at: continue
        anchor=_as_utc(c.created_at)
        if not anchor: continue
        elapsed=max(0,(now-anchor).total_seconds())
        for action,role,label in NURSING_DELAY_ACTIONS:
            notice=db.get(EscalationNotice,(c.id,action))
            # Retain stages already sent by older versions.
            old=_audit_for_call_action(db,c.id,action)
            if notice or old:
                anchor=_as_utc(notice.created_at if notice else old.created_at);continue
            if (now-anchor).total_seconds()<settings[role]: break
            try:
                with db.begin_nested():
                    db.add(EscalationNotice(call_id=c.id,action=action,created_at=now));db.flush()
            except IntegrityError:
                db.expire_all();break
            if role=='nurse_supervisor' and c.status in ['new','acknowledged']:
                c.status='escalated';c.escalated_at=c.escalated_at or now
                c.escalation_reason='Nursing response delayed: nurse has not arrived'
            # Durable unique claim prevents parallel pollers/workers from duplicating a stage.
            db.commit()
            db.refresh(c)
            if not c.arrived_at and not c.resolved_at:
                _notify_nursing_delay_level(db,c,action,role,label,elapsed)
                db.commit()
            break

async def nursing_delay_monitor_loop():
    while True:
        db=SessionLocal()
        try:
            enforce_escalations(db)
        except Exception as exc:
            db.rollback()
            print(f'NURSING_DELAY_MONITOR_ERROR {type(exc).__name__}: {str(exc)[:300]}',flush=True)
            _record_runtime_event('escalation-monitor','/escalation',type(exc).__name__,str(exc),db=db)
        finally:
            db.close()
        await asyncio.sleep(1)

@app.on_event('startup')
async def start_nursing_delay_monitor():
    task=getattr(app.state,'nursing_delay_monitor_task',None)
    if not task or task.done():
        app.state.nursing_delay_monitor_task=asyncio.create_task(nursing_delay_monitor_loop())

@app.on_event('shutdown')
async def stop_nursing_delay_monitor():
    task=getattr(app.state,'nursing_delay_monitor_task',None)
    if task:
        task.cancel()
        try:
            await task
        except asyncio.CancelledError:
            pass
def _recall_state(db,c):
    logs=db.query(AuditLog).filter_by(call_id=c.id,action='PATIENT_RECALL').order_by(AuditLog.created_at.desc()).all()
    count=len(logs); last=logs[0].created_at if logs else c.created_at
    if last and last.tzinfo is None:last=last.replace(tzinfo=timezone.utc)
    available_at=(last+timedelta(seconds=120)) if last else now_utc()
    arrived=bool(c.arrived_at or c.resolved_at)
    can_recall=(not arrived) and count<3 and now_utc()>=available_at
    return {'recall_count':count,'recall_available_at':available_at.isoformat(),'can_recall':can_recall,'recall_limit':3}

def serialize_call(c):
    stop_at=c.arrived_at or c.resolved_at
    return {'id':c.id,'room':c.room.code,'status':c.status,'reason':c.reason,'created_at':_utc_iso(c.created_at),'acknowledged_at':_utc_iso(c.acknowledged_at),'arrived_at':_utc_iso(c.arrived_at),'resolved_at':_utc_iso(c.resolved_at),'stop_at':_utc_iso(stop_at),'elapsed_seconds':_elapsed_seconds(c),'assigned_nurse':c.assigned_nurse.name if c.assigned_nurse else None,'taken_over_by':c.taken_over_by.name if c.taken_over_by else None,'escalation_reason':c.escalation_reason}
def role_home(role): return {'nurse':'/nurse','charge':'/charge','nurse_supervisor':'/charge','manager':'/manager','ed_manager':'/manager','hod':'/manager','admin':'/admin'}.get(role,'/nurse')
PUSH_REQUIRED_ROLES={'nurse','charge','nurse_supervisor','manager','hod'}

async def _mobile_body(request):
    if len(await request.body())>8192: raise HTTPException(413)
    try: body=await request.json()
    except (ValueError,TypeError): raise HTTPException(400,'Invalid JSON')
    if not isinstance(body,dict): raise HTTPException(400,'Expected JSON object')
    return body

def _mobile_user(request,db):
    u=current_user(request,db)
    if not u or not u.active: raise HTTPException(401)
    if u.role not in PUSH_REQUIRED_ROLES: raise HTTPException(403)
    return u

def _mobile_csrf(request):
    expected=request.session.get('mobile_csrf','')
    origin=request.headers.get('origin','')
    allowed=str(request.base_url).rstrip('/')
    if origin!=allowed or not expected or not secrets.compare_digest(request.headers.get('x-csrf-token',''),expected): raise HTTPException(403)

def _mobile_credential(request):
    header=request.headers.get('authorization','')
    if not header.startswith('Bearer ') or len(header)>512: raise HTTPException(401)
    return header[7:]

@app.get('/api/mobile/status')
def mobile_status(request:Request,db:Session=Depends(get_db)):
    u=_mobile_user(request,db)
    token=request.session.setdefault('mobile_csrf',secrets.token_urlsafe(32))
    installation=request.session.get('native_installation')
    device=db.query(NativePushDevice).filter_by(user_id=u.id,installation_id=installation,active=True).first() if installation else None
    return {'csrf_token':token,'registered':bool(device),'installation_id':installation}

@app.post('/api/mobile/challenge')
def mobile_challenge(request:Request,db:Session=Depends(get_db)):
    u=_mobile_user(request,db); _mobile_csrf(request)
    return {'challenge':native_registration.issue_challenge(db,u.id)}

@app.post('/api/mobile/enroll')
async def mobile_enroll(request:Request,db:Session=Depends(get_db)):
    body=await _mobile_body(request)
    result=native_registration.exchange_challenge(db,str(body.get('challenge','')),body.get('installation_id',''),body.get('fcm_token',''),request.headers.get('authorization','')[7:] if request.headers.get('authorization','').startswith('Bearer ') else '')
    # Session linkage requires the same authenticated owner; credential exchange itself is scoped.
    u=current_user(request,db)
    if u and u.id==result['user_id']:
        request.session['native_installation']=result['installation_id']
        request.session.pop('notification_verified',None)
        request.session.pop('native_test',None)
    return result

@app.post('/api/mobile/token')
async def mobile_token(request:Request,db:Session=Depends(get_db)):
    credential=_mobile_credential(request)
    body=await _mobile_body(request)
    return native_registration.refresh_device(db,credential,body.get('fcm_token',''))

@app.post('/api/mobile/ready')
def mobile_ready(request:Request,db:Session=Depends(get_db)):
    credential=_mobile_credential(request)
    row=native_registration.authenticated_device(db,credential)
    u=_mobile_user(request,db)
    if row.user_id!=u.id or request.session.get('native_installation')!=row.installation_id: raise HTTPException(403)
    if not native_fcm.enabled(): raise HTTPException(503,'Native FCM disabled')
    test=request.session.get('native_test',{})
    if test.get('installation')!=row.installation_id or int(now_utc().timestamp())-test.get('at',0)>300: raise HTTPException(409,'Send and observe a device test first')
    verified=db.get(native_registration.Verification,row.installation_id)
    if not verified:
        verified=native_registration.Verification(installation_id=row.installation_id)
        db.add(verified)
    verified.user_id=u.id; verified.confirmed_at=now_utc(); db.commit()
    request.session['notification_verified']=True
    return {'url':role_home(u.role)}

@app.post('/api/mobile/ready-received')
async def mobile_ready_received(request:Request,db:Session=Depends(get_db)):
    row=native_registration.authenticated_device(db,_mobile_credential(request))
    u=_mobile_user(request,db)
    if row.user_id!=u.id or request.session.get('native_installation')!=row.installation_id: raise HTTPException(403)
    body=await _mobile_body(request)
    receipt=str(body.get('receipt_token',''))
    record=db.get(native_registration.Receipt,native_registration.digest(receipt)) if receipt else None
    valid=bool(record and record.installation_id==row.installation_id and record.user_id==u.id and record.credential_hash==row.credential_hash and not record.consumed)
    if valid:
        created=record.created_at
        if created.tzinfo is None: created=created.replace(tzinfo=timezone.utc)
        valid=(now_utc()-created).total_seconds()<=300
    if not valid:
        _record_runtime_event('native-notifications','/api/mobile/ready-received','NativeReceiptRejected','Device test receipt missing, expired or no longer matches enrollment.',user_ref=_diag_user_ref(request))
        raise HTTPException(409,'Device test receipt expired or invalid. Retry verification.')
    if not native_fcm.enabled(): raise HTTPException(503,'Native FCM disabled')
    from sqlalchemy import update as receipt_update
    claimed=db.execute(receipt_update(native_registration.Receipt).where(native_registration.Receipt.digest==record.digest,native_registration.Receipt.consumed.is_(False)).values(consumed=True).execution_options(synchronize_session=False)).rowcount
    if claimed!=1: db.rollback(); raise HTTPException(409,'Receipt already used')
    verified=db.get(native_registration.Verification,row.installation_id)
    if not verified:
        verified=native_registration.Verification(installation_id=row.installation_id)
        db.add(verified)
    verified.user_id=u.id; verified.confirmed_at=now_utc(); db.commit()
    request.session['notification_verified']=True
    request.session.pop('native_test',None)
    return {'url':role_home(u.role)}

@app.post('/api/mobile/resume')
def mobile_resume(request:Request,db:Session=Depends(get_db)):
    row=native_registration.authenticated_device(db,_mobile_credential(request))
    u=_mobile_user(request,db)
    if row.user_id!=u.id or request.session.get('native_installation')!=row.installation_id: raise HTTPException(403)
    if not native_fcm.enabled(): raise HTTPException(503,'Native FCM disabled')
    verified=db.get(native_registration.Verification,row.installation_id)
    if not verified or verified.user_id!=u.id: raise HTTPException(409,'Observe a device test first')
    request.session['notification_verified']=True
    return {'url':role_home(u.role)}

@app.get('/mobile/setup',response_class=HTMLResponse)
def mobile_setup(request:Request,db:Session=Depends(get_db)):
    _mobile_user(request,db)
    html=render_template('login.html',ctx(request,db,current_user=None)).body.decode()
    script="""<script src="/static/native-login.js"></script><script>
    nativeProgress(0,'Credentials verified. Preparing notifications…','');
    async function enrollNative(){
      try{const status=await fetch('/api/mobile/status').then(r=>{if(!r.ok)throw Error();return r.json()});
      const response=await fetch('/api/mobile/challenge',{method:'POST',headers:{'X-CSRF-Token':status.csrf_token}});
      if(!response.ok)throw Error();const body=await response.json();
      if(!window.BurjeelNative)throw Error();
      window.BurjeelNative.postMessage(JSON.stringify({command:'enroll',challenge:body.challenge}));
      }catch(e){nativeProgress(1,'Device registration failed. Check your connection and retry.','Retry');}
    }
    </script>"""
    return HTMLResponse(html.replace('</body>',script+'</body>'))

@app.get('/static/native-login.js')
def native_login_script():
    from native_push.login_ui import SCRIPT
    return Response(SCRIPT,media_type='application/javascript',headers={'Cache-Control':'no-store'})

@app.post('/api/mobile/test')
async def mobile_test(request:Request,db:Session=Depends(get_db)):
    u=_mobile_user(request,db); _mobile_csrf(request)
    body=await _mobile_body(request); installation=body.get('installation_id','')
    row=db.query(NativePushDevice).filter_by(user_id=u.id,installation_id=installation,active=True).first()
    if not row: raise HTTPException(404)
    event={'event_id':secrets.token_hex(16),'sent_at_ms':int(now_utc().timestamp()*1000),'title':'Burjeel ED test','kind':'test','url':role_home(u.role),'receipt_token':secrets.token_urlsafe(32)}
    db.query(native_registration.Receipt).filter_by(installation_id=installation).delete()
    receipt=native_registration.Receipt(digest=native_registration.digest(event['receipt_token']),installation_id=installation,user_id=u.id,credential_hash=row.credential_hash,created_at=now_utc(),consumed=False)
    db.add(receipt); db.commit()
    result=native_fcm.send_native_to_user(db,u.id,event,installation)
    if result.get('sent')!=1:
        db.delete(receipt); db.commit()
        _record_runtime_event('native-notifications','/api/mobile/test','NativeTestSendFailed','Provider did not accept the device test alert.',user_ref=_diag_user_ref(request))
        return JSONResponse({'event_id':event['event_id'],'provider':result},status_code=503)
    if result.get('sent')==1: request.session['native_test']={'installation':installation,'at':int(now_utc().timestamp()),'receipt_digest':native_registration.digest(event['receipt_token'])}
    return {'event_id':event['event_id'],'provider':result}

@app.post('/api/mobile/revoke')
def mobile_revoke(request:Request,db:Session=Depends(get_db)):
    native_registration.revoke_device(db,_mobile_credential(request))
    return {'ok':True}

@app.get('/api/alerts')
def alerts_api(request:Request,db:Session=Depends(get_db)):
    u=current_user(request,db)
    if not u: raise HTTPException(401)
    rows=db.query(InAppNotification).filter_by(user_id=u.id).order_by(InAppNotification.created_at.desc()).limit(30).all()
    unread=db.query(InAppNotification).filter(InAppNotification.user_id==u.id,InAppNotification.read_at.is_(None)).count()
    return {'unread':unread,'items':[{'id':x.id,'kind':x.kind,'title':x.title,'body':x.body or '','url':x.url or '/','created_at':x.created_at.isoformat(),'read':bool(x.read_at)} for x in rows]}

@app.post('/api/alerts/{alert_id}/read')
def alert_read(alert_id:int,request:Request,db:Session=Depends(get_db)):
    u=current_user(request,db)
    if not u: raise HTTPException(401)
    row=db.query(InAppNotification).filter_by(id=alert_id,user_id=u.id).first()
    if not row: raise HTTPException(404)
    row.read_at=row.read_at or now_utc(); db.commit()
    return {'ok':True}

@app.post('/api/alerts/read-all')
def alerts_read_all(request:Request,db:Session=Depends(get_db)):
    u=current_user(request,db)
    if not u: raise HTTPException(401)
    now=now_utc()
    db.query(InAppNotification).filter(InAppNotification.user_id==u.id,InAppNotification.read_at.is_(None)).update({'read_at':now},synchronize_session=False)
    db.commit()
    return {'ok':True}

@app.get('/alerts',response_class=HTMLResponse)
def alerts_page(request:Request,db:Session=Depends(get_db)):
    u=current_user(request,db)
    if not u: return RedirectResponse('/login',303)
    rows=db.query(InAppNotification).filter_by(user_id=u.id).order_by(InAppNotification.created_at.desc()).limit(250).all()
    unread=sum(1 for x in rows if not x.read_at)
    return render_template('alerts.html',ctx(request,db,alerts=rows,unread=unread))

@app.get('/healthz')
def healthz():
    """Render readiness probe: only report healthy when the app can reach its primary DB."""
    started=now_utc()
    db=SessionLocal()
    try:
        db.execute(text('SELECT 1'))
        latency_ms=max(0,round((now_utc()-started).total_seconds()*1000,1))
        return JSONResponse(
            {'status':'ok','database':'ok','latency_ms':latency_ms},
            status_code=200,
            headers={'Cache-Control':'no-store'}
        )
    except Exception as exc:
        try: db.rollback()
        except Exception: pass
        if _is_transient_db_error(exc): _reset_db_pool()
        return JSONResponse(
            {'status':'unhealthy','database':'unavailable'},
            status_code=503,
            headers={'Cache-Control':'no-store'}
        )
    finally:
        try: db.close()
        except Exception: pass
@app.get('/',response_class=HTMLResponse)
def index(request:Request,db:Session=Depends(get_db)):
    u=current_user(request,db)
    return RedirectResponse(role_home(u.role),303) if u else render_template('login.html',ctx(request,db))
@app.get('/login',response_class=HTMLResponse)
def login_page(request:Request,db:Session=Depends(get_db)): return render_template('login.html',ctx(request,db))
@app.post('/login')
async def login(request:Request,db:Session=Depends(get_db)):
    form=await request.form(); email=str(form.get('email','')).strip().lower(); password=str(form.get('password',''))
    u=db.query(User).filter_by(email=email,active=True).first()
    if u and pwd_check(password,u.password_hash):
        request.session['user_id']=u.id
        request.session.pop('notification_verified',None)
        if u.role in PUSH_REQUIRED_ROLES: return RedirectResponse('/notification-setup',303)
        return RedirectResponse(role_home(u.role),303)
    return render_template('login.html',ctx(request,db,error='Invalid credentials'),status_code=401)
@app.get('/logout')
def logout(request:Request,db:Session=Depends(get_db)):
    u=current_user(request,db); installation=request.session.get('native_installation')
    if u and installation:
        db.query(NativePushDevice).filter_by(user_id=u.id,installation_id=installation).update({'active':False})
        db.commit()
    request.session.clear()
    return RedirectResponse('/login',303)

@app.get('/notification-setup',response_class=HTMLResponse)
def notification_setup(request:Request,db:Session=Depends(get_db)):
    u=require_role(request,db,['nurse','charge','nurse_supervisor','manager','ed_manager','hod','admin'])
    next_url=role_home(u.role)
    if request.session.get('notification_verified'): return RedirectResponse(next_url,303)
    return render_template('notification_setup.html',ctx(request,db,next_url=next_url))

@app.get('/api/push/status')
def push_status(request:Request,db:Session=Depends(get_db)):
    u=require_role(request,db,['nurse','charge','nurse_supervisor','manager','ed_manager','hod','admin'])
    return {'verified':bool(request.session.get('notification_verified')),'subscriptions':db.query(PushSubscription).filter_by(user_id=u.id).count(),'permission_required':u.role in PUSH_REQUIRED_ROLES}

@app.get('/api/push/devices')
def push_devices(request:Request,db:Session=Depends(get_db)):
    u=require_role(request,db,['nurse','charge','nurse_supervisor','manager','ed_manager','hod','admin'])
    rows=db.query(PushDevice).filter_by(user_id=u.id).order_by(PushDevice.last_seen.desc()).all()
    return [{'display_mode':r.display_mode,'platform':r.platform,'health':r.health_status,'verified':bool(r.last_verified),'last_seen':r.last_seen.isoformat() if r.last_seen else None,'last_verified':r.last_verified.isoformat() if r.last_verified else None,'endpoint_tail':r.endpoint[-28:] if r.endpoint else ''} for r in rows]

@app.post('/api/push/device-status')
async def push_device_status(request:Request,db:Session=Depends(get_db)):
    u=require_role(request,db,['nurse','charge','nurse_supervisor','manager','ed_manager','hod','admin'])
    try: data=await request.json()
    except Exception: data={}
    endpoint=str(data.get('endpoint','')).strip(); meta=data.get('device') or {}
    mode=str(meta.get('display_mode') or 'browser'); platform=str(meta.get('platform') or '')
    key=_push_device_key(endpoint,mode,platform) if endpoint else ''
    subscribed=bool(endpoint and db.query(PushSubscription).filter_by(user_id=u.id,endpoint=endpoint).first())
    dev=db.query(PushDevice).filter_by(user_id=u.id,device_key=key).first() if key else None
    verified=bool(dev and dev.last_verified and dev.health_status!='expired')
    if dev: dev.last_seen=now_utc(); db.commit()
    if subscribed and verified: request.session['notification_verified']=True
    return JSONResponse({'ok':True,'subscribed':subscribed,'verified':bool(subscribed and verified),'display_mode':mode,'health':dev.health_status if dev else 'unregistered'},headers={'Cache-Control':'no-store'})


@app.post('/api/push/test')
async def push_test(request:Request,db:Session=Depends(get_db)):
    u=require_role(request,db,['nurse','charge','nurse_supervisor','manager','ed_manager','hod','admin'])
    if not VAPID_PUBLIC_KEY or not VAPID_PRIVATE_KEY:
        return JSONResponse({'ok':False,'error':'VAPID push keys are not configured.'},503)
    try:
        data=await request.json()
        if not isinstance(data,dict): data={}
        endpoint=str(data.get('endpoint','')).strip()
        meta=data.get('device') or {}
        sub=db.query(PushSubscription).filter_by(user_id=u.id,endpoint=endpoint).first() if endpoint else db.query(PushSubscription).filter_by(user_id=u.id).order_by(PushSubscription.created_at.desc()).first()
        if not sub:return JSONResponse({'ok':False,'error':'This browser/PWA push subscription is not registered. Please retry setup.'},409)
        endpoint=sub.endpoint
        _upsert_push_device(db,u.id,endpoint,meta,verified=False,health='testing')
        token=secrets.token_urlsafe(32)
        db.query(PushVerification).filter(PushVerification.user_id==u.id,PushVerification.endpoint==endpoint,PushVerification.confirmed_at.is_(None)).delete(synchronize_session=False)
        db.add(PushVerification(token=token,user_id=u.id,endpoint=endpoint));db.commit()
        result=send_push_to_user(db,u.id,'Burjeel ED Call - Test Alert','Notification test successful. This device is ready for nurse calls.',role_home(u.role),extra={'verification_token':token,'verification':True},endpoint=endpoint)
        if result.get('sent',0)<1:
            detail=(result.get('errors') or ['Push service rejected the subscription'])[0]
            return JSONResponse({'ok':False,'error':'Push subscription was rejected. Please retry setup to create a fresh subscription.','detail':detail},502)
        log_action(db,'PUSH_TEST_SENT',f'Test notification sent to {u.name}',user=u);db.commit()
        return {'ok':True,'sent':result.get('sent',0),'verification_token':token}
    except Exception as exc:
        db.rollback()
        print(f'PUSH_TEST_ERROR user={u.id} error={type(exc).__name__}: {str(exc)[:300]}',flush=True)
        return JSONResponse({'ok':False,'error':'Notification verification failed on the server. Please retry.','detail':f'{type(exc).__name__}: {str(exc)[:160]}'},500)

@app.post('/api/push/verify/{token}')
def push_verify_callback(token:str,db:Session=Depends(get_db)):
    rec=db.query(PushVerification).filter_by(token=token).first()
    if not rec:return JSONResponse({'ok':False},404)
    created=rec.created_at
    if created and created.tzinfo is None: created=created.replace(tzinfo=timezone.utc)
    if created and (now_utc()-created).total_seconds()>300:return JSONResponse({'ok':False,'error':'Verification expired'},410)
    if not rec.confirmed_at:rec.confirmed_at=now_utc();db.commit()
    return {'ok':True}

@app.get('/api/push/verification/{token}')
def push_verification_status(token:str,request:Request,db:Session=Depends(get_db)):
    u=require_role(request,db,['nurse','charge','nurse_supervisor','manager','ed_manager','hod','admin'])
    rec=db.query(PushVerification).filter_by(token=token,user_id=u.id).first()
    if not rec:return JSONResponse({'ok':False,'verified':False},404)
    verified=bool(rec.confirmed_at)
    if verified:
        mode=request.query_params.get('mode','browser'); platform=request.query_params.get('platform','')
        _upsert_push_device(db,u.id,rec.endpoint,{'display_mode':mode,'platform':platform,'user_agent':request.headers.get('user-agent','')},verified=True,health='active')
        request.session['notification_verified']=True; db.commit()
    return {'ok':True,'verified':verified}

@app.get('/room/{token}',response_class=HTMLResponse)
def patient_room(token:str,request:Request,db:Session=Depends(get_db)):
    enforce_escalations(db); room=db.query(Room).filter_by(qr_token=token).first()
    if not room: raise HTTPException(404)
    active=db.query(Call).filter_by(room_id=room.id).filter(Call.status.in_(['new','acknowledged','escalated','taken_over','arrived'])).order_by(Call.created_at.desc()).first()
    lang=request.query_params.get('lang','en').lower()
    if lang not in ['en','ar']: lang='en'
    translations={
      'en':{'page_title':'Call Nurse','help_title':'How can we help?','help_note':'Choose a reason, then tap the call button.','general':'General assistance','pain':'Pain','toilet':'Toilet assistance','medication':'IV / Medication','call_nurse':'CALL NURSE','call_active':'Call sent','wait_note':'Your nurse has been notified.','ack_note':'Your nurse acknowledged the call. The timer will continue until the nurse arrives.','arrived_note':'Your nurse has arrived. The response timer is now stopped.','recall_message':'Nurse has not arrived yet. If you still need assistance, please press Re-call.','recall_button':'RE-CALL','recall_sent':'Re-call sent ✓','physical_note':'For urgent or life-threatening needs, use the physical emergency call bell immediately.','error':'Unable to send the call. Please use the physical call bell.'},
      'ar':{'page_title':'استدعاء الممرضة','help_title':'كيف يمكننا مساعدتك؟','help_note':'اختر سبب النداء ثم اضغط زر استدعاء الممرضة.','general':'مساعدة عامة','pain':'ألم','toilet':'مساعدة للحمام','medication':'المحلول / الدواء','call_nurse':'استدعاء الممرضة','call_hint':'اضغط لإرسال النداء','call_active':'تم إرسال النداء','wait_note':'تم إشعار الممرضة المسؤولة عن الغرفة.','ack_note':'تم تأكيد النداء. سيستمر العداد حتى حضور الممرضة للغرفة.','arrived_note':'حضرت الممرضة إلى الغرفة وتم إيقاف عداد الاستجابة.','recall_message':'لم تصل الممرضة بعد. إذا كنت ما زلت بحاجة للمساعدة، يرجى الضغط على إعادة النداء.','recall_button':'إعادة النداء','recall_sent':'تمت إعادة النداء ✓','physical_note':'للحالات العاجلة أو المهددة للحياة استخدم زر النداء الفعلي فورًا.','error':'تعذر إرسال النداء. يرجى استخدام زر النداء الفعلي.'}
    }
    status_en={'new':'Waiting for nurse','acknowledged':'Nurse acknowledged','escalated':'Escalated to nurse in charge','taken_over':'Nurse in charge responding','arrived':'Nurse arrived','resolved':'Resolved'}
    status_ar={'new':'بانتظار استجابة الممرضة','acknowledged':'تم تأكيد النداء','escalated':'تم التصعيد إلى الممرضة المسؤولة','taken_over':'الممرضة المسؤولة تتولى النداء','arrived':'حضرت الممرضة','resolved':'تم إغلاق النداء'}
    status=(status_ar if lang=='ar' else status_en).get(active.status if active else '',translations[lang]['call_active'])
    recall=_recall_state(db,active) if active else {'recall_count':0,'recall_available_at':'','can_recall':False,'recall_limit':3}
    active_elapsed=_elapsed_seconds(active) if active else 0
    active_created_iso=_utc_iso(active.created_at) if active else ''
    active_stop_iso=_utc_iso(active.arrived_at or active.resolved_at) if active else ''
    refresh_key=('idle' if not active else f'{active.id}:{active.status}:{_utc_iso(active.acknowledged_at) or ""}:{_utc_iso(active.arrived_at) or ""}:{_utc_iso(active.resolved_at) or ""}:{recall["recall_count"]}:{recall["recall_available_at"]}')
    return render_template('patient.html',ctx(request,db,room=room,active_call=active,active_elapsed=active_elapsed,active_created_iso=active_created_iso,active_stop_iso=active_stop_iso,lang=lang,t=translations[lang],status_label=status,refresh_key=refresh_key,recall=recall))
@app.post('/room/{token}/call')
async def patient_call_form(token:str,request:Request,db:Session=Depends(get_db)):
    room=db.query(Room).filter_by(qr_token=token,occupied=True).first()
    if not room: raise HTTPException(404)
    form=await request.form()
    reason=str(form.get('reason','General assistance')).strip()
    lang=str(form.get('lang','en')).strip().lower()
    if lang not in ['en','ar']: lang='en'
    allowed=['General assistance','Pain','Toilet assistance','IV / Medication']
    if reason not in allowed: reason='General assistance'
    active=db.query(Call).filter_by(room_id=room.id).filter(Call.status.in_(['new','acknowledged','escalated','taken_over','arrived'])).first()
    if not active:
        c=Call(room_id=room.id,assigned_nurse_id=room.assigned_nurse_id,reason=reason,status='new',created_at=now_utc())
        db.add(c);db.flush();log_action(db,'PATIENT_CALL_CREATED',f'{room.code}: {reason}',room=room,call=c)
        db.commit()
        if c.assigned_nurse_id: send_push_to_user(db,c.assigned_nurse_id,f'Patient Call - {room.code}',reason,'/nurse')
    return RedirectResponse(f'/room/{token}?lang={lang}',303)

@app.post('/api/room/{token}/call')
async def patient_call(token:str,request:Request,db:Session=Depends(get_db)):
    room=db.query(Room).filter_by(qr_token=token).first();
    if not room: raise HTTPException(404)
    if not room.occupied or not room.assigned_nurse_id: return JSONResponse({'ok':False,'error':'Room is not ready for digital call. Please use the physical call bell.'},409)
    active=db.query(Call).filter_by(room_id=room.id).filter(Call.status.in_(['new','acknowledged','escalated','taken_over','arrived'])).first()
    if active:return {'ok':True,'call_id':active.id,'status':active.status,'duplicate':True}
    data=await request.json(); c=Call(room_id=room.id,assigned_nurse_id=room.assigned_nurse_id,reason=data.get('reason','General assistance'),created_at=now_utc()); db.add(c); db.flush(); log_action(db,'CALL_CREATED',f'Patient call: {c.reason}',call=c); db.commit(); send_push_to_user(db,room.assigned_nurse_id,f'Call Bell - {room.code}',f'Patient requests {c.reason}','/nurse'); return {'ok':True,'call_id':c.id,'status':c.status}
def _perform_patient_recall(db,c):
    if not c:return {'ok':False,'status_code':404,'error':'Call not found.'}
    if c.arrived_at or c.resolved_at or c.status=='resolved':
        return {'ok':False,'status_code':409,'error':'Nurse has already arrived.','arrived':True}
    state=_recall_state(db,c)
    if state['recall_count']>=state['recall_limit']:
        return {'ok':False,'status_code':409,'error':'Maximum re-calls reached. The nurse in charge has been notified.','recall_count':state['recall_count']}
    if not state['can_recall']:
        return {'ok':False,'status_code':429,'error':'Re-call is available every 2 minutes until the nurse arrives.','recall_available_at':state['recall_available_at']}
    new_count=state['recall_count']+1
    log_action(db,'PATIENT_RECALL',f'{c.room.code}: patient re-call #{new_count}',call=c,room=c.room)
    if new_count>=3 and c.status!='taken_over':
        c.status='escalated'; c.escalated_at=c.escalated_at or now_utc(); c.escalation_reason='Patient re-called 3 times and nurse has not arrived'
    db.commit()
    if c.assigned_nurse_id:
        send_push_to_user(db,c.assigned_nurse_id,f'Re-call #{new_count} - {c.room.code}',f'Patient is still waiting: {c.reason}','/nurse',extra={'kind':'recall','tag':f'recall-{c.id}-{new_count}'})
    if new_count>=3:
        for charge in db.query(User).filter_by(role='charge',active=True).all():
            send_push_to_user(db,charge.id,f'Patient Re-call Escalation - {c.room.code}',f'Patient re-called 3 times and nurse has not arrived.','/charge',extra={'kind':'recall-escalation','tag':f'recall-escalation-{c.id}'})
        log_action(db,'PATIENT_RECALL_ESCALATED',f'{c.room.code}: third re-call escalated to charge nurse',call=c,room=c.room); db.commit()
    state=_recall_state(db,c)
    return {'ok':True,'status_code':200,'recall_count':state['recall_count'],'recall_available_at':state['recall_available_at'],'can_recall':state['can_recall'],'escalated':new_count>=3}

@app.post('/api/call/{call_id}/recall')
def patient_recall(call_id:int,db:Session=Depends(get_db)):
    result=_perform_patient_recall(db,db.get(Call,call_id))
    code=result.pop('status_code',200)
    return JSONResponse(result,status_code=code)

@app.post('/room/{token}/recall/{call_id}')
def patient_recall_form(token:str,call_id:int,request:Request,db:Session=Depends(get_db)):
    room=db.query(Room).filter_by(qr_token=token).first()
    c=db.get(Call,call_id)
    if not room or not c or c.room_id!=room.id: raise HTTPException(404)
    result=_perform_patient_recall(db,c)
    lang=request.query_params.get('lang','en')
    if lang not in ['en','ar']: lang='en'
    return RedirectResponse(f'/room/{token}?lang={lang}&recall={"sent" if result.get("ok") else "failed"}',303)

@app.get('/api/call/{call_id}/status')
def call_status(call_id:int,db:Session=Depends(get_db)):
    enforce_escalations(db); c=db.get(Call,call_id)
    if not c: raise HTTPException(404)
    data=serialize_call(c); data.update(_recall_state(db,c))
    return JSONResponse(data,headers={'Cache-Control':'no-store, no-cache, must-revalidate, max-age=0'})
@app.get('/nurse',response_class=HTMLResponse)
def nurse_page(request:Request,db:Session=Depends(get_db)):
    u=require_permission(request,db,'my_rooms')
    if u.role in ['nurse','charge','nurse_supervisor'] and not request.session.get('notification_verified'): return RedirectResponse('/notification-setup',303)
    enforce_escalations(db)
    rooms=db.query(Room).filter_by(assigned_nurse_id=u.id).order_by(Room.code).all()
    calls=db.query(Call).filter_by(assigned_nurse_id=u.id).filter(Call.status.in_(['new','acknowledged','escalated','taken_over','arrived'])).order_by(Call.created_at).all()
    colleagues=db.query(User).filter(User.role=='nurse',User.id!=u.id,User.active==True).order_by(User.name).all()
    pending=db.query(Handover).filter_by(to_nurse_id=u.id,status='pending').order_by(Handover.created_at.desc()).all()
    room_key='|'.join(f'{r.id}:{r.assigned_nurse_id}:{int(r.occupied)}' for r in rooms)
    call_key='|'.join(f'{c.id}:{c.status}:{c.assigned_nurse_id}:{c.acknowledged_at.isoformat() if c.acknowledged_at else ""}:{c.arrived_at.isoformat() if c.arrived_at else ""}' for c in calls)
    handover_key='|'.join(f'{h.id}:{h.status}:{h.room_id}:{h.from_nurse_id}:{h.to_nurse_id}' for h in pending)
    refresh_key=hashlib.sha256((room_key+'#'+call_key+'#'+handover_key).encode()).hexdigest()[:20]
    return render_template('nurse.html',ctx(request,db,rooms=rooms,calls=calls,colleagues=colleagues,pending_handovers=pending,refresh_key=refresh_key))

@app.get('/charge',response_class=HTMLResponse)
def charge_page(request:Request,db:Session=Depends(get_db)):
    u=require_permission(request,db,'live_board')
    if u.role in ['charge','nurse_supervisor'] and not request.session.get('notification_verified'): return RedirectResponse('/notification-setup',303)
    enforce_escalations(db)
    calls=db.query(Call).filter(Call.status.in_(['new','acknowledged','escalated','taken_over','arrived'])).order_by(Call.created_at).all()
    rooms=db.query(Room).order_by(Room.code).all()
    nurses=db.query(User).filter_by(role='nurse',active=True).order_by(User.name).all()
    call_key='|'.join(f'{c.id}:{c.status}:{c.assigned_nurse_id}:{c.taken_over_by_id or ""}:{c.acknowledged_at.isoformat() if c.acknowledged_at else ""}:{c.arrived_at.isoformat() if c.arrived_at else ""}' for c in calls)
    room_key='|'.join(f'{r.id}:{r.assigned_nurse_id or ""}:{int(r.occupied)}' for r in rooms)
    refresh_key=hashlib.sha256((call_key+'#'+room_key).encode()).hexdigest()[:20]
    return render_template('charge.html',ctx(request,db,calls=calls,rooms=rooms,nurses=nurses,refresh_key=refresh_key))
def _aware(dt):
    if not dt: return None
    return dt.replace(tzinfo=timezone.utc) if dt.tzinfo is None else dt

def _fmt_seconds(sec):
    if sec is None: return '-'
    sec=max(0,int(sec)); return f'{sec//60:02d}:{sec%60:02d}'

def _management_range(request):
    period=request.query_params.get('period','today')
    now=now_utc()
    if period=='week':
        start=(now-timedelta(days=6)).replace(hour=0,minute=0,second=0,microsecond=0); end=now
    elif period=='month':
        start=now.replace(day=1,hour=0,minute=0,second=0,microsecond=0); end=now
    elif period=='custom':
        try:
            start=datetime.strptime(request.query_params.get('from',''),'%Y-%m-%d').replace(tzinfo=timezone.utc)
            end=datetime.strptime(request.query_params.get('to',''),'%Y-%m-%d').replace(tzinfo=timezone.utc)+timedelta(days=1)-timedelta(microseconds=1)
        except Exception:
            period='today'; start=now.replace(hour=0,minute=0,second=0,microsecond=0); end=now
    else:
        period='today'; start=now.replace(hour=0,minute=0,second=0,microsecond=0); end=now
    return period,start,end

def _management_data(request,db):
    period,start,end=_management_range(request)
    calls=db.query(Call).filter(Call.created_at>=start,Call.created_at<=end).order_by(Call.created_at.desc()).all()
    call_rows=[]; response_secs=[]; completed_secs=[]
    nurse_map={}; room_map={}
    for c in calls:
        created=_aware(c.created_at); ack=_aware(c.acknowledged_at); arrived=_aware(c.arrived_at); resolved=_aware(c.resolved_at)
        response_end=arrived or resolved
        response=(response_end-created).total_seconds() if response_end and created else None
        duration=(resolved-created).total_seconds() if resolved and created else None
        if response is not None: response_secs.append(response)
        if duration is not None: completed_secs.append(duration)
        nurse_name=c.assigned_nurse.name if c.assigned_nurse else 'Unassigned'
        room_code=c.room.code if c.room else '-'; zone=c.room.zone if c.room else '-'
        call_rows.append({'id':c.id,'room':room_code,'zone':zone,'created':created.strftime('%Y-%m-%d %H:%M:%S') if created else '-','nurse':nurse_name,'status':c.status,'status_label':c.status.replace('_',' ').title(),'response':_fmt_seconds(response),'response_seconds':response,'duration':_fmt_seconds(duration),'duration_seconds':duration,'escalated':'Yes' if c.escalated_at else 'No','takeover':c.taken_over_by.name if c.taken_over_by else '-','reason':c.reason,'acknowledged':ack.strftime('%Y-%m-%d %H:%M:%S') if ack else '-','arrived':arrived.strftime('%Y-%m-%d %H:%M:%S') if arrived else '-','resolved':resolved.strftime('%Y-%m-%d %H:%M:%S') if resolved else '-'})
        nm=nurse_map.setdefault(nurse_name,{'name':nurse_name,'calls':0,'responses':[],'over_sla':0,'escalated':0})
        nm['calls']+=1
        if response is not None: nm['responses'].append(response)
        if response is not None and response>SLA_SECONDS: nm['over_sla']+=1
        if c.escalated_at: nm['escalated']+=1
        rm=room_map.setdefault(room_code,{'room':room_code,'zone':zone,'calls':0,'responses':[],'over_sla':0})
        rm['calls']+=1
        if response is not None: rm['responses'].append(response)
        if response is not None and response>SLA_SECONDS: rm['over_sla']+=1
    nurse_perf=[]
    for n in nurse_map.values():
        avg=sum(n['responses'])/len(n['responses']) if n['responses'] else None
        nurse_perf.append({'name':n['name'],'calls':n['calls'],'avg_response':_fmt_seconds(avg),'avg_seconds':avg,'over_sla':n['over_sla'],'escalated':n['escalated']})
    nurse_perf.sort(key=lambda x:(-x['calls'],x['name']))
    room_perf=[]
    for r in room_map.values():
        avg=sum(r['responses'])/len(r['responses']) if r['responses'] else None
        room_perf.append({'room':r['room'],'zone':r['zone'],'calls':r['calls'],'avg_response':_fmt_seconds(avg),'avg_seconds':avg,'over_sla':r['over_sla']})
    room_perf.sort(key=lambda x:(-x['calls'],x['room']))
    avg=sum(response_secs)/len(response_secs) if response_secs else None
    median=None
    if response_secs:
        vals=sorted(response_secs); mid=len(vals)//2
        median=vals[mid] if len(vals)%2 else (vals[mid-1]+vals[mid])/2
    over2=sum(1 for x in response_secs if x>120); over5=sum(1 for x in response_secs if x>300)
    sla_ok=sum(1 for x in response_secs if x<=SLA_SECONDS)
    sla_rate=f'{(sla_ok/len(response_secs)*100):.1f}%' if response_secs else '-'
    kpi={'total':len(calls),'avg_response':_fmt_seconds(avg),'median_response':_fmt_seconds(median),'over_2m':over2,'over_5m':over5,'escalations':sum(1 for c in calls if c.escalated_at),'takeovers':sum(1 for c in calls if c.taken_over_by_id),'sla_rate':sla_rate}
    handovers=db.query(Handover).filter(Handover.created_at>=start,Handover.created_at<=end).order_by(Handover.created_at.desc()).all()
    audits=db.query(AuditLog).filter(AuditLog.created_at>=start,AuditLog.created_at<=end).order_by(AuditLog.created_at.desc()).all()
    filter_qs=f'period={period}'
    from_date=request.query_params.get('from',''); to_date=request.query_params.get('to','')
    if period=='custom' and from_date and to_date: filter_qs+=f'&from={from_date}&to={to_date}'
    range_label=f'{start.strftime("%Y-%m-%d")} to {end.strftime("%Y-%m-%d")}'
    return {'period':period,'start':start,'end':end,'calls':calls,'call_rows':call_rows,'nurse_perf':nurse_perf,'room_perf':room_perf,'kpi':kpi,'handovers':handovers,'audits':audits,'filter_qs':filter_qs,'from_date':from_date,'to_date':to_date,'range_label':range_label}

@app.get('/manager',response_class=HTMLResponse)
def manager_page(request:Request,db:Session=Depends(get_db)):
    require_permission(request,db,'management'); enforce_escalations(db)
    d=_management_data(request,db)
    return render_template('manager.html',ctx(request,db,**d))

@app.get('/manager/export.xlsx')
def manager_export_excel(request:Request,db:Session=Depends(get_db)):
    require_permission(request,db,'export_reports'); d=_management_data(request,db)
    wb=Workbook(); ws=wb.active; ws.title='Executive Dashboard'
    navy='0B2A4A'; blue='0B67C2'; light='EAF2FB'; green='D1FADF'; amber='FEF0C7'; red='FEE4E2'; white='FFFFFF'
    thin=Side(style='thin',color='D0D5DD')
    ws['A1']='Burjeel ED Smart Call Bell - Executive KPI Dashboard'; ws['A1'].font=Font(size=16,bold=True,color=white); ws['A1'].fill=PatternFill('solid',fgColor=navy); ws.merge_cells('A1:F1')
    ws['A2']='Reporting period'; ws['B2']=d['range_label']; ws['A3']='Generated UTC'; ws['B3']=now_utc().strftime('%Y-%m-%d %H:%M:%S')
    metrics=[('Total Calls',d['kpi']['total']),('Average Response',d['kpi']['avg_response']),('Median Response',d['kpi']['median_response']),('Calls >2 min',d['kpi']['over_2m']),('Calls >5 min',d['kpi']['over_5m']),('Escalations',d['kpi']['escalations']),('Charge Takeovers',d['kpi']['takeovers']),('SLA Compliance',d['kpi']['sla_rate'])]
    row=5
    for label,value in metrics:
        ws.cell(row=row,column=1,value=label).font=Font(bold=True); ws.cell(row=row,column=2,value=value); row+=1
    ws['D5']='Calls by Room'; ws['D5'].font=Font(bold=True,color=white); ws['D5'].fill=PatternFill('solid',fgColor=blue); ws['E5']='Calls'; ws['E5'].font=Font(bold=True,color=white); ws['E5'].fill=PatternFill('solid',fgColor=blue)
    for i,r in enumerate(d['room_perf'][:12],start=6): ws.cell(i,4,r['room']); ws.cell(i,5,r['calls'])
    if d['room_perf']:
        chart=BarChart(); chart.title='Call Volume by Room'; chart.y_axis.title='Calls'; chart.x_axis.title='Room'
        chart.add_data(Reference(ws,min_col=5,min_row=5,max_row=5+min(12,len(d['room_perf']))),titles_from_data=True)
        chart.set_categories(Reference(ws,min_col=4,min_row=6,max_row=5+min(12,len(d['room_perf'])))); chart.height=7; chart.width=12; ws.add_chart(chart,'G5')
    for col in range(1,6): ws.column_dimensions[get_column_letter(col)].width=22

    ws2=wb.create_sheet('Call Details')
    headers=['Call ID','Room','Zone','Reason','Primary Nurse','Status','Created','Acknowledged','Arrived','Resolved','Response Time','Response Seconds','Total Duration','Duration Seconds','Escalated','Charge Takeover','SLA Status','Reassignment History']
    ws2.append(headers)
    for cell in ws2[1]: cell.font=Font(bold=True,color=white); cell.fill=PatternFill('solid',fgColor=navy); cell.alignment=Alignment(horizontal='center')
    audit_by_call={}
    for a in d['audits']:
        if a.call_id: audit_by_call.setdefault(a.call_id,[]).append(f'{a.action}: {a.detail or ""}')
    for c in d['call_rows']:
        sla='Within SLA' if c['response_seconds'] is not None and c['response_seconds']<=SLA_SECONDS else ('Over SLA' if c['response_seconds'] is not None else 'Open/No arrival')
        ws2.append([c['id'],c['room'],c['zone'],c['reason'],c['nurse'],c['status_label'],c['created'],c['acknowledged'],c['arrived'],c['resolved'],c['response'],c['response_seconds'],c['duration'],c['duration_seconds'],c['escalated'],c['takeover'],sla,' | '.join(audit_by_call.get(c['id'],[]))])
    ws2.freeze_panes='A2'; ws2.auto_filter.ref=ws2.dimensions

    ws3=wb.create_sheet('Nurse Performance'); ws3.append(['Nurse','Calls','Average Response','Average Seconds','Over SLA','Escalated'])
    for n in d['nurse_perf']: ws3.append([n['name'],n['calls'],n['avg_response'],n['avg_seconds'],n['over_sla'],n['escalated']])
    ws4=wb.create_sheet('Room Performance'); ws4.append(['Room','Zone','Calls','Average Response','Average Seconds','SLA Breaches'])
    for r in d['room_perf']: ws4.append([r['room'],r['zone'],r['calls'],r['avg_response'],r['avg_seconds'],r['over_sla']])
    ws5=wb.create_sheet('SLA & Escalation'); ws5.append(['Call ID','Room','Created','Response','Response Seconds','Escalated','Takeover','SLA Status'])
    for c in d['call_rows']:
        sla='Within SLA' if c['response_seconds'] is not None and c['response_seconds']<=SLA_SECONDS else ('Over SLA' if c['response_seconds'] is not None else 'Open/No arrival')
        ws5.append([c['id'],c['room'],c['created'],c['response'],c['response_seconds'],c['escalated'],c['takeover'],sla])
    ws6=wb.create_sheet('Handover'); ws6.append(['Room','From Nurse','To Nurse','Status','Created','Accepted'])
    for h in d['handovers']: ws6.append([h.room.code,h.from_nurse.name,h.to_nurse.name,h.status,_aware(h.created_at).strftime('%Y-%m-%d %H:%M:%S') if h.created_at else '-',_aware(h.accepted_at).strftime('%Y-%m-%d %H:%M:%S') if h.accepted_at else '-'])
    ws7=wb.create_sheet('Audit Summary'); ws7.append(['Time','Action','User','Room ID','Call ID','Detail'])
    for a in d['audits']: ws7.append([_aware(a.created_at).strftime('%Y-%m-%d %H:%M:%S') if a.created_at else '-',a.action,a.user.name if a.user else '-',a.room_id,a.call_id,a.detail or ''])
    for sheet in [ws3,ws4,ws5,ws6,ws7]:
        for cell in sheet[1]: cell.font=Font(bold=True,color=white); cell.fill=PatternFill('solid',fgColor=navy)
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
    for sheet in wb.worksheets:
        for row_cells in sheet.iter_rows():
            for cell in row_cells:
                cell.border=Border(bottom=thin); cell.alignment=Alignment(vertical='top')
        for column in range(1,min(sheet.max_column,18)+1):
            max_len=max([len(str(sheet.cell(r,column).value or '')) for r in range(1,min(sheet.max_row,200)+1)] or [10]); sheet.column_dimensions[get_column_letter(column)].width=min(max(max_len+2,11),36)
    out=io.BytesIO(); wb.save(out); out.seek(0)
    filename=f'Burjeel_ED_Call_Bell_KPI_{d["period"]}_{now_utc().strftime("%Y%m%d_%H%M")}.xlsx'
    return Response(out.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{filename}"'})

@app.get('/manager/export.pdf')
def manager_export_pdf(request:Request,db:Session=Depends(get_db)):
    require_permission(request,db,'export_reports'); d=_management_data(request,db)
    out=io.BytesIO(); doc=SimpleDocTemplate(out,pagesize=landscape(A4),rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=24); styles=getSampleStyleSheet(); story=[]
    story.append(Paragraph('Burjeel ED Smart Call Bell - KPI Management Report',styles['Title'])); story.append(Paragraph(d['range_label'],styles['Normal'])); story.append(Spacer(1,10))
    k=d['kpi']; kdata=[['KPI','Value'],['Total Calls',k['total']],['Average Response',k['avg_response']],['Median Response',k['median_response']],['Calls >2 min',k['over_2m']],['Calls >5 min',k['over_5m']],['Escalations',k['escalations']],['Charge Takeovers',k['takeovers']],['SLA Compliance',k['sla_rate']]]
    kt=Table(kdata,colWidths=[180,100]); kt.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#0B2A4A')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#D0D5DD')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F8FAFC')]),('PADDING',(0,0),(-1,-1),5)])); story.append(kt); story.append(Spacer(1,12))
    story.append(Paragraph('Nurse Performance',styles['Heading2']))
    ndata=[['Nurse','Calls','Avg Response','Over SLA','Escalated']]+[[n['name'],n['calls'],n['avg_response'],n['over_sla'],n['escalated']] for n in d['nurse_perf']]
    nt=Table(ndata,repeatRows=1); nt.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#0B67C2')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#D0D5DD')),('PADDING',(0,0),(-1,-1),4)])); story.append(nt); story.append(Spacer(1,12))
    story.append(Paragraph('Recent Call Lifecycle',styles['Heading2']))
    cdata=[['Room','Created','Nurse','Status','Response','Escalated','Takeover','Reason']]+[[c['room'],c['created'],c['nurse'],c['status_label'],c['response'],c['escalated'],c['takeover'],c['reason']] for c in d['call_rows'][:60]]
    ct=Table(cdata,repeatRows=1,colWidths=[50,100,80,70,55,55,75,110]); ct.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#0B2A4A')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#D0D5DD')),('FONTSIZE',(0,0),(-1,-1),7),('PADDING',(0,0),(-1,-1),3)])); story.append(ct)
    doc.build(story); out.seek(0)
    filename=f'Burjeel_ED_Call_Bell_KPI_{d["period"]}_{now_utc().strftime("%Y%m%d_%H%M")}.pdf'
    return Response(out.getvalue(),media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="{filename}"'})

def wallboard_payload(db:Session):
    enforce_escalations(db)
    active_status=['new','acknowledged','escalated','taken_over','arrived']
    calls=db.query(Call).filter(Call.status.in_(active_status)).order_by(Call.created_at).all()
    now=now_utc(); call_out=[]; call_by_room={}
    for c in calls:
        created=c.created_at.replace(tzinfo=timezone.utc) if c.created_at and c.created_at.tzinfo is None else c.created_at
        elapsed=max(0,int((now-created).total_seconds())) if created else 0
        if c.status=='taken_over': alert_stage='takeover'
        elif c.status=='arrived': alert_stage='arrived'
        elif c.status=='acknowledged': alert_stage='acknowledged'
        elif c.status=='escalated': alert_stage='critical'
        elif elapsed<=SLA_SECONDS: alert_stage='fresh'
        elif elapsed<=SLA_SECONDS*2: alert_stage='warning'
        else: alert_stage='critical'
        stop_at=c.arrived_at or c.resolved_at
        stop_at=stop_at.replace(tzinfo=timezone.utc) if stop_at and stop_at.tzinfo is None else stop_at
        display_elapsed=max(0,int((stop_at-created).total_seconds())) if stop_at and created else elapsed
        item={'id':c.id,'room':c.room.code,'room_id':c.room_id,'zone':c.room.zone,'status':c.status,'reason':c.reason,'nurse':c.assigned_nurse.name if c.assigned_nurse else 'Unassigned','takeover':c.taken_over_by.name if c.taken_over_by else None,'elapsed_seconds':display_elapsed,'elapsed_label':f'{display_elapsed//60:02d}:{display_elapsed%60:02d}','created_at':created.isoformat() if created else '','timer_stop':stop_at.isoformat() if stop_at else '','sla_progress':min(100,int((display_elapsed/max(1,SLA_SECONDS*2))*100)),'over_sla':c.status in ['new','escalated'] and elapsed>SLA_SECONDS,'alert_stage':alert_stage}
        call_out.append(item); call_by_room[c.room_id]=item
    rooms=db.query(Room).order_by(Room.zone,Room.code).all(); rooms_view=[]
    for r in rooms:
        c=call_by_room.get(r.id)
        if c: rooms_view.append(dict(c))
        else: rooms_view.append({'id':None,'room':r.code,'room_id':r.id,'zone':r.zone,'status':'ready' if r.occupied else 'closed','reason':'Ready' if r.occupied else 'Closed','nurse':r.assigned_nurse.name if r.assigned_nurse else 'Unassigned','takeover':None,'elapsed_seconds':0,'elapsed_label':'--:--','created_at':'','timer_stop':'','sla_progress':0,'over_sla':False,'alert_stage':'idle'})
    recent_calls=db.query(Call).filter(Call.resolved_at.isnot(None)).order_by(Call.resolved_at.desc()).limit(6).all()
    recent=[]
    for c in recent_calls:
        s=c.created_at.replace(tzinfo=timezone.utc) if c.created_at and c.created_at.tzinfo is None else c.created_at
        e=c.arrived_at or c.resolved_at; e=e.replace(tzinfo=timezone.utc) if e and e.tzinfo is None else e
        sec=max(0,int((e-s).total_seconds())) if s and e else 0
        recent.append({'room':c.room.code,'response':f'{sec//60:02d}:{sec%60:02d}','nurse':(c.taken_over_by.name if c.taken_over_by else (c.assigned_nurse.name if c.assigned_nurse else 'Unassigned'))})
    nurses=db.query(User).filter_by(role='nurse',active=True).order_by(User.name).all(); workload=[]
    for n in nurses:
        workload.append({'name':n.name,'rooms':db.query(Room).filter_by(assigned_nurse_id=n.id,occupied=True).count(),'active':sum(1 for c in calls if c.assigned_nurse_id==n.id)})
    hs=db.query(Handover).filter_by(status='pending').order_by(Handover.created_at.desc()).limit(8).all()
    handovers=[{'room':h.room.code,'from_nurse':h.from_nurse.name,'to_nurse':h.to_nurse.name} for h in hs]
    rt=[]
    for c in recent_calls:
        if c.created_at and (c.arrived_at or c.resolved_at):
            s=c.created_at.replace(tzinfo=timezone.utc) if c.created_at.tzinfo is None else c.created_at; e=c.arrived_at or c.resolved_at; e=e.replace(tzinfo=timezone.utc) if e.tzinfo is None else e; rt.append((e-s).total_seconds())
    avg=int(sum(rt)/len(rt)) if rt else 0
    stats={'total_rooms':len(rooms),'active':len(calls),'escalated':sum(1 for c in calls if c.status=='escalated'),'over_sla':sum(1 for c in call_out if c['over_sla']),'avg_response':f'{avg//60:02d}:{avg%60:02d}','unassigned':db.query(Room).filter(Room.occupied==True,Room.assigned_nurse_id.is_(None)).count()}
    return {'sound_repeat_seconds':escalation_settings(db)['sound_repeat'],'calls':call_out,'rooms_view':rooms_view,'recent':recent,'workload':workload,'handovers':handovers,'stats':stats,'zones':[z[0] for z in db.query(Room.zone).distinct().order_by(Room.zone).all()]}

@app.get('/wallboard',response_class=HTMLResponse)
def wallboard_page(request:Request,db:Session=Depends(get_db)):
    u=require_permission(request,db,'wallboard')
    lang=request.query_params.get('lang','en').lower()
    if lang not in ['en','ar']: lang='en'
    p=wallboard_payload(db)
    i18n={
      'en':{'kicker':'Nurse Station Wallboard','title':'ED Call Bell Live Screen','subtitle':'Room calls, SLA escalation, nurse workload and handover status.','enable_sound':'Enable sound','sound_on':'Sound on','tap_sound':'Tap for sound','fullscreen':'Full screen','total_rooms':'Rooms','active_calls':'Active calls','escalated':'Escalated','over_sla':'Over SLA','avg_response':'Avg response','unassigned':'Unassigned rooms','all':'All','connected':'Connected','disconnected':'Offline','no_calls':'No active call bell requests','assigned_nurse':'Assigned nurse','elapsed':'Elapsed','charge_takeover':'Charge takeover','nurse_workload':'Nurse workload','pending_handover':'Pending handover','recent_resolved':'Recent resolved','rooms':'rooms','calls':'calls','none':'None','compact':'Compact','cards':'Cards','list':'List','zones':'Zones','show_idle':'Show ready rooms','legend_green':'0–2 min','legend_amber':'2–4 min','legend_red':'4+ min / escalated','legend_blue':'Charge takeover','status':{'new':'NEW CALL','acknowledged':'ACKNOWLEDGED','escalated':'ESCALATED','taken_over':'CHARGE TAKEOVER','arrived':'ARRIVED','ready':'READY','closed':'CLOSED'},'reason':{'General assistance':'General assistance','Pain':'Pain','Toilet assistance':'Toilet assistance','IV / Medication':'IV / Medication','Ready':'Ready','Closed':'Closed'}},
      'ar':{'kicker':'شاشة محطة التمريض','title':'لوحة نداءات الطوارئ المباشرة','subtitle':'نداءات الغرف والتصعيد وعبء التمريض وحالات التسليم.','enable_sound':'تفعيل الصوت','sound_on':'الصوت مفعل','tap_sound':'اضغط لتفعيل الصوت','fullscreen':'ملء الشاشة','total_rooms':'الغرف','active_calls':'النداءات النشطة','escalated':'تم التصعيد','over_sla':'تجاوز الوقت','avg_response':'متوسط الاستجابة','unassigned':'غرف بدون ممرضة','all':'الكل','connected':'متصل','disconnected':'غير متصل','no_calls':'لا توجد نداءات نشطة','assigned_nurse':'الممرضة المسؤولة','elapsed':'الوقت','charge_takeover':'استلام مسؤول التمريض','nurse_workload':'عبء التمريض','pending_handover':'تسليمات معلقة','recent_resolved':'آخر النداءات المغلقة','rooms':'غرف','calls':'نداءات','none':'لا يوجد','compact':'مضغوط','cards':'بطاقات','list':'قائمة','zones':'مناطق','show_idle':'إظهار الغرف الجاهزة','legend_green':'0–2 دقيقة','legend_amber':'2–4 دقائق','legend_red':'4+ دقائق / تصعيد','legend_blue':'استلام مسؤول التمريض','status':{'new':'نداء جديد','acknowledged':'تم التأكيد','escalated':'تم التصعيد','taken_over':'استلام المسؤول','arrived':'تم الوصول','ready':'جاهزة','closed':'مغلقة'},'reason':{'General assistance':'مساعدة عامة','Pain':'ألم','Toilet assistance':'مساعدة للحمام','IV / Medication':'المحلول / الدواء','Ready':'جاهزة','Closed':'مغلقة'}}
    }
    t=i18n[lang]; calls=[]; rooms_view=[]
    for c in p['calls']:
        x=dict(c); x['status_label']=t['status'].get(c['status'],c['status']); x['reason_label']=t['reason'].get(c['reason'],c['reason']); calls.append(x)
    for r in p['rooms_view']:
        x=dict(r); x['status_label']=t['status'].get(r['status'],r['status']); x['reason_label']=t['reason'].get(r['reason'],r['reason']); rooms_view.append(x)
    return render_template('wallboard.html',ctx(request,db,lang=lang,t=t,calls=calls,rooms_view=rooms_view,recent=p['recent'],workload=p['workload'],handovers=p['handovers'],stats=p['stats'],zones=p['zones'],sound_repeat_seconds=p['sound_repeat_seconds'],wall_i18n_json=json.dumps(t,ensure_ascii=False),call_ids_json=json.dumps([c['id'] for c in p['calls']]),call_stages_json=json.dumps({str(c['id']):c['alert_stage'] for c in p['calls']})))

@app.get('/api/wallboard')
def wallboard_api(request:Request,db:Session=Depends(get_db)):
    require_role(request,db,['nurse','charge','nurse_supervisor','manager','ed_manager','hod','admin']); return wallboard_payload(db)

@app.post('/api/client-error')
async def client_error_report(request:Request):
    try:
        data=await request.json()
        path=str(data.get('path') or '/client')[:240]
        source=str(data.get('source') or 'browser-js')[:120]
        message=str(data.get('message') or 'Client-side JavaScript error')
        exc_type=str(data.get('type') or 'JavaScriptError')[:120]
        trace_text=str(data.get('stack') or '')
        _record_runtime_event(source,path,exc_type,message,trace_text,'client-redacted')
    except Exception:
        pass
    return {'ok':True}

@app.get('/admin/diagnostics',response_class=HTMLResponse)
def admin_diagnostics(request:Request,db:Session=Depends(get_db)):
    require_permission(request,db,'admin_control')
    status=request.query_params.get('status','all').lower()
    if status not in ['all','open','recovered']: status='all'
    q=request.query_params.get('q','').strip().lower()
    try: rows_limit=int(request.query_params.get('rows','250'))
    except Exception: rows_limit=250
    if rows_limit not in [50,100,250]: rows_limit=250
    all_rows=db.query(RuntimeErrorEvent).order_by(RuntimeErrorEvent.last_seen.desc()).all()
    open_errors=sum(1 for x in all_rows if x.status=='open')
    recovered=sum(1 for x in all_rows if x.status=='recovered')
    captured=sum(int(x.count or 0) for x in all_rows)
    filtered=all_rows
    if status!='all': filtered=[x for x in filtered if x.status==status]
    if q:
        filtered=[x for x in filtered if q in ' '.join([x.error_id or '',x.source or '',x.request_path or '',x.exception_type or '',x.message or '']).lower()]
    events=filtered[:rows_limit]
    schema_ready=True
    try: db.query(RuntimeErrorEvent).limit(1).all()
    except Exception: schema_ready=False
    return render_template('diagnostics.html',ctx(request,db,events=events,open_errors=open_errors,recovered=recovered,captured=captured,schema_ready=schema_ready,schema_time=now_utc().strftime('%Y-%m-%dT%H:%M:%SZ'),q=q,status=status,rows_limit=rows_limit))

@app.get('/admin/diagnostics/application-log')
def admin_diagnostics_log(request:Request,db:Session=Depends(get_db)):
    require_permission(request,db,'admin_control')
    rows=db.query(RuntimeErrorEvent).order_by(RuntimeErrorEvent.last_seen.desc()).limit(250).all()
    lines=[]
    for e in rows:
        lines.append(f'{e.last_seen.isoformat() if e.last_seen else ""} [{e.status.upper()}] {e.error_id} count={e.count} source={e.source} request={e.request_path} exception={e.exception_type} message={e.message}')
    return Response('\n'.join(lines) if lines else 'No captured runtime errors.\n',media_type='text/plain',headers={'Content-Disposition':'attachment; filename="burjeel-ed-application-log.txt"'})

@app.get('/admin/diagnostics/runtime-report')
def admin_diagnostics_report(request:Request,db:Session=Depends(get_db)):
    require_permission(request,db,'admin_control')
    rows=db.query(RuntimeErrorEvent).all()
    return JSONResponse({'generated_at':now_utc().isoformat(),'service':'Burjeel ED Smart Call','schema_gate':'READY','database':'connected','open_errors':sum(1 for x in rows if x.status=='open'),'recovered_automatically':sum(1 for x in rows if x.status=='recovered'),'captured_events':sum(int(x.count or 0) for x in rows),'push_subscriptions':db.query(PushSubscription).count(),'push_devices':db.query(PushDevice).count(),'active_calls':db.query(Call).filter(Call.status.in_(['new','acknowledged','escalated','taken_over','arrived'])).count()})

@app.get('/admin/diagnostics/download')
def admin_diagnostics_download(request:Request,db:Session=Depends(get_db)):
    require_permission(request,db,'admin_control')
    status=request.query_params.get('status','all').lower(); q=request.query_params.get('q','').strip().lower()
    try: limit=int(request.query_params.get('rows','250'))
    except Exception: limit=250
    limit=limit if limit in [50,100,250] else 250
    rows=db.query(RuntimeErrorEvent).order_by(RuntimeErrorEvent.last_seen.desc()).all()
    if status in ['open','recovered']: rows=[x for x in rows if x.status==status]
    if q: rows=[x for x in rows if q in ' '.join([x.error_id or '',x.source or '',x.request_path or '',x.exception_type or '',x.message or '']).lower()]
    payload=[]
    for e in rows[:limit]:
        payload.append(json.dumps({'time':e.last_seen.isoformat() if e.last_seen else None,'status':e.status,'count':e.count,'error_id':e.error_id,'source':e.source,'request':e.request_path,'user':e.user_ref,'exception':e.exception_type,'message':e.message,'trace':e.trace},ensure_ascii=False))
    return Response('\n'.join(payload)+(('\n') if payload else ''),media_type='application/x-ndjson',headers={'Content-Disposition':'attachment; filename="burjeel-ed-errors.jsonl"'})

@app.post('/admin/diagnostics/clear')
def admin_diagnostics_clear(request:Request,db:Session=Depends(get_db)):
    admin=require_permission(request,db,'admin_control')
    count=db.query(RuntimeErrorEvent).delete(synchronize_session=False); log_action(db,'ADMIN_DIAGNOSTICS_CLEARED',f'Cleared {count} monitor events',user=admin); db.commit()
    return RedirectResponse('/admin/diagnostics',303)

@app.get('/admin',response_class=HTMLResponse)
def admin_page(request:Request,db:Session=Depends(get_db)):
    require_permission(request,db,'admin_control')
    users=db.query(User).order_by(User.role,User.name).all()
    rooms=db.query(Room).order_by(Room.code).all()
    nurses=db.query(User).filter_by(role='nurse',active=True).order_by(User.name).all()
    audit_logs=db.query(AuditLog).order_by(AuditLog.created_at.desc()).limit(100).all()
    active_calls=db.query(Call).filter(Call.status.in_(['new','acknowledged','escalated','taken_over','arrived'])).count()
    push_count=db.query(PushSubscription).count()
    push_devices=db.query(PushDevice).order_by(PushDevice.last_seen.desc()).all()
    push_healthy=sum(1 for d in push_devices if d.health_status in ['active','registered'])
    push_attention=sum(1 for d in push_devices if d.health_status not in ['active','registered'])
    active_staff=sum(1 for u in users if u.active)
    active_rooms=sum(1 for r in rooms if r.occupied)
    open_errors=db.query(RuntimeErrorEvent).filter_by(status='open').count()
    access_map={x.user_id:x for x in db.query(UserAccessProfile).all()}
    permission_rows=[]
    for usr in users:
        rec=access_map.get(usr.id)
        profile_name=rec.profile_name if rec else ROLE_DEFAULT_PROFILE.get(usr.role,'bedside_nurse')
        permission_rows.append({'user':usr,'role_label':ROLE_LABELS.get(usr.role,usr.role.replace('_',' ').title()),'profile_name':profile_name,'permissions':sorted(get_user_permissions(db,usr))})
    profile_options=[{'key':k,'label':v['label']} for k,v in WORK_PROFILE_PRESETS.items()]
    permission_defs=[{'key':k,'label':label,'description':desc} for k,label,desc in PERMISSION_DEFS]
    permission_presets_json=json.dumps({k:v['permissions'] for k,v in WORK_PROFILE_PRESETS.items()})
    return render_template('admin.html',ctx(request,db,escalation_saved=bool(request.session.pop('escalation_saved',False)),escalation_settings=escalation_settings(db),rooms=rooms,users=users,nurses=nurses,audit_logs=audit_logs,active_calls=active_calls,push_count=push_count,push_devices=push_devices,push_healthy=push_healthy,push_attention=push_attention,active_staff=active_staff,active_rooms=active_rooms,open_errors=open_errors,permission_rows=permission_rows,profile_options=profile_options,permission_defs=permission_defs,permission_presets_json=permission_presets_json,vapid_ready=bool(VAPID_PUBLIC_KEY and VAPID_PRIVATE_KEY),room_msg=request.query_params.get('room_msg',''),room_msg_type=request.query_params.get('room_msg_type','success')))

@app.post('/admin/users/{uid}/permissions')
async def admin_user_permissions(uid:int,request:Request,db:Session=Depends(get_db)):
    admin=require_permission(request,db,'user_permissions')
    usr=db.get(User,uid)
    if not usr: raise HTTPException(404)
    if usr.role=='admin' and usr.id==admin.id:
        return RedirectResponse('/admin#permissions',303)
    form=await request.form()
    profile_name=str(form.get('profile_name','')).strip()
    if profile_name not in WORK_PROFILE_PRESETS: profile_name=ROLE_DEFAULT_PROFILE.get(usr.role,'bedside_nurse')
    selected=[p for p,_,_ in PERMISSION_DEFS if p in form.getlist('permissions')]
    rec=db.query(UserAccessProfile).filter_by(user_id=usr.id).first()
    if not rec:
        rec=UserAccessProfile(user_id=usr.id,profile_name=profile_name,permissions_json='[]',updated_at=now_utc()); db.add(rec)
    rec.profile_name=profile_name; rec.permissions_json=json.dumps(selected); rec.updated_at=now_utc()
    log_action(db,'USER_PERMISSIONS_UPDATED',f'{usr.email}; profile={profile_name}; permissions={",".join(selected)}',user=admin); db.commit()
    return RedirectResponse('/admin#permissions',303)

@app.post('/admin/users/{uid}/permissions/reset')
def admin_user_permissions_reset(uid:int,request:Request,db:Session=Depends(get_db)):
    admin=require_permission(request,db,'user_permissions')
    usr=db.get(User,uid)
    if not usr: raise HTTPException(404)
    rec=db.query(UserAccessProfile).filter_by(user_id=usr.id).first()
    if rec: db.delete(rec)
    log_action(db,'USER_PERMISSIONS_RESET',f'{usr.email}; restored role defaults',user=admin); db.commit()
    return RedirectResponse('/admin#permissions',303)

@app.post('/admin/users')
async def admin_create_user(request:Request,db:Session=Depends(get_db)):
    admin=require_permission(request,db,'admin_control'); form=await request.form()
    name=str(form.get('name','')).strip(); email=str(form.get('email','')).strip().lower(); role=str(form.get('role','')).strip(); password=str(form.get('password',''))
    if role not in ['nurse','charge','nurse_supervisor','manager','ed_manager','hod','admin'] or not name or not email or len(password)<8: raise HTTPException(400)
    if db.query(User).filter_by(email=email).first(): return RedirectResponse('/admin#staff',303)
    u=User(name=name,email=email,role=role,password_hash=pwd_hash(password),active=True); db.add(u); db.flush(); log_action(db,'ADMIN_USER_CREATED',f'{email} role={role}',user=admin); db.commit()
    return RedirectResponse('/admin#staff',303)

@app.post('/admin/users/{uid}/toggle')
def admin_toggle_user(uid:int,request:Request,db:Session=Depends(get_db)):
    admin=require_permission(request,db,'admin_control'); u=db.get(User,uid)
    if not u: raise HTTPException(404)
    if u.id==admin.id: return RedirectResponse('/admin#staff',303)
    u.active=not u.active; log_action(db,'ADMIN_USER_TOGGLED',f'{u.email} active={u.active}',user=admin); db.commit()
    return RedirectResponse('/admin#staff',303)

@app.post('/admin/users/{uid}/reset')
async def admin_reset_password(uid:int,request:Request,db:Session=Depends(get_db)):
    admin=require_permission(request,db,'admin_control'); u=db.get(User,uid)
    if not u: raise HTTPException(404)
    form=await request.form(); password=str(form.get('password',''))
    if len(password)<8: raise HTTPException(400)
    u.password_hash=pwd_hash(password); log_action(db,'ADMIN_PASSWORD_RESET',u.email,user=admin); db.commit()
    return RedirectResponse('/admin#staff',303)

@app.post('/admin/rooms')
async def admin_create_room(request:Request,db:Session=Depends(get_db)):
    admin=require_permission(request,db,'admin_control'); form=await request.form()
    code=str(form.get('code','')).strip().upper(); zone=str(form.get('zone','')).strip() or 'ED Main'; nurse_id=str(form.get('nurse_id','')).strip()
    if not code or db.query(Room).filter_by(code=code).first(): return RedirectResponse('/admin#rooms',303)
    nurse=db.get(User,int(nurse_id)) if nurse_id else None
    if nurse and nurse.role!='nurse': nurse=None
    room=Room(code=code,zone=zone,qr_token=secrets.token_urlsafe(18),occupied=True,assigned_nurse_id=nurse.id if nurse else None); db.add(room); db.flush(); log_action(db,'ADMIN_ROOM_CREATED',f'{code} zone={zone}',room=room,user=admin); db.commit()
    return RedirectResponse('/admin#rooms',303)

@app.post('/admin/rooms/{room_id}/edit')
async def admin_edit_room(room_id:int,request:Request,db:Session=Depends(get_db)):
    admin=require_permission(request,db,'admin_control'); room=db.get(Room,room_id)
    if not room: raise HTTPException(404)
    form=await request.form()
    code=str(form.get('code','')).strip().upper()
    zone=str(form.get('zone','')).strip()
    if not code or not zone:
        return RedirectResponse('/admin?room_msg=Room+name+and+zone+are+required&room_msg_type=error#rooms',303)
    duplicate=db.query(Room).filter(Room.code==code,Room.id!=room.id).first()
    if duplicate:
        return RedirectResponse('/admin?room_msg=Room+name+already+exists&room_msg_type=error#rooms',303)
    old_code,old_zone=room.code,room.zone
    room.code=code; room.zone=zone
    log_action(db,'ADMIN_ROOM_EDITED',f'{old_code}/{old_zone} -> {code}/{zone}',room=room,user=admin)
    db.commit()
    return RedirectResponse('/admin?room_msg=Room+updated+successfully#rooms',303)

@app.post('/admin/rooms/{room_id}/delete')
def admin_delete_room(room_id:int,request:Request,db:Session=Depends(get_db)):
    admin=require_permission(request,db,'admin_control'); room=db.get(Room,room_id)
    if not room: raise HTTPException(404)
    active=db.query(Call).filter_by(room_id=room.id).filter(Call.status.in_(['new','acknowledged','escalated','taken_over','arrived'])).first()
    if active:
        return RedirectResponse('/admin?room_msg=Cannot+delete+a+room+with+an+active+call&room_msg_type=error#rooms',303)
    if db.query(Call).filter_by(room_id=room.id).first() or db.query(Handover).filter_by(room_id=room.id).first():
        return RedirectResponse('/admin?room_msg=Room+has+call+or+handover+history.+Close+it+instead+to+preserve+audit+records&room_msg_type=error#rooms',303)
    code,zone=room.code,room.zone
    db.query(AuditLog).filter_by(room_id=room.id).update({AuditLog.room_id:None},synchronize_session=False)
    db.delete(room)
    db.add(AuditLog(action='ADMIN_ROOM_DELETED',detail=f'{code} zone={zone}',user_id=admin.id))
    db.commit()
    return RedirectResponse('/admin?room_msg=Room+deleted+successfully#rooms',303)

@app.get('/admin/rooms/{room_id}/qr',response_class=HTMLResponse)
def admin_room_qr(room_id:int,request:Request,db:Session=Depends(get_db)):
    require_permission(request,db,'admin_control'); room=db.get(Room,room_id)
    if not room: raise HTTPException(404)
    patient_url=str(request.base_url).rstrip('/')+f'/room/{room.qr_token}'
    image=qrcode.make(patient_url,image_factory=qrcode.image.svg.SvgPathImage)
    buf=io.BytesIO(); image.save(buf); qr_data=base64.b64encode(buf.getvalue()).decode('ascii')
    auto_print=request.query_params.get('print')=='1'
    return render_template('qr.html',ctx(request,db,room=room,qr_data=qr_data,patient_url=patient_url,auto_print=auto_print))

@app.post('/admin/rooms/{room_id}/toggle')
def admin_toggle_room(room_id:int,request:Request,db:Session=Depends(get_db)):
    admin=require_permission(request,db,'admin_control'); room=db.get(Room,room_id)
    if not room: raise HTTPException(404)
    room.occupied=not room.occupied; log_action(db,'ADMIN_ROOM_TOGGLED',f'{room.code} occupied={room.occupied}',room=room,user=admin); db.commit()
    return RedirectResponse('/admin#rooms',303)

@app.post('/admin/rooms/{room_id}/assign')
async def admin_assign_room(room_id:int,request:Request,db:Session=Depends(get_db)):
    admin=require_permission(request,db,'admin_control'); room=db.get(Room,room_id)
    if not room: raise HTTPException(404)
    form=await request.form(); nurse_id=str(form.get('nurse_id','')).strip(); nurse=db.get(User,int(nurse_id)) if nurse_id else None
    if nurse and nurse.role!='nurse': raise HTTPException(400)
    old=room.assigned_nurse.name if room.assigned_nurse else 'Unassigned'; room.assigned_nurse_id=nurse.id if nurse else None
    log_action(db,'ADMIN_ROOM_ASSIGNED',f'{room.code}: {old} -> {nurse.name if nurse else "Unassigned"}',room=room,user=admin); db.commit()
    if nurse: send_push_to_user(db,nurse.id,f'Room assigned - {room.code}',f'{room.code} ({room.zone}) is now under your responsibility.','/nurse')
    return RedirectResponse('/admin#rooms',303)

@app.post('/admin/rooms/{room_id}/token')
def admin_regenerate_room_token(room_id:int,request:Request,db:Session=Depends(get_db)):
    admin=require_permission(request,db,'admin_control'); room=db.get(Room,room_id)
    if not room: raise HTTPException(404)
    room.qr_token=secrets.token_urlsafe(18); log_action(db,'ADMIN_ROOM_TOKEN_REGENERATED',room.code,room=room,user=admin); db.commit()
    return RedirectResponse('/admin#rooms',303)
def call_action(call_id:int,action:str,request:Request,db:Session=Depends(get_db)):
    u=require_role(request,db,['nurse','charge','nurse_supervisor','admin']); c=db.get(Call,call_id)
    if not c: raise HTTPException(404)
    if action=='ack':
        if u.role=='nurse' and c.assigned_nurse_id!=u.id: raise HTTPException(403)
        c.status='acknowledged'; c.acknowledged_at=c.acknowledged_at or now_utc(); log_action(db,'CALL_ACKNOWLEDGED',f'By {u.name}',call=c,user=u)
    elif action=='takeover':
        if 'takeover_calls' not in get_user_permissions(db,u): raise HTTPException(403)
        c.status='taken_over'; c.taken_over_by_id=u.id
        if not c.acknowledged_at: c.acknowledged_at=now_utc()
        if not c.escalated_at:c.escalated_at=now_utc();c.escalation_reason='Charge nurse manual takeover'
        log_action(db,'CALL_TAKEN_OVER',f'By {u.name}; primary={c.assigned_nurse.name if c.assigned_nurse else "Unassigned"}',call=c,user=u)
    elif action=='arrive':
        if not(u.id in [c.assigned_nurse_id,c.taken_over_by_id] or 'takeover_calls' in get_user_permissions(db,u)): raise HTTPException(403)
        arrived=now_utc()
        c.arrived_at=arrived
        if not c.acknowledged_at: c.acknowledged_at=arrived
        c.resolved_at=arrived
        c.status='resolved'
        log_action(db,'NURSE_ARRIVED',f'By {u.name}; call auto-closed on arrival',call=c,user=u)
        log_action(db,'CALL_AUTO_RESOLVED',f'Auto-closed at arrival by {u.name}',call=c,user=u)
    elif action=='resolve':
        if not(u.id in [c.assigned_nurse_id,c.taken_over_by_id] or u.role in ['charge','manager','ed_manager','admin']): raise HTTPException(403)
        c.status='resolved'; c.resolved_at=now_utc(); log_action(db,'CALL_RESOLVED',f'By {u.name}',call=c,user=u)
    else: raise HTTPException(400)
    db.commit(); return {'ok':True,'call':serialize_call(c)}
@app.post('/api/call/{call_id}/reassign')
async def reassign_active_call(call_id:int,request:Request,db:Session=Depends(get_db)):
    supervisor=require_permission(request,db,'reassign_calls')
    c=db.get(Call,call_id)
    if not c: raise HTTPException(404)
    if c.status not in ['new','acknowledged','escalated','taken_over']:
        return JSONResponse({'ok':False,'error':'Only an open call that has not reached Arrived can be reassigned.'},409)
    data=await request.json()
    try: nurse_id=int(data.get('nurse_id'))
    except: return JSONResponse({'ok':False,'error':'Select a nurse.'},400)
    reason=str(data.get('reason','')).strip()
    if reason not in ['Workload balancing','Break','Shift change','No response','Clinical priority','Other']:
        return JSONResponse({'ok':False,'error':'Select a valid reassignment reason.'},400)
    nurse=db.get(User,nurse_id)
    if not nurse or nurse.role!='nurse' or not nurse.active:
        return JSONResponse({'ok':False,'error':'Selected nurse is not active.'},400)
    if c.assigned_nurse_id==nurse.id:
        return JSONResponse({'ok':False,'error':'This nurse is already assigned.'},409)
    old_nurse=c.assigned_nurse
    old_name=old_nurse.name if old_nurse else 'Unassigned'
    old_id=c.assigned_nurse_id
    c.assigned_nurse_id=nurse.id
    c.room.assigned_nurse_id=nurse.id
    if c.status in ['new','acknowledged']:
        c.status='new'; c.acknowledged_at=None
    elif c.status in ['escalated','taken_over']:
        c.status='escalated'; c.taken_over_by_id=None
    detail=f'{c.room.code}: {old_name} -> {nurse.name}; reason={reason}; by={supervisor.name}; original_call_time={c.created_at.isoformat()}'
    log_action(db,'CALL_REASSIGNED',detail,room=c.room,call=c,user=supervisor)
    db.commit()
    send_push_to_user(db,nurse.id,f'Reassigned Call - {c.room.code}',f'{reason}. Patient call requires your response.','/nurse')
    if old_id:
        send_push_to_user(db,old_id,f'Call Reassigned - {c.room.code}',f'Call moved to {nurse.name} by {supervisor.name}.','/nurse')
    return {'ok':True,'call':serialize_call(c),'from_nurse':old_name,'to_nurse':nurse.name,'reason':reason,'timer_preserved':True}

app.post('/api/call/{call_id}/{action}')(call_action)

@app.post('/api/room/{room_id}/assign')
async def assign_room(room_id:int,request:Request,db:Session=Depends(get_db)):
    require_permission(request,db,'assign_rooms'); data=await request.json(); room=db.get(Room,room_id); nurse=db.get(User,int(data['nurse_id']))
    if not room or not nurse: raise HTTPException(404)
    if nurse.role!='nurse': return JSONResponse({'ok':False,'error':'Only nurse users can be assigned'},400)
    old=room.assigned_nurse; room.assigned_nurse_id=nurse.id; log_action(db,'ROOM_ASSIGNED',f'{room.code}: {old.name if old else "Unassigned"} -> {nurse.name}',room=room); db.commit()
    send_push_to_user(db,nurse.id,f'Room assigned - {room.code}',f'{room.code} ({room.zone}) is now under your responsibility.','/nurse')
    if old and old.id!=nurse.id: send_push_to_user(db,old.id,f'Room reassigned - {room.code}',f'{room.code} was reassigned to {nurse.name}.','/nurse')
    return {'ok':True}

@app.post('/api/handover')
async def create_handover(request:Request,db:Session=Depends(get_db)):
    u=require_role(request,db,['nurse','charge','nurse_supervisor','admin']); data=await request.json(); room=db.get(Room,int(data['room_id'])); to_nurse=db.get(User,int(data['to_nurse_id']))
    if not room or not to_nurse: raise HTTPException(404)
    if to_nurse.role!='nurse': return JSONResponse({'ok':False,'error':'Target must be a nurse'},400)
    if u.role=='nurse' and room.assigned_nurse_id!=u.id: raise HTTPException(403)
    existing=db.query(Handover).filter_by(room_id=room.id,status='pending').first()
    if existing: return {'ok':True,'handover_id':existing.id,'duplicate':True}
    h=Handover(room_id=room.id,from_nurse_id=room.assigned_nurse_id or u.id,to_nurse_id=to_nurse.id,status='pending'); db.add(h); db.flush(); log_action(db,'HANDOVER_CREATED',f'{room.code} -> {to_nurse.name}',room=room,user=u); db.commit(); send_push_to_user(db,to_nurse.id,f'Handover request - {room.code}',f'{u.name} wants to hand over {room.code} ({room.zone}) to you. Open My Rooms to accept.','/nurse'); return {'ok':True,'handover_id':h.id}

@app.post('/api/handover/{hid}/accept')
def accept_handover(hid:int,request:Request,db:Session=Depends(get_db)):
    u=require_role(request,db,['nurse','charge','nurse_supervisor','admin']); h=db.get(Handover,hid)
    if not h: raise HTTPException(404)
    if u.role=='nurse' and h.to_nurse_id!=u.id: raise HTTPException(403)
    h.status='accepted'; h.accepted_at=now_utc(); h.room.assigned_nurse_id=h.to_nurse_id; log_action(db,'HANDOVER_ACCEPTED',f'{h.room.code}: {h.from_nurse.name} -> {h.to_nurse.name}',room=h.room,user=u); db.commit()
    send_push_to_user(db,h.from_nurse_id,f'Handover accepted - {h.room.code}',f'{h.to_nurse.name} accepted responsibility for {h.room.code}.','/nurse')
    return {'ok':True}

@app.post('/api/push/health')
async def push_health(request:Request,db:Session=Depends(get_db)):
    u=require_role(request,db,['nurse','charge','nurse_supervisor','manager','ed_manager','hod','admin'])
    try: data=await request.json()
    except Exception: data={}
    endpoint=str(data.get('endpoint','')).strip(); meta=data.get('device') or {}
    total=db.query(PushSubscription).filter_by(user_id=u.id).count()
    exact=db.query(PushSubscription).filter_by(user_id=u.id,endpoint=endpoint).first() if endpoint else None
    dev=_upsert_push_device(db,u.id,endpoint,meta,verified=False,health='active' if exact else 'needs_resubscribe') if endpoint else None
    db.commit()
    return JSONResponse({'ok':True,'healthy':bool(exact),'subscriptions':total,'force_resubscribe':False,'needs_registration':bool(endpoint and not exact),'display_mode':dev.display_mode if dev else None},headers={'Cache-Control':'no-store'})

@app.post('/api/push/subscribe')
async def push_subscribe(request:Request,db:Session=Depends(get_db)):
    u=require_role(request,db,['nurse','charge','nurse_supervisor','manager','ed_manager','hod','admin']); data=await request.json()
    payload=data.get('subscription') if isinstance(data,dict) and data.get('subscription') else data
    meta=data.get('device') if isinstance(data,dict) else {}
    endpoint=payload.get('endpoint') if isinstance(payload,dict) else None
    if not endpoint:return JSONResponse({'ok':False},400)
    mode=str((meta or {}).get('display_mode') or 'browser'); platform=str((meta or {}).get('platform') or '')
    key=_push_device_key(endpoint,mode,platform)
    dev=db.query(PushDevice).filter_by(user_id=u.id,device_key=key).first()
    already_verified=bool(dev and dev.last_verified and dev.health_status!='expired')
    rec=db.query(PushSubscription).filter_by(endpoint=endpoint).first()
    if rec:rec.user_id=u.id;rec.payload=json.dumps(payload)
    else:db.add(PushSubscription(user_id=u.id,endpoint=endpoint,payload=json.dumps(payload)))
    _upsert_push_device(db,u.id,endpoint,meta,verified=already_verified,health='active')
    if already_verified: request.session['notification_verified']=True
    db.commit(); return {'ok':True,'verified':bool(request.session.get('notification_verified')),'display_mode':mode}
@app.get('/api/live')
def live(request:Request,db:Session=Depends(get_db)):
    u=require_role(request,db,['nurse','charge','nurse_supervisor','manager','ed_manager','hod','admin']); enforce_escalations(db); q=db.query(Call).filter(Call.status.in_(['new','acknowledged','escalated','taken_over','arrived'])); q=q.filter_by(assigned_nurse_id=u.id) if u.role=='nurse' else q; return [serialize_call(c) for c in q.order_by(Call.created_at).all()]
@app.get('/manifest.webmanifest')
def manifest(): return Response(MANIFEST_JSON,media_type='application/manifest+json')
@app.get('/sw.js')
def sw(): return Response(SERVICE_WORKER_JS,media_type='application/javascript',headers={'Service-Worker-Allowed':'/','Cache-Control':'no-cache'})

def seed_demo():
    db=SessionLocal()
    try:
        if db.query(User).first():return
        created={}
        for name,email,role in [('Nurse Sara','sara@demo.local','nurse'),('Nurse Mona','mona@demo.local','nurse'),('Charge Nurse Fatima','charge@demo.local','charge'),('Nurse Manager','manager@demo.local','manager'),('ED Department Manager','edmanager@demo.local','ed_manager'),('System Admin','admin@demo.local','admin')]:
            u=User(name=name,email=email,role=role,password_hash=pwd_hash('Demo123!'));db.add(u);db.flush();created[email]=u
        for i in range(1,9):db.add(Room(code=f'ED-{i:02d}',zone='ED Main' if i<=4 else 'Fast Track',qr_token=f'room-{i:02d}-secure-token',occupied=True,assigned_nurse_id=created['sara@demo.local'].id if i<=4 else created['mona@demo.local'].id))
        db.commit()
    finally:db.close()
seed_demo()
